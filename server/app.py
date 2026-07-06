#!/usr/bin/env python3
"""
AnalytIQ API  —  Flask + SQLite (sqlite3)

Run:   python server/app.py
Test:  python -m pytest tests/
API:   http://localhost:3001
"""
import functools
import json
import logging
import os

# Load .env from project root (one level up from server/)
from pathlib import Path
_env_path = Path(__file__).resolve().parent.parent / '.env'
if _env_path.exists():
    for _line in _env_path.read_text().splitlines():
        _line = _line.strip()
        if _line and not _line.startswith('#') and '=' in _line:
            _k, _, _v = _line.partition('=')
            os.environ[_k.strip()] = _v.strip()
import threading
import time
from datetime import date, datetime, timedelta
from queue import Empty, Queue

import sqlite3
import resend
import stripe
from cryptography.fernet import Fernet
from croniter import croniter
from flask import Flask, Response, g, jsonify, request, stream_with_context
from flask_cors import CORS
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address

log = logging.getLogger(__name__)

# ─────────────────────────────────────────────────────────
# App setup
# ─────────────────────────────────────────────────────────
app = Flask(__name__)
CORS(app, origins=['http://localhost:5173'], supports_credentials=True)

REDIS_URL = os.environ.get('REDIS_URL')
limiter = Limiter(
    get_remote_address,
    app=app,
    default_limits=['120/minute'],
    storage_uri=REDIS_URL or 'memory://',
)


@app.errorhandler(429)
def ratelimit_handler(e):
    return jsonify({'error': 'Rate limit exceeded'}), 429


@app.errorhandler(404)
def not_found_handler(e):
    return jsonify({'error': 'Not found'}), 404


@app.errorhandler(405)
def method_not_allowed_handler(e):
    return jsonify({'error': 'Method not allowed'}), 405


@app.errorhandler(500)
def internal_error_handler(e):
    return jsonify({'error': 'Internal server error'}), 500

_DB_PATH = os.environ.get('DATABASE_PATH', str(Path(__file__).parent.parent / 'analytiq.db'))
PORT = int(os.environ.get('PORT', 3001))

RESEND_API_KEY = os.environ.get('RESEND_API_KEY')
FROM_EMAIL = os.environ.get('FROM_EMAIL', 'AnalytIQ <noreply@analytiq.dev>')
if RESEND_API_KEY:
    resend.api_key = RESEND_API_KEY

STRIPE_SECRET_KEY = os.environ.get('STRIPE_SECRET_KEY')
STRIPE_WEBHOOK_SECRET = os.environ.get('STRIPE_WEBHOOK_SECRET')
STRIPE_PRO_PRICE_ID = os.environ.get('STRIPE_PRO_PRICE_ID')
if STRIPE_SECRET_KEY:
    stripe.api_key = STRIPE_SECRET_KEY

_credential_key = os.environ.get('CREDENTIAL_ENCRYPTION_KEY')
if not _credential_key:
    raise RuntimeError('CREDENTIAL_ENCRYPTION_KEY env var is required')
_fernet = Fernet(_credential_key.encode())



# ─────────────────────────────────────────────────────────
# Credential encryption helpers
# ─────────────────────────────────────────────────────────
def encrypt(plaintext: str | None) -> str | None:
    if plaintext is None:
        return None
    return _fernet.encrypt(plaintext.encode()).decode()


def decrypt(ciphertext: str | None) -> str | None:
    if ciphertext is None:
        return None
    return _fernet.decrypt(ciphertext.encode()).decode()


MASKED_PASSWORD = '••••••••'


# ─────────────────────────────────────────────────────────
# External connector test helpers
# ─────────────────────────────────────────────────────────
def _test_postgres(config):
    host = config.get('host', '')
    # Auto-apply SSL for known cloud database hosts
    _cloud = ('neon.tech', 'supabase.co', 'amazonaws.com', 'azure.com', 'cockroachlabs.cloud', 'planetscale.com')
    is_cloud = any(p in host for p in _cloud)
    params = {
        'host': host,
        'port': int(config.get('port', 5432)),
        'dbname': config['database_name'],
        'user': config['username'],
        'password': config['password'],
        'connect_timeout': 10,
        'sslmode': config.get('sslmode', 'require' if is_cloud else 'prefer'),
    }
    # Neon requires channel_binding=require; honour explicit override otherwise
    if 'channel_binding' in config:
        params['channel_binding'] = config['channel_binding']
    elif 'neon.tech' in host:
        params['channel_binding'] = 'require'
    try:
        import psycopg2 as _psycopg2
    except ImportError:
        raise RuntimeError('psycopg2 is not installed. Run: pip install psycopg2-binary')
    conn = _psycopg2.connect(**params)
    try:
        cur = conn.cursor()
        cur.execute('SELECT 1')
        cur.close()
    finally:
        conn.close()


def _introspect_postgres(sqlite_conn, run_id):
    """Pull real schema from a PostgreSQL connection. Returns (tables, defs)."""
    try:
        import psycopg2 as _pg
    except ImportError:
        raise RuntimeError('psycopg2 not installed')

    row = sqlite_conn.execute(
        'SELECT c.* FROM connections c '
        'JOIN governance_runs g ON g.connection_id = c.id WHERE g.id = ?',
        (run_id,)
    ).fetchone()
    if not row:
        raise ValueError('Connection not found')

    host     = row['account'] or ''
    port     = int(row['warehouse'] or 5432)
    dbname   = row['database_name'] or ''
    username = decrypt(row['username'])
    password = decrypt(row['password'])
    schema   = row['schema_name'] or 'public'

    _cloud = ('neon.tech', 'supabase.co', 'amazonaws.com', 'azure.com', 'cockroachlabs.cloud')
    sslmode = 'require' if any(p in host for p in _cloud) else 'prefer'
    extra   = {'channel_binding': 'require'} if 'neon.tech' in host else {}

    pg = _pg.connect(host=host, port=port, dbname=dbname, user=username, password=password,
                     connect_timeout=15, sslmode=sslmode, **extra)
    try:
        cur = pg.cursor()

        # Tables + fast row estimates via pg_stat_user_tables
        cur.execute("""
            SELECT t.table_name, t.table_schema,
                   COALESCE(s.n_live_tup, 0) AS row_est
            FROM information_schema.tables t
            LEFT JOIN pg_stat_user_tables s
              ON s.relname = t.table_name AND s.schemaname = t.table_schema
            WHERE t.table_schema = %s AND t.table_type = 'BASE TABLE'
            ORDER BY t.table_name LIMIT 100
        """, (schema,))
        raw_tables = cur.fetchall()

        # PKs
        cur.execute("""
            SELECT table_name FROM information_schema.table_constraints
            WHERE table_schema = %s AND constraint_type = 'PRIMARY KEY'
        """, (schema,))
        pk_tables = {r[0] for r in cur.fetchall()}

        # Freshness columns (first match per table)
        cur.execute("""
            SELECT table_name, MIN(column_name) FROM information_schema.columns
            WHERE table_schema = %s
              AND lower(column_name) IN ('updated_at','created_at','modified_at','_updated_at','last_modified')
            GROUP BY table_name
        """, (schema,))
        freshness_cols = dict(cur.fetchall())

        # All columns for semantic defs
        cur.execute("""
            SELECT table_name, column_name, data_type
            FROM information_schema.columns
            WHERE table_schema = %s
            ORDER BY table_name, ordinal_position
        """, (schema,))
        all_cols = cur.fetchall()

        from datetime import datetime, timezone
        tables = []
        for tname, tschema, row_est in raw_tables:
            n = row_est
            row_str = f'{n:,}' if n < 1_000_000 else f'{n/1_000_000:.1f}M'

            # Freshness
            freshness = 'N/A'
            if tname in freshness_cols:
                try:
                    cur.execute(f'SELECT MAX("{freshness_cols[tname]}") FROM "{tschema}"."{tname}"')
                    mx = cur.fetchone()[0]
                    if mx:
                        now = datetime.now(timezone.utc)
                        ts  = mx.replace(tzinfo=timezone.utc) if (hasattr(mx, 'tzinfo') and mx.tzinfo is None) else mx
                        h   = int((now - ts).total_seconds() / 3600)
                        freshness = ('<1h ago' if h < 1 else f'{h}h ago' if h < 24 else f'{h//24}d ago')
                except Exception:
                    pass

            has_pk      = tname in pk_tables
            pk_gate     = 'pass' if has_pk else 'warn'
            row_min_gate= 'pass' if n >= 500 else ('warn' if n >= 100 else 'fail')
            fresh_gate  = 'pass' if freshness != 'N/A' else 'warn'
            health      = max(0, 100 - (0 if has_pk else 15) - (10 if freshness == 'N/A' else 0)
                              - (30 if n < 100 else 10 if n < 500 else 0))
            ml_ready    = 1 if health >= 70 else 0

            tables.append((tname, tschema, health, freshness, row_str,
                           pk_gate, 'pass', fresh_gate, 'pass', row_min_gate, ml_ready))

        # Semantic defs from column patterns
        METRIC_WORDS  = ('revenue','amount','total','count','price','cost','sales',
                         'value','quantity','rate','score','spend','profit','margin')
        DIM_WORDS     = ('category','type','status','region','country','city','tier',
                         'segment','channel','source','platform','device','brand','label')
        NUMERIC_TYPES = ('integer','bigint','numeric','real','double precision',
                         'money','smallint','decimal','float')

        defs, seen = [], set()
        for tname, cname, dtype in all_cols:
            key = cname.lower()
            if key in seen:
                continue
            label   = cname.replace('_', ' ').title()
            explore = tname.replace('_', ' ').title()
            if any(w in key for w in METRIC_WORDS) and any(t in dtype for t in NUMERIC_TYPES):
                conf = round(0.60 + (abs(hash(cname)) % 25) / 100, 2)
                defs.append(('Metric', label,
                             f'{label} aggregated from {tname} (type: {dtype}).',
                             conf, explore))
                seen.add(key)
            elif any(w in key for w in DIM_WORDS):
                conf = round(0.58 + (abs(hash(cname)) % 25) / 100, 2)
                defs.append(('Dimension', label,
                             f'Categorical dimension — {label} from {tname}.',
                             conf, explore))
                seen.add(key)

        return tables[:50], defs[:30]
    finally:
        pg.close()


def _test_bigquery(config):
    try:
        from google.cloud import bigquery
        from google.oauth2 import service_account
    except ImportError:
        raise RuntimeError(
            'google-cloud-bigquery is not installed. '
            'Run: pip install google-cloud-bigquery'
        )
    creds_info = json.loads(config['credentials_json'])
    credentials = service_account.Credentials.from_service_account_info(creds_info)
    client = bigquery.Client(project=config['project_id'], credentials=credentials)
    try:
        list(client.query('SELECT 1').result())
    finally:
        client.close()


# ─────────────────────────────────────────────────────────
# Email helper
# ─────────────────────────────────────────────────────────
def _outbox_record(to, subject, html, status):
    """R1S2E6: every email lands in the SQLite outbox (inspectable offline)."""
    try:
        conn = _new_conn()
        try:
            conn.execute('INSERT INTO email_outbox (recipient, subject, body_html, status) '
                         'VALUES (?,?,?,?)', (to, subject, html, status))
            conn.commit()
        finally:
            conn.close()
    except Exception:
        log.exception('outbox write failed')


def send_email(to: str, subject: str, html: str):
    if not RESEND_API_KEY:
        log.info('EMAIL (dev) to=%s subject=%s', to, subject)
        _outbox_record(to, subject, html, 'queued')
        return
    try:
        resend.Emails.send({
            'from': FROM_EMAIL,
            'to': [to],
            'subject': subject,
            'html': html,
        })
        _outbox_record(to, subject, html, 'sent')
    except Exception:
        log.exception('Failed to send email to %s', to)
        _outbox_record(to, subject, html, 'failed')


# ─────────────────────────────────────────────────────────
# Database helpers
# ─────────────────────────────────────────────────────────
def _new_conn():
    conn = sqlite3.connect(_DB_PATH, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    conn.execute('PRAGMA journal_mode=WAL')
    conn.execute('PRAGMA foreign_keys=ON')
    conn.execute('PRAGMA busy_timeout=5000')  # R9S1E2: concurrent node workers
    return conn


def get_db():
    if 'db' not in g:
        g.db = _new_conn()
    return g.db


def thread_db():
    return _new_conn()


def put_db(conn):
    try:
        conn.rollback()
    except Exception:
        pass


@app.teardown_appcontext
def close_db(e=None):
    db = g.pop('db', None)
    if db:
        try:
            db.close()
        except Exception:
            pass


def one(sql, params=()):
    cur = get_db().execute(sql, params)
    row = cur.fetchone()
    return dict(row) if row else None


def many(sql, params=()):
    cur = get_db().execute(sql, params)
    return [dict(r) for r in cur.fetchall()]


def execute(sql, params=()):
    db = get_db()
    cur = db.execute(sql, params)
    lid = cur.lastrowid if sql.strip().upper().startswith('INSERT') else None
    db.commit()
    return lid


# ─────────────────────────────────────────────────────────
# Schema
# ─────────────────────────────────────────────────────────
SCHEMA = """
CREATE TABLE IF NOT EXISTS connections (
    id            INTEGER PRIMARY KEY,
    name          TEXT NOT NULL,
    type          TEXT NOT NULL DEFAULT 'snowflake',
    account       TEXT, username TEXT, password TEXT,
    warehouse TEXT, database_name TEXT, schema_name TEXT,
    owner_email   TEXT,
    status        TEXT DEFAULT 'active',
    poll_interval_minutes INTEGER,
    next_poll_at  TEXT,
    created_at    TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS poll_runs (
    id            INTEGER PRIMARY KEY,
    connection_id INTEGER REFERENCES connections(id),
    status        TEXT DEFAULT 'done',
    rows_ingested INTEGER DEFAULT 0,
    mode          TEXT,
    ran_at        TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS governance_runs (
    id                   INTEGER PRIMARY KEY,
    connection_id        INTEGER REFERENCES connections(id),
    status               TEXT DEFAULT 'pending',
    current_step         INTEGER DEFAULT 0,
    tables_count         INTEGER,
    definitions_count    INTEGER,
    low_confidence_count INTEGER,
    started_at           TEXT DEFAULT (datetime('now')),
    completed_at         TEXT
);
CREATE TABLE IF NOT EXISTS cataloged_tables (
    id             INTEGER PRIMARY KEY,
    run_id         INTEGER REFERENCES governance_runs(id),
    name           TEXT NOT NULL, schema_name TEXT,
    health_score   INTEGER, freshness TEXT, row_count TEXT,
    pk_gate        TEXT DEFAULT 'pass',
    null_gate      TEXT DEFAULT 'pass',
    freshness_gate TEXT DEFAULT 'pass',
    pii_gate       TEXT DEFAULT 'pass',
    row_min_gate   TEXT DEFAULT 'pass',
    ml_ready       INTEGER DEFAULT 1
);
CREATE TABLE IF NOT EXISTS semantic_definitions (
    id         INTEGER PRIMARY KEY,
    run_id     INTEGER REFERENCES governance_runs(id),
    type       TEXT NOT NULL, name TEXT NOT NULL,
    definition TEXT, confidence REAL, explore TEXT,
    status     TEXT DEFAULT 'pending'
);
CREATE TABLE IF NOT EXISTS dq_rule_settings (
    id            INTEGER PRIMARY KEY,
    connection_id INTEGER NOT NULL,
    rule_id       TEXT NOT NULL,
    enabled       INTEGER NOT NULL DEFAULT 1,
    block_on_failure INTEGER,
    updated_at    TEXT DEFAULT (datetime('now')),
    UNIQUE(connection_id, rule_id)
);
CREATE TABLE IF NOT EXISTS sessions (
    id             INTEGER PRIMARY KEY,
    parent_session_id INTEGER,
    connection_id  INTEGER, run_id INTEGER,
    is_sandbox     INTEGER NOT NULL DEFAULT 0,
    metric         TEXT NOT NULL,
    grain          TEXT NOT NULL DEFAULT 'Location · Day',
    horizon        INTEGER DEFAULT 14,
    training_start TEXT DEFAULT '2023-01-01',
    training_end   TEXT DEFAULT '2023-12-31',
    status         TEXT DEFAULT 'pending',
    created_at     TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS pipeline_runs (
    id             INTEGER PRIMARY KEY,
    session_id     INTEGER, status TEXT DEFAULT 'running',
    current_step   INTEGER DEFAULT 0, mape REAL,
    features_count INTEGER DEFAULT 0, rows_count INTEGER DEFAULT 0,
    log_entries    TEXT DEFAULT '[]',
    started_at     TEXT DEFAULT (datetime('now')),
    completed_at   TEXT
);
CREATE TABLE IF NOT EXISTS artifacts (
    id              INTEGER PRIMARY KEY,
    is_sandbox     INTEGER NOT NULL DEFAULT 0,
    confidence     REAL,
    confidence_json TEXT,
    layout_json    TEXT,
    title           TEXT NOT NULL, type TEXT DEFAULT 'Predictive',
    mape            REAL, owner TEXT DEFAULT 'analyst@acme.com',
    dq_status       TEXT DEFAULT 'pass',
    pipeline_run_id INTEGER,
    favorite        INTEGER DEFAULT 0,
    tags_json       TEXT DEFAULT '[]',
    created_at      TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS artifact_shares (
    id          INTEGER PRIMARY KEY,
    artifact_id INTEGER, email TEXT NOT NULL,
    role        TEXT DEFAULT 'Viewer',
    shared_at   TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS chart_data (
    id              INTEGER PRIMARY KEY,
    pipeline_run_id INTEGER, day_index INTEGER,
    date TEXT, actual REAL, predicted REAL,
    ci_low REAL, ci_high REAL, is_forecast INTEGER DEFAULT 0
);
CREATE TABLE IF NOT EXISTS artifact_schedules (
    id            INTEGER PRIMARY KEY,
    artifact_id   INTEGER REFERENCES artifacts(id) UNIQUE,
    org_id        TEXT,
    timezone      TEXT DEFAULT 'UTC',
    cron_expr     TEXT NOT NULL,
    next_run_at   TEXT NOT NULL,
    last_run_at   TEXT,
    enabled       INTEGER DEFAULT 1,
    created_at    TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS subscriptions (
    id                      INTEGER PRIMARY KEY,
    org_id                  TEXT NOT NULL,
    stripe_customer_id      TEXT,
    stripe_subscription_id  TEXT,
    plan                    TEXT NOT NULL DEFAULT 'free',
    status                  TEXT NOT NULL DEFAULT 'active',
    current_period_end      TEXT
);
CREATE TABLE IF NOT EXISTS health_history (
    id            INTEGER PRIMARY KEY,
    connection_id INTEGER REFERENCES connections(id),
    run_id        INTEGER REFERENCES governance_runs(id),
    table_name    TEXT NOT NULL,
    health_score  INTEGER,
    recorded_at   TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS artifact_insights (
    id             INTEGER PRIMARY KEY,
    artifact_id    INTEGER REFERENCES artifacts(id),
    kind           TEXT NOT NULL,
    summary        TEXT NOT NULL,
    drill_question TEXT,
    detail_json    TEXT DEFAULT '{}',
    dismissed      INTEGER DEFAULT 0,
    created_at     TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS artifact_activity (
    id          INTEGER PRIMARY KEY,
    artifact_id INTEGER REFERENCES artifacts(id),
    kind        TEXT NOT NULL,
    actor       TEXT,
    detail      TEXT,
    created_at  TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS share_links (
    id            INTEGER PRIMARY KEY,
    artifact_id   INTEGER REFERENCES artifacts(id),
    token_hash    TEXT NOT NULL,
    password_hash TEXT,
    expires_at    TEXT,
    view_count    INTEGER DEFAULT 0,
    snapshot_html TEXT,
    snapshot_at   TEXT,
    created_by    TEXT,
    created_at    TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS workspace_branding (
    workspace_id  TEXT PRIMARY KEY,
    primary_color TEXT, logo_text TEXT, font_family TEXT,
    updated_at    TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS artifact_annotations (
    id          INTEGER PRIMARY KEY,
    artifact_id INTEGER REFERENCES artifacts(id),
    grain_value TEXT, timestamp TEXT, text TEXT NOT NULL,
    author      TEXT,
    created_at  TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS metric_subscriptions (
    id          INTEGER PRIMARY KEY,
    artifact_id INTEGER REFERENCES artifacts(id),
    metric      TEXT NOT NULL,
    threshold   REAL NOT NULL,
    direction   TEXT DEFAULT 'above',
    subscriber  TEXT,
    created_at  TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS gold_predictions (
    id INTEGER PRIMARY KEY, pipeline_run_id INTEGER, session_id INTEGER,
    day_index INTEGER, date TEXT, actual REAL, predicted REAL,
    ci_low REAL, ci_high REAL,
    created_at TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS gold_forecast (
    id INTEGER PRIMARY KEY, pipeline_run_id INTEGER, session_id INTEGER,
    day_index INTEGER, date TEXT, predicted REAL, ci_low REAL, ci_high REAL,
    created_at TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS gold_model_insights (
    id            INTEGER PRIMARY KEY,
    session_id    INTEGER,
    model_card_id INTEGER,
    feature       TEXT NOT NULL,
    importance    REAL,
    shap_mean     REAL,
    rank          INTEGER,
    created_at    TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS custom_features (
    id           INTEGER PRIMARY KEY,
    session_id   INTEGER REFERENCES sessions(id),
    name         TEXT NOT NULL,
    expr         TEXT NOT NULL,
    status       TEXT DEFAULT 'pending_review',
    leakage_json TEXT DEFAULT '{}',
    applied      INTEGER DEFAULT 0,
    created_at   TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS leakage_confirmations (
    id            INTEGER PRIMARY KEY,
    session_id    INTEGER REFERENCES sessions(id),
    feature       TEXT NOT NULL,
    justification TEXT,
    confirmed_by  TEXT,
    created_at    TEXT DEFAULT (datetime('now')),
    UNIQUE (session_id, feature)
);
CREATE TABLE IF NOT EXISTS holiday_calendars (
    id         INTEGER PRIMARY KEY,
    name       TEXT NOT NULL,
    dates_json TEXT NOT NULL DEFAULT '[]',
    created_at TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS pdts (
    id                INTEGER PRIMARY KEY,
    workspace_id      TEXT NOT NULL DEFAULT 'default',
    name              TEXT NOT NULL,
    sql               TEXT NOT NULL,
    row_count         INTEGER,
    last_refreshed_at TEXT,
    created_at        TEXT DEFAULT (datetime('now')),
    UNIQUE (workspace_id, name)
);
CREATE TABLE IF NOT EXISTS dq_custom_tests (
    id            INTEGER PRIMARY KEY,
    connection_id INTEGER NOT NULL,
    table_name    TEXT NOT NULL,
    expression    TEXT NOT NULL,
    compiled_sql  TEXT NOT NULL,
    last_status   TEXT,
    last_violations INTEGER,
    last_run_at   TEXT,
    created_at    TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS data_contracts (
    id                    INTEGER PRIMARY KEY,
    connection_id         INTEGER NOT NULL,
    table_name            TEXT NOT NULL,
    required_columns_json TEXT DEFAULT '[]',
    min_rows              INTEGER,
    max_age_hours         REAL,
    updated_at            TEXT DEFAULT (datetime('now')),
    UNIQUE (connection_id, table_name)
);
CREATE TABLE IF NOT EXISTS freshness_slas (
    id            INTEGER PRIMARY KEY,
    connection_id INTEGER NOT NULL,
    table_name    TEXT NOT NULL,
    max_age_hours REAL NOT NULL,
    updated_at    TEXT DEFAULT (datetime('now')),
    UNIQUE (connection_id, table_name)
);
CREATE TABLE IF NOT EXISTS governance_thresholds (
    id            INTEGER PRIMARY KEY,
    connection_id INTEGER UNIQUE,
    min_health    INTEGER NOT NULL,
    updated_at    TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS alerts (
    id            INTEGER PRIMARY KEY,
    type          TEXT NOT NULL,
    connection_id INTEGER,
    subject       TEXT,
    detail_json   TEXT DEFAULT '{}',
    created_at    TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS webhook_events (
    id            INTEGER PRIMARY KEY,
    connection_id INTEGER REFERENCES connections(id),
    payload_json  TEXT NOT NULL,
    received_at   TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS jobs (
    id           INTEGER PRIMARY KEY,
    kind         TEXT NOT NULL,
    payload_json TEXT DEFAULT '{}',
    priority     INTEGER NOT NULL DEFAULT 0,
    status       TEXT DEFAULT 'queued',
    retries      INTEGER DEFAULT 0,
    max_retries  INTEGER DEFAULT 1,
    error        TEXT,
    created_at   TEXT DEFAULT (datetime('now')),
    started_at   TEXT, completed_at TEXT
);
CREATE TABLE IF NOT EXISTS service_logs (
    id          INTEGER PRIMARY KEY,
    method      TEXT, path TEXT, status INTEGER,
    duration_ms REAL,
    user_email  TEXT,
    created_at  TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS email_outbox (
    id         INTEGER PRIMARY KEY,
    recipient  TEXT NOT NULL,
    subject    TEXT,
    body_html  TEXT,
    status     TEXT DEFAULT 'queued',
    created_at TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS secrets (
    id              INTEGER PRIMARY KEY,
    name            TEXT NOT NULL UNIQUE,
    value_encrypted TEXT NOT NULL,
    updated_at      TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS resource_acls (
    id            INTEGER PRIMARY KEY,
    resource_type TEXT NOT NULL,
    resource_id   TEXT NOT NULL,
    principal     TEXT NOT NULL,
    role          TEXT NOT NULL DEFAULT 'viewer',
    created_at    TEXT DEFAULT (datetime('now')),
    UNIQUE (resource_type, resource_id, principal)
);
CREATE TABLE IF NOT EXISTS users (
    id            INTEGER PRIMARY KEY,
    email         TEXT NOT NULL UNIQUE,
    password_hash TEXT NOT NULL,
    role          TEXT NOT NULL DEFAULT 'analyst',
    created_at    TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS api_tokens (
    id         INTEGER PRIMARY KEY,
    user_id    INTEGER REFERENCES users(id),
    token_hash TEXT NOT NULL,
    expires_at TEXT NOT NULL,
    created_at TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS dq_gate_results (
    id               INTEGER PRIMARY KEY,
    connection_id    INTEGER,
    manifest_version TEXT,
    outcome          TEXT NOT NULL,
    result_hash      TEXT NOT NULL,
    trace_id         TEXT NOT NULL,
    rules_json       TEXT NOT NULL DEFAULT '[]',
    evaluated_at     TEXT,
    created_at       TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS artifact_files (
    id             INTEGER PRIMARY KEY,
    artifact_id    INTEGER REFERENCES artifacts(id),
    version        INTEGER NOT NULL DEFAULT 1,
    html           TEXT NOT NULL,
    size_bytes     INTEGER,
    sha256         TEXT,
    storage_uri    TEXT,
    validator_json TEXT DEFAULT '{}',
    created_at     TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS model_registry (
    id            INTEGER PRIMARY KEY,
    session_id    INTEGER,
    model_card_id INTEGER REFERENCES model_cards(id),
    model_id      TEXT NOT NULL,
    version       INTEGER NOT NULL DEFAULT 1,
    artifact_uri  TEXT,
    status        TEXT DEFAULT 'active',
    created_at    TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS training_jobs (
    id            INTEGER PRIMARY KEY,
    session_id    INTEGER REFERENCES sessions(id),
    gold_table_id INTEGER REFERENCES gold_tables(id),
    status        TEXT DEFAULT 'queued',
    priority      INTEGER DEFAULT 0,
    error         TEXT,
    model_card_id INTEGER,
    created_at    TEXT DEFAULT (datetime('now')),
    started_at    TEXT, completed_at TEXT
);
CREATE TABLE IF NOT EXISTS model_cards (
    id                       INTEGER PRIMARY KEY,
    session_id               INTEGER REFERENCES sessions(id),
    job_id                   INTEGER REFERENCES training_jobs(id),
    algorithm                TEXT,
    gold_table_name          TEXT,
    gold_output_hash         TEXT,
    feature_manifest_version TEXT,
    hyperparams_json         TEXT DEFAULT '{}',
    metrics_json             TEXT DEFAULT '{}',
    gates_json               TEXT DEFAULT '{}',
    target_type              TEXT DEFAULT 'regression',
    lineage_json             TEXT DEFAULT '{}',
    status                   TEXT DEFAULT 'candidate',
    trained_at               TEXT,
    created_at               TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS model_trials (
    id          INTEGER PRIMARY KEY,
    job_id      INTEGER REFERENCES training_jobs(id),
    session_id  INTEGER,
    params_json TEXT NOT NULL,
    mape        REAL,
    folds_json  TEXT DEFAULT '[]',
    created_at  TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS gold_tables (
    id             INTEGER PRIMARY KEY,
    session_id     INTEGER REFERENCES sessions(id),
    table_name     TEXT NOT NULL,
    physical_table TEXT,
    version        INTEGER NOT NULL DEFAULT 1,
    ddl            TEXT, insert_sql TEXT,
    output_hash    TEXT,
    row_count      INTEGER,
    status         TEXT DEFAULT 'generated',
    manifest_version TEXT,
    split_config_json TEXT, dq_json TEXT,
    created_at     TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS feature_manifests (
    id                INTEGER PRIMARY KEY,
    workspace_id      TEXT NOT NULL DEFAULT 'default',
    session_id        INTEGER REFERENCES sessions(id),
    gold_table_name   TEXT NOT NULL,
    gold_output_hash  TEXT,
    manifest_version  TEXT NOT NULL,
    feature_list_json TEXT NOT NULL DEFAULT '[]',
    enrichment_status TEXT DEFAULT 'scaffold',
    generated_at      TEXT,
    created_at        TEXT DEFAULT (datetime('now'))
);
CREATE TRIGGER IF NOT EXISTS feature_manifests_immutable
BEFORE UPDATE ON feature_manifests
BEGIN
    SELECT RAISE(ABORT, 'feature_manifests rows are immutable — create a new version');
END;
CREATE TABLE IF NOT EXISTS pipeline_steps (
    node_key            TEXT,
    id                 INTEGER PRIMARY KEY,
    run_id             INTEGER REFERENCES pipeline_runs(id),
    step               INTEGER NOT NULL,
    label              TEXT, description TEXT,
    input_schema_json  TEXT DEFAULT '[]',
    output_schema_json TEXT DEFAULT '[]',
    flagged            INTEGER DEFAULT 0,
    flag_reason        TEXT,
    created_at         TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS session_events (
    id           INTEGER PRIMARY KEY,
    session_id   INTEGER REFERENCES sessions(id),
    type         TEXT NOT NULL,
    payload_json TEXT DEFAULT '{}',
    created_at   TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS session_templates (
    id         INTEGER PRIMARY KEY,
    name       TEXT NOT NULL,
    spec_json  TEXT NOT NULL,
    created_by TEXT,
    created_at TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS session_specs (
    id              INTEGER PRIMARY KEY,
    session_id      INTEGER REFERENCES sessions(id),
    spec_version    INTEGER NOT NULL,
    idempotency_key TEXT,
    payload_hash    TEXT,
    spec_json       TEXT NOT NULL,
    created_at      TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS semantic_schemas (
    id           INTEGER PRIMARY KEY,
    workspace_id TEXT NOT NULL DEFAULT 'default',
    version      TEXT NOT NULL,
    schema_json  TEXT NOT NULL,
    change_note  TEXT,
    created_at   TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS governance_manifests (
    id            INTEGER PRIMARY KEY,
    connection_id INTEGER REFERENCES connections(id),
    run_id        INTEGER REFERENCES governance_runs(id),
    version       TEXT NOT NULL,
    manifest_json TEXT NOT NULL,
    created_at    TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS ingestion_profiles (
    id            INTEGER PRIMARY KEY,
    connection_id INTEGER REFERENCES connections(id),
    table_name    TEXT,
    status        TEXT DEFAULT 'pending',
    row_count     INTEGER,
    sampled_rows  INTEGER,
    columns_json  TEXT DEFAULT '[]',
    created_at    TEXT DEFAULT (datetime('now')),
    completed_at  TEXT
);
CREATE TABLE IF NOT EXISTS benchmark_refs (
    id         INTEGER PRIMARY KEY AUTOINCREMENT,
    metric     TEXT NOT NULL,
    kind       TEXT NOT NULL,
    value      REAL NOT NULL,
    created_at TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS plugin_validators (
    id         INTEGER PRIMARY KEY AUTOINCREMENT,
    name       TEXT UNIQUE NOT NULL,
    table_name TEXT NOT NULL,
    min_rows   INTEGER NOT NULL DEFAULT 1,
    created_at TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS outbound_actions (
    id         INTEGER PRIMARY KEY AUTOINCREMENT,
    kind       TEXT NOT NULL,
    target     TEXT NOT NULL,
    message    TEXT NOT NULL,
    status     TEXT NOT NULL DEFAULT 'queued',
    created_at TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS dashboard_templates (
    id            INTEGER PRIMARY KEY AUTOINCREMENT,
    name          TEXT NOT NULL,
    plan_template_json TEXT NOT NULL,
    created_by    TEXT,
    created_at    TEXT DEFAULT (datetime('now'))
);

CREATE TABLE IF NOT EXISTS rls_policies (
    id         INTEGER PRIMARY KEY AUTOINCREMENT,
    table_name TEXT NOT NULL,
    expression TEXT NOT NULL,
    status     TEXT NOT NULL DEFAULT 'on',
    created_at TEXT DEFAULT (datetime('now'))
);

CREATE TABLE IF NOT EXISTS notifications (
    id         INTEGER PRIMARY KEY AUTOINCREMENT,
    user_id    TEXT NOT NULL DEFAULT 'default',
    kind       TEXT NOT NULL,
    message    TEXT NOT NULL,
    link       TEXT,
    read       INTEGER NOT NULL DEFAULT 0,
    created_at TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS comments (
    id         INTEGER PRIMARY KEY AUTOINCREMENT,
    artifact_id INTEGER REFERENCES artifacts(id),
    parent_id  INTEGER,
    section_id TEXT,
    author     TEXT NOT NULL DEFAULT 'admin@acme.com',
    body       TEXT NOT NULL,
    resolved   INTEGER NOT NULL DEFAULT 0,
    created_at TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS team_invites (
    id         INTEGER PRIMARY KEY AUTOINCREMENT,
    email      TEXT NOT NULL,
    role       TEXT NOT NULL DEFAULT 'analyst',
    token      TEXT UNIQUE NOT NULL,
    status     TEXT NOT NULL DEFAULT 'pending',
    invited_by TEXT,
    created_at TEXT DEFAULT (datetime('now')),
    accepted_at TEXT
);

CREATE TABLE IF NOT EXISTS query_contracts (
    id                    INTEGER PRIMARY KEY AUTOINCREMENT,
    run_id                INTEGER NOT NULL,
    component_id          TEXT NOT NULL,
    sql                   TEXT NOT NULL,
    warehouse_dialect     TEXT NOT NULL DEFAULT 'sqlite',
    expected_columns_json TEXT NOT NULL DEFAULT '[]',
    row_limit             INTEGER,
    status                TEXT NOT NULL DEFAULT 'pending',
    execution_time_ms     INTEGER,
    created_at            TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS component_data_contracts (
    id                        INTEGER PRIMARY KEY AUTOINCREMENT,
    run_id                    INTEGER NOT NULL,
    component_id              TEXT NOT NULL,
    query_contract_id         INTEGER,
    actual_columns_json       TEXT NOT NULL DEFAULT '[]',
    row_count                 INTEGER,
    numeric_ranges_json       TEXT NOT NULL DEFAULT '{}',
    high_cardinality_json     TEXT NOT NULL DEFAULT '[]',
    empty_result              INTEGER NOT NULL DEFAULT 0,
    contract_version          TEXT DEFAULT '1.0',
    created_at                TEXT DEFAULT (datetime('now'))
);

CREATE TABLE IF NOT EXISTS platform_signals (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    signal_kind TEXT NOT NULL,
    subject     TEXT NOT NULL,
    consumer    TEXT NOT NULL,
    detail_json TEXT NOT NULL DEFAULT '{}',
    status      TEXT NOT NULL DEFAULT 'delivered',
    fingerprint TEXT UNIQUE,
    created_at  TEXT DEFAULT (datetime('now'))
);

CREATE TABLE IF NOT EXISTS recommendation_feedback (
    id         INTEGER PRIMARY KEY AUTOINCREMENT,
    rec_type   TEXT NOT NULL,
    rec_id     TEXT,
    category   TEXT,
    decision   TEXT NOT NULL,
    user_id    TEXT NOT NULL DEFAULT 'default',
    signal     REAL,
    created_at TEXT DEFAULT (datetime('now'))
);

CREATE TABLE IF NOT EXISTS opportunities (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    artifact_id INTEGER REFERENCES artifacts(id),
    kind        TEXT NOT NULL,
    headline    TEXT NOT NULL,
    question    TEXT NOT NULL,
    detail_json TEXT NOT NULL DEFAULT '{}',
    status      TEXT NOT NULL DEFAULT 'open',
    fingerprint TEXT UNIQUE,
    decided_at  TEXT,
    created_at  TEXT DEFAULT (datetime('now'))
);

CREATE TABLE IF NOT EXISTS repair_attempts (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    scope       TEXT NOT NULL,
    artifact_id INTEGER,
    run_id      INTEGER,
    cycle       INTEGER NOT NULL,
    detail_json TEXT NOT NULL DEFAULT '{}',
    resolved    INTEGER NOT NULL DEFAULT 0,
    created_at  TEXT DEFAULT (datetime('now'))
);

CREATE TABLE IF NOT EXISTS semantic_proposals (
    id            INTEGER PRIMARY KEY AUTOINCREMENT,
    kind          TEXT NOT NULL,
    subject       TEXT NOT NULL,
    suggestion    TEXT NOT NULL,
    evidence_json TEXT NOT NULL DEFAULT '{}',
    status        TEXT NOT NULL DEFAULT 'proposed',
    fingerprint   TEXT UNIQUE,
    decided_by    TEXT,
    decided_at    TEXT,
    created_at    TEXT DEFAULT (datetime('now'))
);

CREATE TABLE IF NOT EXISTS intent_history (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    user_id     TEXT NOT NULL DEFAULT 'default',
    step_kind   TEXT NOT NULL,
    ref_id      TEXT,
    session_id  INTEGER,
    detail_json TEXT DEFAULT '{}',
    created_at  TEXT DEFAULT (datetime('now'))
);

CREATE TABLE IF NOT EXISTS kg_edges (
    id         INTEGER PRIMARY KEY AUTOINCREMENT,
    edge_type  TEXT NOT NULL,
    src_node   TEXT NOT NULL,
    dst_node   TEXT NOT NULL,
    weight     REAL NOT NULL DEFAULT 1.0,
    created_at TEXT DEFAULT (datetime('now')),
    updated_at TEXT DEFAULT (datetime('now')),
    UNIQUE(edge_type, src_node, dst_node)
);

CREATE TABLE IF NOT EXISTS agent_memory (
    id           INTEGER PRIMARY KEY AUTOINCREMENT,
    workspace_id TEXT NOT NULL DEFAULT 'default',
    user_id      TEXT NOT NULL DEFAULT 'default',
    agent        TEXT NOT NULL,
    category     TEXT NOT NULL,
    mem_key      TEXT NOT NULL,
    value        TEXT NOT NULL,
    weight       REAL NOT NULL DEFAULT 1.0,
    last_used    TEXT NOT NULL DEFAULT (datetime('now')),
    created_at   TEXT DEFAULT (datetime('now')),
    UNIQUE(workspace_id, user_id, agent, category, mem_key)
);

CREATE TABLE IF NOT EXISTS optimization_proposals (
    id             INTEGER PRIMARY KEY AUTOINCREMENT,
    kind           TEXT NOT NULL,
    target         TEXT NOT NULL,
    recommendation TEXT NOT NULL,
    evidence_json  TEXT NOT NULL DEFAULT '{}',
    status         TEXT NOT NULL DEFAULT 'proposed',
    fingerprint    TEXT UNIQUE,
    decided_by     TEXT,
    decided_at     TEXT,
    created_at     TEXT DEFAULT (datetime('now'))
);

CREATE TABLE IF NOT EXISTS agent_consultations (
    id            INTEGER PRIMARY KEY AUTOINCREMENT,
    from_agent    TEXT NOT NULL,
    to_agent      TEXT NOT NULL,
    question_json TEXT NOT NULL DEFAULT '{}',
    answer_json   TEXT NOT NULL DEFAULT '{}',
    run_id        INTEGER,
    created_at    TEXT DEFAULT (datetime('now'))
);

CREATE TABLE IF NOT EXISTS meta_decisions (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    kind        TEXT NOT NULL,
    rule        TEXT NOT NULL,
    winner      TEXT,
    loser       TEXT,
    run_id      INTEGER,
    detail_json TEXT DEFAULT '{}',
    created_at  TEXT DEFAULT (datetime('now'))
);

CREATE TABLE IF NOT EXISTS platform_events (
    id           INTEGER PRIMARY KEY AUTOINCREMENT,
    event_type   TEXT NOT NULL,
    payload_json TEXT NOT NULL DEFAULT '{}',
    status       TEXT NOT NULL DEFAULT 'queued',
    processed_at TEXT,
    created_at   TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS opportunity_investigations (
    id           INTEGER PRIMARY KEY AUTOINCREMENT,
    event_id     INTEGER,
    payload_json TEXT NOT NULL DEFAULT '{}',
    status       TEXT NOT NULL DEFAULT 'open',
    created_at   TEXT DEFAULT (datetime('now'))
);

CREATE TABLE IF NOT EXISTS platform_settings (
    key        TEXT PRIMARY KEY,
    value      TEXT,
    updated_at TEXT DEFAULT (datetime('now'))
);

CREATE TABLE IF NOT EXISTS task_dispatches (
    id             INTEGER PRIMARY KEY AUTOINCREMENT,
    task_kind      TEXT NOT NULL,
    tier           TEXT NOT NULL,
    est_cost       REAL NOT NULL DEFAULT 0,
    est_latency_ms INTEGER NOT NULL DEFAULT 0,
    signature      TEXT,
    workspace_id   TEXT NOT NULL DEFAULT 'default',
    created_at     TEXT DEFAULT (datetime('now'))
);
CREATE TABLE IF NOT EXISTS task_templates (
    id            INTEGER PRIMARY KEY AUTOINCREMENT,
    task_kind     TEXT NOT NULL,
    pattern_key   TEXT NOT NULL,
    template_json TEXT NOT NULL,
    created_at    TEXT DEFAULT (datetime('now')),
    UNIQUE(task_kind, pattern_key)
);

CREATE TABLE IF NOT EXISTS dag_nodes (
    id              INTEGER PRIMARY KEY AUTOINCREMENT,
    run_id          INTEGER NOT NULL REFERENCES pipeline_runs(id),
    node_key        TEXT NOT NULL,
    node_type       TEXT NOT NULL DEFAULT 'stage',
    content_hash    TEXT NOT NULL,
    status          TEXT NOT NULL DEFAULT 'pending',
    cached          INTEGER NOT NULL DEFAULT 0,
    prior_run_id    INTEGER,
    uas_artifact_id TEXT,
    started_at      TEXT,
    completed_at    TEXT,
    created_at      TEXT DEFAULT (datetime('now')),
    UNIQUE(run_id, node_key)
);
CREATE INDEX IF NOT EXISTS idx_dag_nodes_hash ON dag_nodes(node_key, content_hash);
CREATE TABLE IF NOT EXISTS dag_edges (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    run_id      INTEGER NOT NULL REFERENCES pipeline_runs(id),
    from_key    TEXT NOT NULL,
    to_key      TEXT NOT NULL,
    gate_name   TEXT NOT NULL,
    gate_status TEXT NOT NULL,
    gate_detail TEXT DEFAULT '{}',
    created_at  TEXT DEFAULT (datetime('now'))
);

CREATE TABLE IF NOT EXISTS cache_entries (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    layer       TEXT NOT NULL,
    cache_key   TEXT NOT NULL,
    value_json  TEXT NOT NULL,
    gov_version TEXT,
    sem_version TEXT,
    expires_at  REAL NOT NULL,
    created_at  TEXT DEFAULT (datetime('now')),
    UNIQUE(layer, cache_key)
);
CREATE TABLE IF NOT EXISTS cache_stats (
    layer  TEXT PRIMARY KEY,
    hits   INTEGER NOT NULL DEFAULT 0,
    misses INTEGER NOT NULL DEFAULT 0
);

CREATE TABLE IF NOT EXISTS uas_artifacts (
    id                          INTEGER PRIMARY KEY AUTOINCREMENT,
    artifact_uid                TEXT UNIQUE NOT NULL,
    logical_key                 TEXT NOT NULL,
    artifact_type               TEXT NOT NULL,
    version                     INTEGER NOT NULL DEFAULT 1,
    content_hash                TEXT NOT NULL,
    upstream_artifact_ids       TEXT NOT NULL DEFAULT '[]',
    governance_manifest_version TEXT,
    semantic_layer_version      TEXT,
    created_by_agent            TEXT,
    workspace_id                TEXT NOT NULL DEFAULT 'default',
    run_id                      INTEGER,
    payload_json                TEXT NOT NULL,
    created_at                  TEXT DEFAULT (datetime('now')),
    UNIQUE(logical_key, version)
);
CREATE INDEX IF NOT EXISTS idx_uas_logical ON uas_artifacts(logical_key);
CREATE INDEX IF NOT EXISTS idx_uas_run ON uas_artifacts(run_id);
CREATE TRIGGER IF NOT EXISTS uas_immutable_update
BEFORE UPDATE ON uas_artifacts
BEGIN
    SELECT RAISE(ABORT, 'uas_artifacts is immutable (append-only)');
END;
CREATE TRIGGER IF NOT EXISTS uas_immutable_delete
BEFORE DELETE ON uas_artifacts
BEGIN
    SELECT RAISE(ABORT, 'uas_artifacts is immutable (append-only)');
END;

CREATE TABLE IF NOT EXISTS audit_logs (
    id            INTEGER PRIMARY KEY,
    org_id        TEXT,
    user_email    TEXT,
    action        TEXT NOT NULL,
    resource_type TEXT NOT NULL,
    resource_id   TEXT,
    metadata      TEXT DEFAULT '{}',
    created_at    TEXT DEFAULT (datetime('now'))
);
CREATE INDEX IF NOT EXISTS idx_audit_logs_org ON audit_logs (org_id);
CREATE TRIGGER IF NOT EXISTS audit_logs_no_update
BEFORE UPDATE ON audit_logs
BEGIN
    SELECT RAISE(ABORT, 'audit_logs is append-only');
END;
CREATE TRIGGER IF NOT EXISTS audit_logs_no_delete
BEFORE DELETE ON audit_logs
BEGIN
    SELECT RAISE(ABORT, 'audit_logs is append-only');
END;
CREATE INDEX IF NOT EXISTS idx_audit_logs_resource ON audit_logs (resource_type, resource_id);
"""

# ─────────────────────────────────────────────────────────
# Chart data (deterministic seeded RNG — matches JS version)
# ─────────────────────────────────────────────────────────
def seeded_rng(seed=42):
    s = seed

    def rand():
        nonlocal s
        s = (s * 9301 + 49297) % 233280
        return s / 233280

    return rand


def generate_chart_data(pipeline_run_id: int) -> list[dict]:
    rand    = seeded_rng(42)
    weekly  = [0.88, 0.94, 1.0, 1.06, 1.13, 0.80, 0.73]
    start   = date(2024, 1, 15)
    rows    = []
    for i in range(90):
        trend  = 1 + (i / 90) * 0.10
        base   = 620 * weekly[i % 7] * trend
        actual = round(base + (rand() - 0.5) * 95) if i < 76 else None
        pred   = round(base + (rand() - 0.5) * 55)
        ci     = 48 + i * 1.9
        d      = start + timedelta(days=i)
        rows.append({
            'pipeline_run_id': pipeline_run_id,
            'day_index': i,
            'date': d.strftime('%b ') + str(d.day),
            'actual': actual,
            'predicted': pred,
            'ci_low': round(pred - ci),
            'ci_high': round(pred + ci),
            'is_forecast': 1 if i >= 76 else 0,
        })
    return rows


def compute_kpis(rows: list[dict]) -> dict:
    hist = [r for r in rows if 46 <= r['day_index'] < 76 and r['actual'] is not None]
    fcast = [r for r in rows if r['is_forecast']]
    avg_actual   = round(sum(r['actual'] for r in hist) / len(hist)) if hist else 0
    mape         = round(sum(abs(r['actual'] - r['predicted']) / r['actual'] for r in hist) / len(hist) * 100, 1) if hist else 0
    forecast14   = round(sum(r['predicted'] for r in fcast) / len(fcast)) if fcast else 0
    return {'avgActual': avg_actual, 'mape': mape, 'forecast14Avg': forecast14}


def init_db():
    """Create schema on first run. No demo data is seeded."""
    conn = _new_conn()
    try:
        conn.executescript(SCHEMA)
        # lightweight migrations for pre-existing dev databases
        for ddl in ('ALTER TABLE task_dispatches ADD COLUMN tokens INTEGER NOT NULL DEFAULT 0',
                    'ALTER TABLE audit_logs ADD COLUMN severity TEXT DEFAULT \'info\'',
                    'ALTER TABLE dq_custom_tests ADD COLUMN enabled INTEGER NOT NULL DEFAULT 1',
                    'ALTER TABLE artifacts ADD COLUMN layout_json TEXT',
                    'ALTER TABLE artifacts ADD COLUMN confidence REAL',
                    'ALTER TABLE artifacts ADD COLUMN confidence_json TEXT',
                    'ALTER TABLE sessions ADD COLUMN is_sandbox INTEGER NOT NULL DEFAULT 0',
                    'ALTER TABLE artifacts ADD COLUMN is_sandbox INTEGER NOT NULL DEFAULT 0',
                    'ALTER TABLE jobs ADD COLUMN priority INTEGER NOT NULL DEFAULT 0',
                    'ALTER TABLE pipeline_steps ADD COLUMN node_key TEXT',
                    'ALTER TABLE artifact_files ADD COLUMN storage_uri TEXT',
                    'ALTER TABLE connections ADD COLUMN poll_interval_minutes INTEGER',
                    'ALTER TABLE connections ADD COLUMN next_poll_at TEXT',
                    'ALTER TABLE sessions ADD COLUMN parent_session_id INTEGER',
                    "ALTER TABLE model_cards ADD COLUMN target_type TEXT DEFAULT 'regression'",
                    "ALTER TABLE model_cards ADD COLUMN lineage_json TEXT DEFAULT '{}'",
                    'ALTER TABLE artifacts ADD COLUMN favorite INTEGER DEFAULT 0',
                    "ALTER TABLE artifact_schedules ADD COLUMN timezone TEXT DEFAULT 'UTC'",
                    "ALTER TABLE artifacts ADD COLUMN tags_json TEXT DEFAULT '[]'"):
            try:
                conn.execute(ddl)
            except sqlite3.OperationalError:
                pass
        print('Schema ready')
    finally:
        conn.close()


# ─────────────────────────────────────────────────────────
# SSE client registries (in-memory, per-process)
# ─────────────────────────────────────────────────────────
_gov_clients: dict[int, list[Queue]]  = {}
_gov_lock    = threading.Lock()
_pipe_clients: dict[int, list[Queue]] = {}
_pipe_lock   = threading.Lock()


def _broadcast(registry, lock, run_id, data):
    with lock:
        for q in list(registry.get(run_id, [])):
            try:
                q.put_nowait(data)
            except Exception:
                pass


def broadcast_gov(run_id, data):  _broadcast(_gov_clients,  _gov_lock,  run_id, data)
def broadcast_pipe(run_id, data): _broadcast(_pipe_clients, _pipe_lock, run_id, data)


# ─────────────────────────────────────────────────────────
# Audit log helpers
# ─────────────────────────────────────────────────────────
def log_action(action, resource_type, resource_id=None, metadata=None):
    user_email = getattr(g, 'user_email', None)
    org_id = getattr(g, 'org_id', None)
    db = get_db()
    db.execute(
        'INSERT INTO audit_logs (org_id, user_email, action, resource_type, resource_id, metadata) '
        'VALUES (?, ?, ?, ?, ?, ?)',
        (org_id, user_email, action, resource_type, str(resource_id) if resource_id is not None else None,
         json.dumps(metadata or {})),
    )
    db.commit()


def log_action_bg(conn, action, resource_type, resource_id=None, metadata=None):
    conn.execute(
        'INSERT INTO audit_logs (org_id, user_email, action, resource_type, resource_id, metadata) '
        'VALUES (?, ?, ?, ?, ?, ?)',
        (None, None, action, resource_type, str(resource_id) if resource_id is not None else None,
         json.dumps(metadata or {})),
    )
    conn.commit()


# ─────────────────────────────────────────────────────────
# Identity & RBAC (R1S1E1) — bearer tokens preferred; legacy
# X-User-Role header kept for dev compat (absent header = admin)
# ─────────────────────────────────────────────────────────
@app.before_request
def _start_timer():
    g._req_start = time.time()


@app.after_request
def _log_request(resp):
    try:
        if request.path.startswith('/api/') and not request.path.startswith('/api/platform/logs'):
            dur = (time.time() - getattr(g, '_req_start', time.time())) * 1000
            db = get_db()
            db.execute('INSERT INTO service_logs (method, path, status, duration_ms, user_email) '
                       'VALUES (?,?,?,?,?)',
                       (request.method, request.path, resp.status_code, round(dur, 2),
                        getattr(g, 'user_email', None)))
            db.commit()
    except Exception:
        pass
    return resp


@app.before_request
def _resolve_identity():
    import authn
    g.auth_user = None
    g.identity_source = 'header'
    auth = request.headers.get('Authorization', '')
    if auth.startswith('Embed '):
        import embed_tokens as et
        payload = et.verify(auth[6:].strip(),
                            et.workspace_secret('default'))
        g.embed_payload = payload
        g.identity_source = 'embed'
        return
    if auth.startswith('Bearer '):
        user = authn.resolve_token(get_db(), auth[7:].strip())
        if user:
            g.auth_user = user
            g.identity_source = 'token'
            g.user_email = user['email']


def current_role():
    if getattr(g, 'identity_source', None) == 'embed':
        return 'embed'   # R7S1E2: embed tokens are never write-authorized
    user = getattr(g, 'auth_user', None)
    if user:
        return (user.get('role') or 'viewer').lower()
    return (request.headers.get('X-User-Role') or 'admin').strip().lower()


ACL_ROLES = ('viewer', 'editor', 'owner')


def acl_entries(rtype, rid):
    return many('SELECT * FROM resource_acls WHERE resource_type=? AND resource_id=?',
                (rtype, str(rid)))


def acl_allows(rtype, rid, need='read'):
    """Restrict-only ACLs (R1S1E2): when a resource has ACL entries, callers
    must appear on them (admins bypass). ACL role caps the workspace role.
    Returns None when allowed, else a (response, status) tuple."""
    entries = acl_entries(rtype, rid)
    if not entries:
        return None                      # unrestricted → workspace defaults apply
    if current_role() == 'admin':
        return None
    user = getattr(g, 'auth_user', None)
    principal = user['email'] if user else None
    mine = next((e for e in entries if e['principal'] == principal), None)
    if not mine:
        return jsonify({'error': 'Forbidden — this resource is restricted',
                        'identity_source': getattr(g, 'identity_source', 'header')}), 403
    if need == 'write' and mine['role'] not in ('editor', 'owner'):
        return jsonify({'error': 'Forbidden — your access to this resource is read-only',
                        'acl_role': mine['role']}), 403
    return None


def acl_permits(rtype, rid, need='read'):
    return acl_allows(rtype, rid, need) is None


def require_role(*allowed):
    """403 unless the caller's role is in `allowed` (case-insensitive)."""
    def decorator(f):
        @functools.wraps(f)
        def wrapper(*args, **kwargs):
            role = current_role()
            if role not in allowed:
                return jsonify({
                    'error': 'Forbidden',
                    'required_role': allowed[0],
                    'current_role': role,
                    'identity_source': getattr(g, 'identity_source', 'header'),
                }), 403
            return f(*args, **kwargs)
        return wrapper
    return decorator


# ─────────────────────────────────────────────────────────
# Connection payload validation
# ─────────────────────────────────────────────────────────
CONNECTOR_REQUIRED_FIELDS = {
    'snowflake': ('account', 'username', 'password'),
    'postgres':  ('host', 'database_name', 'username', 'password'),
    'bigquery':  ('project_id', 'credentials_json'),
    'gsheet':    ('sheet_url',),
    'webhook':   (),
    'rest_api':  ('endpoint_url',),
    'mysql':      ('host', 'database_name', 'username', 'password'),
    'redshift':   ('host', 'database_name', 'username', 'password'),
    'databricks': ('host', 'http_path', 'access_token'),
    'duckdb':     ('database_path',),
}


def _valid_sheet_url(url):
    import re as _re
    return bool(_re.match(r'^https://docs\.google\.com/spreadsheets/d/[\w-]+', url or ''))


def validate_connection_payload(b):
    """Return dict of field → error for an invalid payload, {} when valid."""
    errors = {}
    conn_type = b.get('type', 'snowflake')
    if conn_type not in CONNECTOR_REQUIRED_FIELDS:
        return {'type': f'Unknown connector type: {conn_type}'}
    for f in CONNECTOR_REQUIRED_FIELDS[conn_type]:
        v = b.get(f)
        if v is None or (isinstance(v, str) and not v.strip()):
            errors[f] = 'This field is required'
    if conn_type == 'gsheet' and b.get('sheet_url') and not _valid_sheet_url(b['sheet_url']):
        errors['sheet_url'] = 'Must be a docs.google.com/spreadsheets URL'
    if conn_type == 'rest_api' and b.get('endpoint_url') and \
            not str(b['endpoint_url']).startswith(('http://', 'https://')):
        errors['endpoint_url'] = 'Must be an http(s) URL'
    return errors


# ─────────────────────────────────────────────────────────
# Subscription / plan helpers
# ─────────────────────────────────────────────────────────
_PLAN_RANK = {'free': 0, 'pro': 1, 'enterprise': 2}


def get_org_plan(org_id):
    if not org_id:
        return 'pro' if not STRIPE_SECRET_KEY else 'free'
    row = one(
        "SELECT plan FROM subscriptions "
        "WHERE org_id=? AND status='active' "
        "ORDER BY current_period_end DESC LIMIT 1",
        (org_id,),
    )
    return row['plan'] if row else 'free'


def plan_gate(required_plan):
    def decorator(f):
        @functools.wraps(f)
        def wrapper(*args, **kwargs):
            org_id = request.headers.get('X-Org-Id')
            current = get_org_plan(org_id)
            if _PLAN_RANK.get(current, 0) < _PLAN_RANK.get(required_plan, 0):
                return jsonify({
                    'error': 'Upgrade required',
                    'required_plan': required_plan,
                    'current_plan': current,
                }), 403
            return f(*args, **kwargs)
        return wrapper
    return decorator


# ─────────────────────────────────────────────────────────
# Schema drift detection (R3S1E3)
# ─────────────────────────────────────────────────────────
def _save_manifest_with_drift(conn, connection_id, manifest_dict):
    """Persist a manifest and raise a drift alert when the schema changed."""
    import manifest as manifest_mod
    prev_row = conn.execute(
        'SELECT manifest_json, version FROM governance_manifests WHERE connection_id=? '
        'ORDER BY id DESC LIMIT 1', (connection_id,)).fetchone()
    prev = json.loads(prev_row['manifest_json']) if prev_row else None
    saved = manifest_mod.save_manifest(conn, connection_id, manifest_dict)
    if prev and manifest_mod.schema_fingerprint(prev) != manifest_mod.schema_fingerprint(saved):
        prev_tables = {t['name'] for t in prev.get('tables', [])}
        new_tables = {t['name'] for t in saved.get('tables', [])}
        detail = {
            'from_version': prev_row['version'],
            'to_version': saved['manifest_version'],
            'added_tables': sorted(new_tables - prev_tables),
            'removed_tables': sorted(prev_tables - new_tables),
        }
        conn.execute('INSERT INTO alerts (type, connection_id, subject, detail_json) '
                     "VALUES ('drift', ?, ?, ?)",
                     (connection_id,
                      f"Schema drift detected on integration {connection_id}",
                      json.dumps(detail)))
        conn.commit()
        send_email(to='workspace-admins@analytiq.dev',
                   subject='Schema drift detected',
                   html=(f"<p>Manifest {detail['from_version']} → {detail['to_version']}: "
                         f"+{detail['added_tables']} -{detail['removed_tables']}</p>"))
    return saved


# ─────────────────────────────────────────────────────────
# Governance simulation
# ─────────────────────────────────────────────────────────
# SIM_DELAY_SCALE: multiplier applied to all simulation delays. The UI test
# server (tests/ui/boot_server.py) sets it to 0 so Playwright runs are
# deterministic-fast; backend tests keep monkeypatching module constants.
_SIM_SCALE = float(os.environ.get('SIM_DELAY_SCALE', '1'))
GOV_DELAYS = [d * _SIM_SCALE for d in (2.0, 3.0, 4.0, 2.5)]

# Simulated per-table column samples (drive PII scan + manifest columns)
SIM_TABLE_COLUMNS = {
    'fact_revenue': [
        {'name': 'revenue_id',  'semantic_type': 'id'},
        {'name': 'location_id', 'semantic_type': 'id'},
        {'name': 'day',         'semantic_type': 'date'},
        {'name': 'net_revenue', 'semantic_type': 'measure'},
    ],
    'dim_location': [
        {'name': 'location_id', 'semantic_type': 'id'},
        {'name': 'city',        'semantic_type': 'dimension'},
        {'name': 'tier',        'semantic_type': 'dimension'},
    ],
    'fact_sessions': [
        {'name': 'session_id',   'semantic_type': 'id'},
        {'name': 'customer_id',  'semantic_type': 'id', 'null_pct': 4.1},
        {'name': 'duration_sec', 'semantic_type': 'measure'},
    ],
    'dim_customer': [
        {'name': 'customer_id', 'semantic_type': 'id'},
        {'name': 'email',       'semantic_type': 'text',
         'samples': ['ava.chen@example.com', 'liam.p@corp.io', 'noah.k@mail.net',
                     'mia.r@example.com', 'ethan.w@corp.io']},
        {'name': 'phone',       'semantic_type': 'text',
         'samples': ['555-201-7788', '555-882-1034', '555-449-2210']},
        {'name': 'segment',     'semantic_type': 'dimension',
         'samples': ['SMB', 'Enterprise', 'SMB', 'Consumer']},
    ],
    'staging_events': [
        {'name': 'event_id',   'semantic_type': 'id'},
        {'name': 'event_type', 'semantic_type': 'dimension',
         'samples': ['click', 'view', 'purchase']},
        {'name': 'ip_address', 'semantic_type': 'text',
         'samples': ['10.4.22.19', '172.16.4.88', '192.168.7.3']},
    ],
    'raw_clickstream': [
        {'name': 'raw_id',  'semantic_type': 'id'},
        {'name': 'payload', 'semantic_type': 'text',
         'samples': ['{"e":"click"}', '{"e":"view"}']},
    ],
}


def simulate_governance(run_id: int):
    def _run():
        conn = thread_db()
        try:
            creds = conn.execute(
                'SELECT c.username, c.password FROM governance_runs g '
                'JOIN connections c ON c.id = g.connection_id WHERE g.id = ?',
                (run_id,),
            ).fetchone()
            if creds:
                _sf_username = decrypt(creds['username'])
                _sf_password = decrypt(creds['password'])
            for step in range(1, 5):
                time.sleep(GOV_DELAYS[step - 1])
                status = 'done' if step == 4 else 'running'
                conn.execute('UPDATE governance_runs SET current_step=?, status=? WHERE id=?', (step, status, run_id))
                conn.commit()
                if step < 4:
                    broadcast_gov(run_id, {'step': step, 'status': status})

            # Try real introspection for postgres connections
            conn_row = conn.execute(
                'SELECT c.type FROM connections c '
                'JOIN governance_runs g ON g.connection_id = c.id WHERE g.id = ?',
                (run_id,)
            ).fetchone()
            conn_type = conn_row['type'] if conn_row else None

            real_tables, real_defs = None, None
            if conn_type == 'postgres':
                try:
                    real_tables, real_defs = _introspect_postgres(conn, run_id)
                except Exception as exc:
                    print(f'[gov] introspect failed, falling back to sim: {exc}')

            if real_tables is not None and len(real_tables) > 0:
                sim_tables = [(t[0], t[1], t[2], t[3], t[4], t[5], t[6], t[7], t[8], t[9], t[10])
                              for t in real_tables]
                sim_defs   = real_defs or []
            else:
                sim_tables = [
                    ('fact_revenue',    'CORE',    98, '2h ago',  '4.2M',  'pass', 'pass', 'pass', 'pass', 'pass', 1),
                    ('dim_location',    'CORE',    94, '6h ago',  '12.8K', 'pass', 'pass', 'pass', 'pass', 'pass', 1),
                    ('fact_sessions',   'CORE',    87, '1h ago',  '2.1M',  'pass', 'warn', 'pass', 'pass', 'pass', 1),
                    ('dim_customer',    'CORE',    71, '3d ago',  '84.2K', 'pass', 'warn', 'warn', 'flag', 'pass', 0),
                    ('staging_events',  'STAGING', 90, '30m ago', '890K',  'warn', 'pass', 'pass', 'pass', 'pass', 1),
                    ('raw_clickstream', 'RAW',     44, '12d ago', '124',   'fail', 'warn', 'fail', 'pass', 'fail', 0),
                ]
                sim_defs = [
                    ('Metric',    'Net Revenue',          'Total revenue after refunds and discounts, aggregated daily per location.',   0.71, 'Revenue'),
                    ('Metric',    'Conversion Rate',      'Percentage of sessions that resulted in at least one purchase.',              0.64, 'Revenue'),
                    ('Dimension', 'Location Tier',        'Operational tier classification assigned to a physical location.',            0.68, 'Location Perf.'),
                    ('Metric',    'Avg Session Duration', 'Mean time in seconds a user spent in an active session.',                     0.59, 'Engagement'),
                    ('Dimension', 'Customer Segment',     'Behavioral segment label assigned to a customer by the ML pipeline.',         0.73, 'Customer'),
                ]
            if not conn.execute('SELECT 1 FROM cataloged_tables WHERE run_id=?', (run_id,)).fetchone():
                conn.executemany(
                    'INSERT INTO cataloged_tables (run_id,name,schema_name,health_score,freshness,row_count,'
                    'pk_gate,null_gate,freshness_gate,pii_gate,row_min_gate,ml_ready) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)',
                    [(run_id,) + t for t in sim_tables],
                )
                conn.executemany(
                    'INSERT INTO semantic_definitions (run_id,type,name,definition,confidence,explore,status) VALUES (?,?,?,?,?,?,?)',
                    [(run_id,) + d + ('pending',) for d in sim_defs],
                )
            tables_count = len(sim_tables)
            defs_count   = len(sim_defs)
            low_conf     = sum(1 for d in sim_defs if d[3] < 0.70)
            conn.execute(
                "UPDATE governance_runs SET tables_count=?,definitions_count=?,low_confidence_count=?,"
                "completed_at=datetime('now') WHERE id=?",
                (tables_count, defs_count, low_conf, run_id),
            )
            conn.commit()
            log_action_bg(conn, 'governance.completed', 'governance_run', run_id,
                          {'tables_count': tables_count, 'definitions_count': defs_count})

            # Build + persist versioned governance manifest (Sprint 2 / F-007)
            import manifest as manifest_mod
            conn_id_row = conn.execute(
                'SELECT connection_id FROM governance_runs WHERE id=?', (run_id,)).fetchone()
            connection_id = conn_id_row['connection_id'] if conn_id_row else None
            mtables = [{
                'name': t[0], 'schema_name': t[1], 'health_score': t[2], 'freshness': t[3],
                'row_count': t[4], 'pk_gate': t[5], 'null_gate': t[6], 'freshness_gate': t[7],
                'pii_gate': t[8], 'row_min_gate': t[9],
                'columns': SIM_TABLE_COLUMNS.get(t[0], []),
            } for t in sim_tables]
            mdefs = [{'type': d[0], 'name': d[1], 'definition': d[2], 'confidence': d[3]}
                     for d in sim_defs]
            # R3S1E2: configured freshness SLAs override the static heuristic
            import dq as dq_mod
            slas = {r['table_name']: r['max_age_hours'] for r in conn.execute(
                'SELECT table_name, max_age_hours FROM freshness_slas '
                'WHERE connection_id=?', (connection_id,)).fetchall()}
            sla_violations = []
            for t in mtables:
                sla = slas.get(t['name'])
                if sla is None:
                    continue
                age = dq_mod.parse_freshness_age(t.get('freshness'))
                if age is None or age > sla:
                    t['freshness_gate'] = 'warn'
                    sla_violations.append({'table': t['name'], 'age_hours': age,
                                           'sla_hours': sla})
                else:
                    t['freshness_gate'] = 'pass'

            m = manifest_mod.build_manifest(connection_id, run_id, mtables, mdefs)

            # R3S1E4: data contract enforcement
            import dq as _dq_mod
            contracts = {r['table_name']: dict(r) for r in conn.execute(
                'SELECT * FROM data_contracts WHERE connection_id=?',
                (connection_id,)).fetchall()}
            contract_alerts = []
            for t in m['tables']:
                c = contracts.get(t['name'])
                if not c:
                    continue
                violations = []
                have = {col['name'] for col in t.get('columns') or []}
                for req in json.loads(c['required_columns_json'] or '[]'):
                    if req not in have:
                        violations.append(f'required column {req!r} is missing')
                if c['min_rows'] and _dq_mod.row_count_to_int(t.get('row_count')) < c['min_rows']:
                    violations.append(f"row count below contract minimum {c['min_rows']}")
                if c['max_age_hours']:
                    age = _dq_mod.parse_freshness_age(t.get('freshness'))
                    if age is None or age > c['max_age_hours']:
                        violations.append(f"freshness exceeds contract SLA {c['max_age_hours']}h")
                t['contract_violations'] = violations
                if violations:
                    t['dq_gate_status'] = 'BLOCK'
                    m['human_review_required'] = True
                    m['dq_gate_status'] = 'BLOCK'
                    contract_alerts.append((t['name'], violations))
            for tname, violations in contract_alerts:
                conn.execute('INSERT INTO alerts (type, connection_id, subject, detail_json) '
                             "VALUES ('contract', ?, ?, ?)",
                             (connection_id, f'Data contract violated on {tname}',
                              json.dumps({'table': tname, 'violations': violations,
                                          'run_id': run_id})))
            conn.commit()

            _save_manifest_with_drift(conn, connection_id, m)

            for v in sla_violations:
                conn.execute('INSERT INTO alerts (type, connection_id, subject, detail_json) '
                             "VALUES ('freshness', ?, ?, ?)",
                             (connection_id,
                              f"Table {v['table']} exceeds its freshness SLA",
                              json.dumps({**v, 'run_id': run_id})))
            conn.commit()
            if sla_violations:
                send_email(to='workspace-admins@analytiq.dev',
                           subject=f"{len(sla_violations)} freshness SLA violation(s)",
                           html='<p>' + ', '.join(v['table'] for v in sla_violations) + '</p>')

            # R3S1E1: health trend history + threshold alerting
            for t in m['tables']:
                conn.execute('INSERT INTO health_history (connection_id, run_id, table_name, '
                             'health_score) VALUES (?,?,?,?)',
                             (connection_id, run_id, t['name'], t['health_score']))
            conn.commit()
            thr = conn.execute('SELECT min_health FROM governance_thresholds '
                               'WHERE connection_id=?', (connection_id,)).fetchone()
            if thr:
                breaches = [t for t in m['tables'] if (t['health_score'] or 0) < thr['min_health']]
                for t in breaches:
                    conn.execute('INSERT INTO alerts (type, connection_id, subject, detail_json) '
                                 "VALUES ('health', ?, ?, ?)",
                                 (connection_id,
                                  f"Table {t['name']} health {t['health_score']} below "
                                  f"threshold {thr['min_health']}",
                                  json.dumps({'table': t['name'],
                                              'health_score': t['health_score'],
                                              'threshold': thr['min_health'],
                                              'run_id': run_id})))
                conn.commit()
                if breaches:
                    send_email(to='workspace-admins@analytiq.dev',
                               subject=f'{len(breaches)} table(s) below health threshold',
                               html='<p>' + ', '.join(
                                   f"{t['name']} ({t['health_score']})" for t in breaches) +
                                    f" fell below {thr['min_health']}.</p>")
            log_action_bg(conn, 'governance.manifest_created', 'manifest', connection_id,
                          {'version': m['manifest_version'], 'run_id': run_id,
                           'dq_gate_status': m['dq_gate_status']})

            # Emit typed agent events before the final step event closes streams
            for t in m['tables']:
                broadcast_gov(run_id, {'type': 'dq_gate', 'run_id': run_id,
                                       'table': t['name'], 'gate_status': t['dq_gate_status'],
                                       'gates': t['gates']})
            if m['human_review_required']:
                broadcast_gov(run_id, {'type': 'human_required', 'run_id': run_id,
                                       'reason': 'low-confidence definitions or blocked tables'})
            broadcast_gov(run_id, {'type': 'agent_complete', 'run_id': run_id,
                                   'manifest_version': m['manifest_version']})
            broadcast_gov(run_id, {'step': 4, 'status': 'done'})

            row = conn.execute(
                'SELECT c.owner_email, c.name AS conn_name '
                'FROM governance_runs g JOIN connections c ON c.id = g.connection_id '
                'WHERE g.id = ?', (run_id,),
            ).fetchone()
            if row and row[0]:
                send_email(
                    to=row[0],
                    subject='Governance scan complete',
                    html=(
                        f'<h2>Governance scan finished</h2>'
                        f'<p>Connection <strong>{row[1]}</strong> has been scanned.</p>'
                        f'<ul>'
                        f'<li><strong>{tables_count}</strong> tables cataloged</li>'
                        f'<li><strong>{defs_count}</strong> semantic definitions generated</li>'
                        f'<li><strong>{low_conf}</strong> low-confidence definitions need review</li>'
                        f'</ul>'
                        f'<p>Log in to AnalytIQ to review the results.</p>'
                    ),
                )
        finally:
            put_db(conn)

    threading.Thread(target=_run, daemon=True).start()


# ─────────────────────────────────────────────────────────
# Pipeline simulation
# ─────────────────────────────────────────────────────────
PIPE_DELAYS = [d * _SIM_SCALE for d in (3.0, 5.0, 4.5, 2.5)]
PIPE_LOG_DELAY = 0.28 * _SIM_SCALE

# R4S1E3: plain-English step cards with input/output schemas
PIPE_STEP_CARDS = {
    1: {'label': 'Build gold table & features',
        'description': 'Joined fact_revenue with dim_location at the location · day '
                       'grain and verified the time-based split.',
        'input_schema': ['fact_revenue(revenue_id, location_id, day, net_revenue)',
                         'dim_location(location_id, tier)'],
        'output_schema': ['gold(location_id, day, features…, target_net_revenue)']},
    2: {'label': 'Train candidate models',
        'description': 'Fitted the model on the training window with the confirmed '
                       'hyperparameter grid.',
        'input_schema': ['gold train split (Jan–Sep 2023)'],
        'output_schema': ['fitted model + validation predictions']},
    3: {'label': 'Walk-forward validation',
        'description': 'Backtested across 5 expanding windows and checked the '
                       'MAPE and overfit gates.',
        'input_schema': ['fitted model', 'validation folds'],
        'output_schema': ['fold metrics (MAPE per window)']},
    4: {'label': 'Score & generate artifact',
        'description': 'Wrote predictions to the gold layer and generated the '
                       'self-contained dashboard artifact.',
        'input_schema': ['promoted model', 'gold table'],
        'output_schema': ['gold.predictions', 'artifact HTML']},
}
PIPE_LOGS   = {
    1: [
        '[INFO] Gold table: fact_revenue × dim_location · grain=location_day',
        '[INFO] 12,847 rows · 34 features · time-based split verified',
        '[DQ]  PK uniqueness: ✓ PASS',
        '[DQ]  Key null rate < 5%: ✓ PASS (max 1.2%)',
        '[INFO] Leakage scan: clean (no future-date features)',
    ],
    2: [
        '[INFO] XGBoost training started · n_estimators=500 · max_depth=6',
        '[INFO] Train split: 8,993 rows (Jan–Sep 2023)',
        '[INFO] Val split:   1,927 rows (Oct 2023)',
        '[INFO] Test split:  1,927 rows (Nov–Dec 2023)',
    ],
    3: [
        '[INFO] Fold 1/5 → MAPE 7.4%',
        '[INFO] Fold 2/5 → MAPE 8.1%',
        '[INFO] Fold 3/5 → MAPE 9.8%',
        '[INFO] Fold 4/5 → MAPE 8.3%',
        '[INFO] Fold 5/5 → MAPE 7.9%',
        '[INFO] Validation MAPE 8.9% < 15.0% threshold: ✓ PASS',
        '[INFO] Overfit check: test 9.1% ≤ val 8.9% × 1.2: ✓ PASS',
    ],
    4: [
        '[INFO] Writing 12,847 predictions → gold.net_revenue_preds_v1',
        '[DQ]  Distribution gate: ✓ PASS (KS p=0.42)',
        '[INFO] Generating self-contained dashboard artifact...',
        '[DONE] Pipeline complete · 3m 41s · model_id=xgb-locrev-v1',
    ],
}


def _uas_ns(sess) -> str:
    """R9S2E6: sandbox sessions live in an isolated 'sandbox:'-prefixed
    namespace — excluded from production indices and caches."""
    return 'sandbox:default' if (sess and sess.get('is_sandbox')) else 'default'


def _uas_context_versions(conn, connection_id=None):
    """Current governed context for UAS content addressing (R8S1E1)."""
    gov = '0.0.0'
    if connection_id:
        row = conn.execute('SELECT version FROM governance_manifests WHERE connection_id=? '
                           'ORDER BY id DESC LIMIT 1', (connection_id,)).fetchone()
        if row:
            gov = row['version']
    sem_row = conn.execute("SELECT version FROM semantic_schemas WHERE workspace_id='default' "
                           'ORDER BY id DESC LIMIT 1').fetchone()
    return gov, (sem_row['version'] if sem_row else '0.0.0')


def _register_pipeline_uas_chain(conn, run_id, sess_id):
    """R8S1E1-US1: register session_spec -> dashboard_plan ->
    (gold_predictions_ref, gold_forecast_ref) in the Unified Artifact Store.
    The artifact_html_ref node is appended at artifact save/render time."""
    import uas
    sess = conn.execute('SELECT * FROM sessions WHERE id=?', (sess_id,)).fetchone() if sess_id else None
    if not sess:
        return
    gov, sem = _uas_context_versions(conn, sess['connection_id'])
    ns = _uas_ns(dict(sess))
    spec_payload = {'metric': sess['metric'], 'grain': sess['grain'],
                    'horizon': sess['horizon'], 'training_start': sess['training_start'],
                    'training_end': sess['training_end'],
                    'connection_id': sess['connection_id']}
    spec = uas.register(conn, 'session_spec', spec_payload,
                        logical_key=f"{ns}:session_spec:s{sess_id}",
                        gov_version=gov, sem_version=sem, agent='session_planner',
                        workspace_id=ns, run_id=run_id)
    plan_payload = {'title': f"{sess['metric']} Forecast", 'metric': sess['metric'],
                    'grain': sess['grain'], 'horizon': sess['horizon'],
                    'panels': ['kpi_row', 'timeseries_ci', 'feature_importance',
                               'dimension_breakdown', 'forecast', 'trial_leaderboard',
                               'dq_lineage_footer', 'header_bar']}
    plan = uas.register(conn, 'dashboard_plan', plan_payload,
                        logical_key=f"{ns}:dashboard_plan:s{sess_id}",
                        upstream=[spec['artifact_uid']],
                        gov_version=gov, sem_version=sem, agent='dashboard_planner',
                        workspace_id=ns, run_id=run_id)
    counts = {t: conn.execute(f'SELECT COUNT(*) c FROM {t} WHERE pipeline_run_id=?',
                              (run_id,)).fetchone()['c']
              for t in ('gold_predictions', 'gold_forecast')}
    for table, atype, agent in (('gold_predictions', 'gold_predictions_ref', 'data_modeler'),
                                ('gold_forecast', 'gold_forecast_ref', 'data_modeler')):
        uas.register(conn, atype,
                     {'table': table, 'row_count': counts[table], 'run_id': run_id},
                     logical_key=f"{ns}:{atype}:r{run_id}",
                     upstream=[plan['artifact_uid']],
                     gov_version=gov, sem_version=sem, agent=agent,
                     workspace_id=ns, run_id=run_id)


def _persist_repair_attempts(conn, scope, artifact_id, run_id, attempts, resolved):
    """R11S2E3: failed repair attempts are retained even when a later attempt
    succeeded — the replay debugger shows what the repair loop tried."""
    for a in attempts:
        conn.execute('INSERT INTO repair_attempts (scope, artifact_id, run_id, cycle, '
                     'detail_json, resolved) VALUES (?,?,?,?,?,?)',
                     (scope, artifact_id, run_id, a['cycle'],
                      json.dumps({'failed_checks': a['failed_checks']}),
                      1 if resolved else 0))
    if attempts:
        conn.commit()


def _register_artifact_html_uas(conn, artifact_id, run_id, title, file_info):
    """R8S1E1-US1: terminal artifact_html_ref node, upstream of the run's
    gold refs. Called from artifact save and re-render paths."""
    import uas
    sess_row = conn.execute('SELECT s.* FROM sessions s JOIN pipeline_runs p ON p.session_id=s.id '
                            'WHERE p.id=?', (run_id,)).fetchone()
    gov, sem = _uas_context_versions(conn, sess_row['connection_id'] if sess_row else None)
    ns = _uas_ns(dict(sess_row) if sess_row else None)
    upstream = []
    for atype in ('gold_predictions_ref', 'gold_forecast_ref'):
        node = conn.execute('SELECT artifact_uid FROM uas_artifacts WHERE logical_key=? '
                            'ORDER BY version DESC LIMIT 1',
                            (f"{ns}:{atype}:r{run_id}",)).fetchone()
        if node:
            upstream.append(node['artifact_uid'])
    payload = {'artifact_id': artifact_id, 'title': title}
    if file_info:
        payload.update({'sha256': file_info.get('sha256'),
                        'size_bytes': file_info.get('size_bytes'),
                        'validation_status': file_info.get('validation_status'),
                        'file_version': file_info.get('version')})
    uas.register(conn, 'artifact_html_ref', payload,
                 logical_key=f"{ns}:artifact_html_ref:a{artifact_id}",
                 upstream=upstream, gov_version=gov, sem_version=sem,
                 agent='artifact_assembler', workspace_id=ns, run_id=run_id)


def _concurrency_budget(conn) -> int:
    """R9S1E2: per-workspace worker budget (default 4, admin-configurable)."""
    row = conn.execute("SELECT value FROM platform_settings WHERE key='concurrency_budget'").fetchone()
    try:
        return max(1, int(row['value'])) if row else 4
    except (TypeError, ValueError):
        return 4


def _register_viz_specs_uas(conn, run_id, sess_id):
    """viz_specs node output: deterministic Vega-Lite spec summaries for the
    descriptive panels, registered in the UAS (validated by the
    spec_validation edge gate before assembly)."""
    import uas
    sess = conn.execute('SELECT * FROM sessions WHERE id=?', (sess_id,)).fetchone() if sess_id else None
    gov, sem = _uas_context_versions(conn, sess['connection_id'] if sess else None)
    ns = _uas_ns(dict(sess) if sess else None)
    plan_uid = conn.execute('SELECT artifact_uid FROM uas_artifacts WHERE logical_key=? '
                            'ORDER BY version DESC LIMIT 1',
                            (f'{ns}:dashboard_plan:s{sess_id}',)).fetchone()
    # R9S2E5: consult the semantic-layer agent for the canonical display
    # format instead of guessing and risking a spec-validation repair cycle.
    import agent_bus
    fmt = agent_bus.consult(
        conn, 'visualization_agent', 'semantic_layer_agent',
        {'kind': 'metric_format', 'metric': (sess['metric'] if sess else 'value')},
        run_id=run_id,
        broadcaster=lambda payload: broadcast_pipe(run_id, {**payload, 'step': 0,
                                                            'status': 'running'}))
    specs = [{'panel': p, 'mark': m, 'schema': 'https://vega.github.io/schema/vega-lite/v5.json'}
             for p, m in (('timeseries_ci', 'line'), ('dimension_breakdown', 'bar'),
                          ('forecast', 'area'), ('feature_importance', 'bar'))]
    uas.register(conn, 'vega_lite_specs',
                 {'specs': specs, 'validated': True, 'metric_format': fmt['format']},
                 logical_key=f'{ns}:vega_lite_specs:s{sess_id}',
                 upstream=[plan_uid['artifact_uid']] if plan_uid else [],
                 gov_version=gov, sem_version=sem, agent='visualization_agent',
                 workspace_id=ns, run_id=run_id)


def _persist_component_contracts(conn, run_id):
    """R17S1E1 (arch §7.2/§7.3): per-component query + data contracts,
    computed from the rows the run actually produced."""
    comps = {
        'kpi_row': ('SELECT SUM(actual) AS total, AVG(actual) AS avg_daily FROM gold_predictions '
                    f'WHERE pipeline_run_id={run_id}', 100,
                    'SELECT actual FROM gold_predictions WHERE pipeline_run_id=? AND actual IS NOT NULL'),
        'timeseries_ci': ('SELECT day_index, date, actual, predicted, ci_low, ci_high '
                          f'FROM gold_predictions WHERE pipeline_run_id={run_id} ORDER BY day_index',
                          10000,
                          'SELECT actual, predicted FROM gold_predictions WHERE pipeline_run_id=?'),
        'forecast': ('SELECT day_index, date, predicted, ci_low, ci_high FROM gold_forecast '
                     f'WHERE pipeline_run_id={run_id} ORDER BY day_index', 10000,
                     'SELECT predicted FROM gold_forecast WHERE pipeline_run_id=?'),
    }
    import time as _t
    for comp, (sql, limit, stats_sql) in comps.items():
        t0 = _t.time()
        rows = conn.execute(stats_sql, (run_id,)).fetchall()
        elapsed = int((_t.time() - t0) * 1000)
        cols = list(rows[0].keys()) if rows else []
        expected = [{'name': c, 'type': 'number'} for c in cols]
        cur = conn.execute(
            'INSERT INTO query_contracts (run_id, component_id, sql, warehouse_dialect, '
            "expected_columns_json, row_limit, status, execution_time_ms) "
            "VALUES (?,?,?,?,?,?, 'executed', ?)",
            (run_id, comp, sql, 'sqlite', json.dumps(expected), limit, elapsed))
        qc_id = cur.lastrowid
        ranges = {}
        for c in cols:
            vals = [r[c] for r in rows if isinstance(r[c], (int, float))]
            if vals:
                ranges[c] = {'min': min(vals), 'max': max(vals),
                             'mean': round(sum(vals) / len(vals), 2)}
        actual_cols = [{'name': c, 'type': 'number',
                        'null_count': sum(1 for r in rows if r[c] is None),
                        'distinct_count': len({r[c] for r in rows})} for c in cols]
        conn.execute(
            'INSERT INTO component_data_contracts (run_id, component_id, query_contract_id, '
            'actual_columns_json, row_count, numeric_ranges_json, empty_result) '
            'VALUES (?,?,?,?,?,?,?)',
            (run_id, comp, qc_id, json.dumps(actual_cols), len(rows),
             json.dumps(ranges), 1 if not rows else 0))
    conn.commit()


def _write_gold_outputs(conn, run_id, sess_id):
    """gold_build node work: chart data + PRD gold output tables (R6S1E1)."""
    conn.executemany(
        'INSERT INTO chart_data (pipeline_run_id,day_index,date,actual,predicted,ci_low,ci_high,is_forecast) '
        'VALUES (:pipeline_run_id,:day_index,:date,:actual,:predicted,:ci_low,:ci_high,:is_forecast)',
        generate_chart_data(run_id),
    )
    for r in generate_chart_data(run_id):
        if r['actual'] is not None:
            conn.execute('INSERT INTO gold_predictions (pipeline_run_id, session_id, '
                         'day_index, date, actual, predicted, ci_low, ci_high) '
                         'VALUES (?,?,?,?,?,?,?,?)',
                         (run_id, sess_id, r['day_index'], r['date'], r['actual'],
                          r['predicted'], r['ci_low'], r['ci_high']))
        if r['is_forecast']:
            conn.execute('INSERT INTO gold_forecast (pipeline_run_id, session_id, '
                         'day_index, date, predicted, ci_low, ci_high) '
                         'VALUES (?,?,?,?,?,?,?)',
                         (run_id, sess_id, r['day_index'], r['date'],
                          r['predicted'], r['ci_low'], r['ci_high']))
    conn.commit()


def _copy_gold_outputs(conn, run_id, sess_id, prior_run_id):
    """Cached gold_build node: serve outputs from the store — copy the
    content-identical prior run's rows instead of recomputing (§17.2.1)."""
    conn.execute('INSERT INTO chart_data (pipeline_run_id,day_index,date,actual,predicted,'
                 'ci_low,ci_high,is_forecast) '
                 'SELECT ?,day_index,date,actual,predicted,ci_low,ci_high,is_forecast '
                 'FROM chart_data WHERE pipeline_run_id=?', (run_id, prior_run_id))
    conn.execute('INSERT INTO gold_predictions (pipeline_run_id,session_id,day_index,date,'
                 'actual,predicted,ci_low,ci_high) '
                 'SELECT ?,?,day_index,date,actual,predicted,ci_low,ci_high '
                 'FROM gold_predictions WHERE pipeline_run_id=?', (run_id, sess_id, prior_run_id))
    conn.execute('INSERT INTO gold_forecast (pipeline_run_id,session_id,day_index,date,'
                 'predicted,ci_low,ci_high) '
                 'SELECT ?,?,day_index,date,predicted,ci_low,ci_high '
                 'FROM gold_forecast WHERE pipeline_run_id=?', (run_id, sess_id, prior_run_id))
    conn.commit()


def _arbitrate_gold_grain(conn, run_id, sess):
    """R9S2E4: if the data-modeler's latest gold table for this session was
    built at a different grain than the session spec, the meta-orchestrator
    arbitrates deterministically (session grain canonical) instead of letting
    the disagreement become a repair-loop failure."""
    import meta_orchestrator as meta
    from modeler import _slug
    gt = conn.execute("SELECT * FROM gold_tables WHERE session_id=? "
                      'ORDER BY id DESC LIMIT 1', (sess.get('id'),)).fetchone()
    if not gt:
        return
    sess_slug = _slug(sess.get('grain') or '')
    if not sess_slug or sess_slug in (gt['table_name'] or ''):
        return
    import re as _re
    m = _re.search(r'gold_[a-z0-9_]*?_((?:location|region|store|customer|product|day|week|month|hour)'
                   r'(?:_[a-z0-9]+)*)_v\d+$', gt['table_name'] or '')
    modeler_slug = m.group(1) if m else (gt['table_name'] or 'unknown')
    meta.arbitrate_grain(conn, run_id, dict(sess), modeler_slug, sess_slug)


def _notify_systemic_alert(conn):
    """Email the workspace admin exactly once per systemic-failure alert."""
    row = conn.execute("SELECT id, detail_json FROM alerts WHERE type='meta.systemic_failure' "
                       'ORDER BY id DESC LIMIT 1').fetchone()
    if not row:
        return
    sent = conn.execute("SELECT 1 FROM email_outbox WHERE subject LIKE '%systemic%' "
                        "AND body_html LIKE ?", (f'%alert #{row["id"]}%',)).fetchone()
    if sent:
        return
    send_email(to='admin@acme.com',
               subject='AnalytIQ: systemic pipeline failure pattern detected',
               html=f'<p>Meta-orchestrator triage raised alert #{row["id"]}: '
                    f'repeated gate exhaustion across runs. Detail: {row["detail_json"]}</p>')


def simulate_pipeline(run_id: int):
    """R8S2E3-US1: DAG-native pipeline execution (§17.2.1). Topological walk
    over content-addressed nodes; gates are edge contracts; cached nodes are
    served from the store. SSE/step/log contracts preserved."""
    def _run():
        import dag
        conn     = thread_db()
        all_logs = []
        try:
            sess_row = conn.execute('SELECT session_id FROM pipeline_runs WHERE id=?',
                                    (run_id,)).fetchone()
            sess_id = sess_row['session_id'] if sess_row else None
            sess = dict(conn.execute('SELECT * FROM sessions WHERE id=?', (sess_id,)).fetchone()
                        or {'metric': 'Net Revenue', 'grain': 'Location · Day', 'horizon': 14,
                            'training_start': None, 'training_end': None, 'connection_id': None})
            gov, sem = _uas_context_versions(conn, sess.get('connection_id'))
            hashes = dag.compute_hashes(sess, gov, sem, namespace=_uas_ns(sess))
            cache_plan = dag.create_run_nodes(conn, run_id, hashes)
            gate_ctx = {'session': sess}

            # R9S1E2: wave-based topological executor — independent branches
            # run concurrently in a worker pool under the workspace
            # concurrency budget; a BLOCK halts only true descendants.
            from concurrent.futures import ThreadPoolExecutor
            budget = _concurrency_budget(conn)
            status_map = {k: 'pending' for k in dag.NODE_KEYS}
            log_lock = threading.Lock()

            def append_log(line, step=None, node_key=None, extra=None):
                with log_lock:
                    all_logs.append(line)
                    conn.execute('UPDATE pipeline_runs SET log_entries=? WHERE id=?',
                                 (json.dumps(all_logs), run_id))
                    conn.commit()
                    payload = {'step': step or 0, 'status': 'running', 'log': list(all_logs)}
                    if node_key:
                        payload['node_key'] = node_key
                    if extra:
                        payload.update(extra)
                    broadcast_pipe(run_id, payload)

            def block_from(node_key, results):
                reason = '; '.join(f"{g['gate']}={g['status']}" for g in results)
                append_log(f'✗ Edge gate BLOCK before {node_key}: {reason}', node_key=node_key)
                for key in [node_key] + dag.downstream_of(node_key):
                    status_map[key] = 'blocked'
                    dag.mark(conn, run_id, key, 'blocked')
                log_action_bg(conn, 'pipeline.gate_blocked', 'pipeline_run', run_id,
                              {'blocked_node': node_key, 'gates': reason})
                import meta_orchestrator as meta
                meta.triage_failure(conn, run_id, reason)
                _notify_systemic_alert(conn)

            def execute_node(node_key):
                # own connection per worker thread (WAL + busy_timeout)
                wconn = thread_db()
                try:
                    node = wconn.execute('SELECT * FROM dag_nodes WHERE run_id=? AND node_key=?',
                                         (run_id, node_key)).fetchone()
                    step = next((s for s, k in dag.STEP_NODES.items() if k == node_key), None)
                    dag.mark(wconn, run_id, node_key, 'running', started=True)
                    if step is not None:
                        wconn.execute('UPDATE pipeline_runs SET current_step=?, status=? WHERE id=?',
                                      (step, 'running', run_id))
                        card = PIPE_STEP_CARDS[step]
                        wconn.execute('INSERT INTO pipeline_steps (run_id, step, node_key, label, '
                                      'description, input_schema_json, output_schema_json) '
                                      'VALUES (?,?,?,?,?,?,?)',
                                      (run_id, step, node_key, card['label'], card['description'],
                                       json.dumps(card['input_schema']),
                                       json.dumps(card['output_schema'])))
                        wconn.commit()

                    work = dag.NODE_WORK.get(node_key)
                    if node['cached'] and node['prior_run_id'] is not None:
                        if node_key == 'gold_build':
                            _copy_gold_outputs(wconn, run_id, sess_id, node['prior_run_id'])
                            _persist_component_contracts(wconn, run_id)
                        if node_key == 'viz_specs':
                            # store semantics: content is identical — materialize
                            # this session's logical mapping (idempotent register)
                            _register_viz_specs_uas(wconn, run_id, sess_id)
                        append_log(f'⚡ {node_key}: served from store '
                                   f"(content-identical to run {node['prior_run_id']})",
                                   step=step, node_key=node_key, extra={'cached': True})
                    else:
                        if node_key == 'gold_build':
                            _arbitrate_gold_grain(wconn, run_id, sess)
                            _write_gold_outputs(wconn, run_id, sess_id)
                            _persist_component_contracts(wconn, run_id)
                        if node_key == 'viz_specs':
                            _register_viz_specs_uas(wconn, run_id, sess_id)
                        if work:
                            work(wconn, run_id, gate_ctx)
                        if step is not None:
                            for line in PIPE_LOGS.get(step, []):
                                append_log(line, step=step, node_key=node_key)
                                time.sleep(PIPE_LOG_DELAY)
                            time.sleep(PIPE_DELAYS[step - 1])
                    dag.mark(wconn, run_id, node_key, 'done')
                    status_map[node_key] = 'done'
                finally:
                    put_db(wconn)
                    try:
                        wconn.close()
                    except Exception:
                        pass

            with ThreadPoolExecutor(max_workers=budget) as pool:
                while True:
                    ready = [k for k in dag.NODE_KEYS
                             if status_map[k] == 'pending'
                             and all(status_map[u] == 'done' for u in dag.upstreams_of(k))]
                    if not ready:
                        break
                    runnable = []
                    for key in ready:
                        blocked = False
                        for f in dag.upstreams_of(key):
                            worst, results = dag.evaluate_edge(conn, run_id, f, key, gate_ctx)
                            if worst == 'BLOCK':
                                block_from(key, results)
                                blocked = True
                                break
                        if not blocked and status_map[key] == 'pending':
                            runnable.append(key)
                    if runnable:
                        list(pool.map(execute_node, runnable))

            if any(v == 'blocked' for v in status_map.values()):
                conn.execute("UPDATE pipeline_runs SET status='failed', log_entries=?, "
                             "completed_at=datetime('now') WHERE id=?",
                             (json.dumps(all_logs), run_id))
                conn.commit()
                broadcast_pipe(run_id, {'step': 0, 'status': 'error', 'log': list(all_logs)})
                return

            conn.execute(
                "UPDATE pipeline_runs SET status='done',current_step=4,mape=8.9,features_count=34,"
                "rows_count=12847,log_entries=?,completed_at=datetime('now') WHERE id=?",
                (json.dumps(all_logs), run_id),
            )
            conn.commit()
            _register_pipeline_uas_chain(conn, run_id, sess_id)
            # lineage = execution: link nodes to their UAS artifacts
            for node_key, atype in (('session_plan', 'dashboard_plan'),
                                    ('gold_build', 'gold_predictions_ref')):
                row = conn.execute('SELECT artifact_uid FROM uas_artifacts WHERE run_id=? '
                                   'AND artifact_type=? ORDER BY id DESC LIMIT 1',
                                   (run_id, atype)).fetchone()
                if row is None and atype == 'dashboard_plan':
                    row = conn.execute('SELECT artifact_uid FROM uas_artifacts WHERE '
                                       "artifact_type='dashboard_plan' AND logical_key=? "
                                       'ORDER BY id DESC LIMIT 1',
                                       (f'default:dashboard_plan:s{sess_id}',)).fetchone()
                if row:
                    dag.set_uas_ref(conn, run_id, node_key, row['artifact_uid'])
            log_action_bg(conn, 'pipeline.completed', 'pipeline_run', run_id,
                          {'mape': 8.9, 'features_count': 34, 'rows_count': 12847})
            broadcast_pipe(run_id, {'step': 4, 'status': 'done',
                                    'node_key': 'artifact_ready', 'log': all_logs})
        finally:
            put_db(conn)

    threading.Thread(target=_run, daemon=True).start()


# ─────────────────────────────────────────────────────────
# Event-driven job handlers (R9S1E3) — targeted recompute + retrain
# ─────────────────────────────────────────────────────────
def _job_event_recompute(conn, payload):
    sid = payload.get('session_id')
    if not sid:
        return
    lid = conn.execute('INSERT INTO pipeline_runs (session_id,status,current_step,log_entries) '
                       "VALUES (?, 'running', 0, '[]')", (sid,)).lastrowid
    conn.commit()
    log_action_bg(conn, 'pipeline.event_recompute', 'pipeline_run', lid,
                  {'session_id': sid, 'event_id': payload.get('event_id')})
    simulate_pipeline(lid)


def _job_event_retrain(conn, payload):
    sid = payload.get('session_id')
    if not sid:
        return
    log_action_bg(conn, 'model.event_retrain_queued', 'session', sid,
                  {'event_id': payload.get('event_id')})


def _job_optimization_scan(conn, payload):
    import optimizer
    n = optimizer.analyze(conn)
    log_action_bg(conn, 'optimization.scanned', 'platform', None, {'new_proposals': n})


def _register_event_job_handlers():
    import jobs
    jobs.register('event_recompute', _job_event_recompute)
    jobs.register('event_retrain', _job_event_retrain)
    jobs.register('optimization_scan', _job_optimization_scan)
    jobs.register('self_improve_scan',
                  lambda conn, payload: __import__('self_improve').mine(conn))


_register_event_job_handlers()


# ─────────────────────────────────────────────────────────
# SSE helper
# ─────────────────────────────────────────────────────────
def sse_response(generator_fn):
    resp = Response(stream_with_context(generator_fn()), mimetype='text/event-stream')
    resp.headers['Cache-Control']      = 'no-cache'
    resp.headers['X-Accel-Buffering']  = 'no'
    resp.headers['Transfer-Encoding']  = 'chunked'
    return resp


# ─────────────────────────────────────────────────────────
# Routes — Health
# ─────────────────────────────────────────────────────────
@app.get('/api/health')
def health():
    return jsonify({'ok': True})


@app.get('/api/platform/status')
def platform_status():
    """Which managed tool vs local fallback each platform service is using."""
    import authn
    import secrets_store

    def svc(mode):
        return {'mode': mode, 'fallback_active': mode == 'local'}

    return jsonify({
        'auth':    svc(authn.provider_mode() if authn.provider_mode() == 'local'
                       else authn.provider_mode()),
        'secrets': svc(secrets_store.provider_mode()),
        'queue':   svc('upstash' if os.environ.get('UPSTASH_REDIS_REST_URL') else 'local'),
        'storage': svc('r2' if os.environ.get('R2_ACCOUNT_ID') else 'local'),
        'search':  svc('meilisearch' if os.environ.get('MEILI_HOST') else 'local'),
        'email':   svc('resend' if RESEND_API_KEY else 'local'),
        'logging': svc('betterstack' if os.environ.get('LOGTAIL_TOKEN') else 'local'),
        'cache':   svc(__import__('cache_hier').provider_mode()),
    })


@app.post('/api/semantic/evolve')
def semantic_evolve():
    """R10S2E5: run the semantic-evolution proposal engine now."""
    import semantic_evolution as se
    n = se.propose(get_db())
    log_action('semantic.evolution_scanned', 'platform', None, {'new_proposals': n})
    return jsonify({'new_proposals': n})


@app.get('/api/semantic/proposals')
def list_semantic_proposals():
    return jsonify({'proposals': many('SELECT * FROM semantic_proposals ORDER BY id DESC LIMIT 100')})


def _decide_semantic_proposal(pid, decision):
    row = one('SELECT * FROM semantic_proposals WHERE id=?', (pid,))
    if not row:
        return jsonify({'error': 'Proposal not found'}), 404
    if row['status'] != 'proposed':
        return jsonify({'error': f"Proposal already {row['status']}"}), 409
    execute("UPDATE semantic_proposals SET status=?, decided_by=?, decided_at=datetime('now') "
            'WHERE id=?', (decision, getattr(g, 'user_email', None) or 'admin', pid))
    import feedback_loop as fb
    fb.record(get_db(), 'semantic_proposal', pid,
              'accept' if decision == 'approved' else 'dismiss',
              category=row['kind'], user=getattr(g, 'user_email', None) or 'default')
    # §17.3.4: approval records the decision — the canonical schema is only
    # ever changed through the explicit semantic-layer editing flow.
    log_action(f'semantic.proposal_{decision}', 'semantic_proposal', pid,
               {'kind': row['kind'], 'subject': row['subject']})
    return jsonify({'id': pid, 'status': decision})


@app.post('/api/semantic/proposals/<int:pid>/approve')
@require_role('admin')
def approve_semantic_proposal(pid):
    return _decide_semantic_proposal(pid, 'approved')


@app.post('/api/semantic/proposals/<int:pid>/reject')
@require_role('admin')
def reject_semantic_proposal(pid):
    return _decide_semantic_proposal(pid, 'rejected')


@app.get('/api/kg/related')
def kg_related():
    """R10S1E2: related metrics ranked by knowledge-graph edge weight."""
    import knowledge_graph as kg
    metric = request.args.get('metric')
    if not metric:
        return jsonify({'error': 'metric required'}), 400
    return jsonify({'metric': metric, 'related': kg.related_metrics(get_db(), metric)})


@app.get('/api/kg/co_analysis')
def kg_co_analysis():
    import knowledge_graph as kg
    metric = request.args.get('metric')
    if not metric:
        return jsonify({'error': 'metric required'}), 400
    return jsonify(kg.co_analysis(get_db(), metric))


@app.post('/api/kg/rebuild')
def kg_rebuild():
    import knowledge_graph as kg
    n = kg.rebuild(get_db())
    log_action('kg.rebuilt', 'platform', None, {'edges_touched': n})
    return jsonify({'edges_touched': n})


@app.post('/api/memory')
def post_memory():
    """R10S1E1: explicit memory write (PII-gated, audited)."""
    import agent_memory as am
    b = request.get_json() or {}
    try:
        row = am.remember(get_db(), b.get('agent', 'planner'), b.get('category', ''),
                          b.get('key', ''), b.get('value', ''),
                          user=b.get('user', 'default'))
    except ValueError as e:
        return jsonify({'error': str(e)}), 400
    return jsonify(row), 201


@app.get('/api/memory')
def list_memory():
    q, args = 'SELECT * FROM agent_memory WHERE 1=1', []
    for field, col in (('agent', 'agent'), ('category', 'category'), ('user', 'user_id')):
        if request.args.get(field):
            q += f' AND {col}=?'; args.append(request.args[field])
    q += ' ORDER BY last_used DESC LIMIT 200'
    return jsonify({'memories': many(q, tuple(args))})


@app.delete('/api/memory/<int:mid>')
def delete_memory(mid):
    import agent_memory as am
    if not am.forget(get_db(), mid):
        return jsonify({'error': 'Memory not found'}), 404
    return '', 204


@app.post('/api/platform/optimize')
def run_optimization_scan():
    """R9S2E7: run the autonomous optimization analysis now (also registered
    as a background job kind)."""
    import optimizer
    n = optimizer.analyze(get_db())
    log_action('optimization.scanned', 'platform', None, {'new_proposals': n})
    return jsonify({'new_proposals': n})


@app.get('/api/platform/optimizations')
def list_optimizations():
    rows = many('SELECT * FROM optimization_proposals ORDER BY id DESC LIMIT 100')
    return jsonify({'proposals': rows})


def _decide_optimization(pid, decision):
    row = one('SELECT * FROM optimization_proposals WHERE id=?', (pid,))
    if not row:
        return jsonify({'error': 'Proposal not found'}), 404
    if row['status'] != 'proposed':
        return jsonify({'error': f"Proposal already {row['status']}"}), 409
    execute("UPDATE optimization_proposals SET status=?, decided_by=?, "
            "decided_at=datetime('now') WHERE id=?",
            (decision, getattr(g, 'user_email', None) or 'admin', pid))
    import feedback_loop as fb
    fb.record(get_db(), 'optimization', pid,
              'accept' if decision == 'approved' else 'dismiss',
              category=row['kind'], user=getattr(g, 'user_email', None) or 'default')
    log_action(f'optimization.{decision}', 'optimization_proposal', pid,
               {'kind': row['kind'], 'target': row['target']})
    # NOTE (§14.3): approval records intent only — nothing is ever applied
    # to warehouse or schema state automatically.
    return jsonify({'id': pid, 'status': decision})


@app.post('/api/platform/optimizations/<int:pid>/approve')
@require_role('admin')
def approve_optimization(pid):
    return _decide_optimization(pid, 'approved')


@app.post('/api/platform/optimizations/<int:pid>/reject')
@require_role('admin')
def reject_optimization(pid):
    return _decide_optimization(pid, 'rejected')


@app.get('/api/agents/consultations')
def agent_consultations():
    """R9S2E5: first-class record of agent-to-agent consultations."""
    q, args = 'SELECT * FROM agent_consultations', []
    if request.args.get('run_id'):
        q += ' WHERE run_id=?'; args.append(request.args.get('run_id', type=int))
    q += ' ORDER BY id DESC LIMIT ?'; args.append(min(200, request.args.get('limit', 50, type=int)))
    rows = many(q, tuple(args))
    for r in rows:
        r['question'] = json.loads(r.pop('question_json') or '{}')
        r['answer'] = json.loads(r.pop('answer_json') or '{}')
    return jsonify({'consultations': rows})


@app.get('/api/meta/decisions')
def meta_decisions():
    """R9S2E4: arbitration decisions + systemic alerts."""
    return jsonify({
        'decisions': many('SELECT * FROM meta_decisions ORDER BY id DESC LIMIT 50'),
        'alerts': many("SELECT * FROM alerts WHERE type LIKE 'meta.%' ORDER BY id DESC LIMIT 20"),
    })


@app.post('/api/meta/reprioritize')
def meta_reprioritize():
    """R9S2E4: queue reprioritization sweep (user-facing work first)."""
    import meta_orchestrator as meta
    changed = meta.reprioritize(get_db())
    return jsonify({'changed': changed})


@app.post('/api/meta/override')
def meta_override():
    """R9S2E4/§18.3: any attempt to skip a human checkpoint is refused."""
    import meta_orchestrator as meta
    b = request.get_json() or {}
    body = meta.refuse_override(get_db(), b.get('session_id'), b.get('action', 'unknown'))
    return jsonify(body), 409


@app.post('/api/platform/events')
def post_platform_event():
    """R9S1E3: emit a platform event (schema/data/drift/business) — triggers
    fire synchronously; heavy work lands on the job queue."""
    import events
    b = request.get_json() or {}
    etype = b.get('event_type')
    if not etype:
        return jsonify({'error': 'event_type required'}), 400
    ev = events.emit(get_db(), etype, b.get('payload') or {})
    return jsonify(ev), 201


@app.get('/api/platform/events')
def list_platform_events():
    limit = min(200, request.args.get('limit', 50, type=int))
    q, args = 'SELECT * FROM platform_events', []
    if request.args.get('type'):
        q += ' WHERE event_type=?'; args.append(request.args['type'])
    q += ' ORDER BY id DESC LIMIT ?'; args.append(limit)
    return jsonify({'events': many(q, tuple(args))})


@app.post('/api/platform/jobs/drain')
def drain_jobs():
    """Deterministic worker step: process queued jobs now (tests + admin ops)."""
    import jobs
    n = jobs.process_pending(get_db())
    return jsonify({'processed': n})


@app.put('/api/platform/concurrency')
@require_role('admin')
def put_concurrency():
    """R9S1E2: per-workspace concurrency budget (default 4)."""
    b = request.get_json() or {}
    try:
        budget = int(b.get('budget'))
        assert 1 <= budget <= 32
    except (TypeError, ValueError, AssertionError):
        return jsonify({'error': 'budget must be an integer 1–32'}), 400
    execute("INSERT INTO platform_settings (key, value) VALUES ('concurrency_budget', ?) "
            "ON CONFLICT(key) DO UPDATE SET value=excluded.value, updated_at=datetime('now')",
            (str(budget),))
    log_action('platform.concurrency_set', 'platform', None, {'budget': budget})
    return jsonify({'budget': budget})


@app.get('/api/platform/dispatches')
def platform_dispatches():
    """R9S1E1: cost-ladder telemetry — dispatch counts per tier/task + est cost."""
    import orchestrator
    return jsonify(orchestrator.aggregate(get_db()))


@app.get('/api/platform/cache')
def platform_cache_stats():
    """R8S1E2: per-layer cache telemetry (entries, hits, misses, hit_rate)."""
    import cache_hier
    return jsonify({'mode': cache_hier.provider_mode(),
                    'layers': cache_hier.stats(get_db())})


@app.get('/api/search')
def workspace_search():
    import search as search_mod
    q = (request.args.get('q') or '').strip()
    if not q:
        return jsonify({'error': 'q (search query) is required'}), 400
    hits = search_mod.search(get_db(), q)
    out = []
    for h in hits:
        art = one('SELECT * FROM artifacts WHERE id=?', (h['artifact_id'],))
        if art:
            out.append({**art, 'score': h['score']})
    return jsonify(out)


@app.get('/api/platform/jobs')
def platform_jobs():
    kind, status = request.args.get('kind'), request.args.get('status')
    clauses, params = [], []
    if kind:
        clauses.append('kind=?'); params.append(kind)
    if status:
        clauses.append('status=?'); params.append(status)
    where = ('WHERE ' + ' AND '.join(clauses)) if clauses else ''
    rows = many(f'SELECT * FROM jobs {where} ORDER BY id DESC LIMIT 200', tuple(params))
    for r in rows:
        r['payload'] = json.loads(r.pop('payload_json') or '{}')
    return jsonify(rows)


@app.get('/api/platform/logs')
def platform_logs():
    try:
        limit = min(int(request.args.get('limit', 50)), 500)
    except ValueError:
        return jsonify({'error': 'limit must be an integer'}), 400
    return jsonify(many('SELECT * FROM service_logs ORDER BY id DESC LIMIT ?', (limit,)))


@app.get('/api/platform/outbox')
def platform_outbox():
    return jsonify(many('SELECT id, recipient, subject, status, created_at '
                        'FROM email_outbox ORDER BY id DESC LIMIT 100'))


@app.get('/api/platform/metrics')
def platform_metrics():
    rows = many('SELECT path, duration_ms FROM service_logs')
    durs = sorted(r['duration_ms'] for r in rows if r['duration_ms'] is not None)

    def pct(q):
        if not durs:
            return 0.0
        return round(durs[min(len(durs) - 1, int(len(durs) * q))], 2)

    by_path = {}
    for r in rows:
        by_path.setdefault(r['path'], []).append(r['duration_ms'] or 0)
    return jsonify({
        'requests': len(rows),
        'latency_ms': {'p50': pct(0.50), 'p95': pct(0.95)},
        'by_path': {p: {'count': len(v),
                        'avg_ms': round(sum(v) / len(v), 2)} for p, v in by_path.items()},
    })


@app.get('/api/workspace/status')
def workspace_status():
    counts = {
        'connections': one('SELECT COUNT(*) c FROM connections')['c'],
        'governance_runs': one("SELECT COUNT(*) c FROM governance_runs WHERE status='done'")['c'],
        'pending_reviews': one(
            "SELECT COUNT(*) c FROM semantic_definitions "
            "WHERE status='pending' AND confidence < 0.70")['c'],
        'sessions': one('SELECT COUNT(*) c FROM sessions')['c'],
        'pipeline_runs': one("SELECT COUNT(*) c FROM pipeline_runs WHERE status='done'")['c'],
        'artifacts': one('SELECT COUNT(*) c FROM artifacts')['c'],
    }
    if counts['connections'] == 0:
        next_step = 'connect'
    elif counts['governance_runs'] == 0:
        next_step = 'scan'
    elif counts['pending_reviews'] > 0:
        next_step = 'review'
    elif counts['artifacts'] == 0:
        next_step = 'analyze'
    else:
        next_step = 'explore'
    return jsonify({**counts, 'next_step': next_step,
                    'onboarding_complete': counts['artifacts'] > 0})


# ─────────────────────────────────────────────────────────
# Routes — Auth (R1S1E1; local fallback provider)
# ─────────────────────────────────────────────────────────
@app.post('/api/auth/register')
@limiter.limit('20/minute')
def auth_register():
    import authn
    b = request.get_json() or {}
    email = (b.get('email') or '').strip().lower()
    password = b.get('password') or ''
    role = (b.get('role') or 'analyst').lower()
    err = authn.validate_registration(email, password, role)
    if err:
        return jsonify({'error': err}), 400
    if one('SELECT id FROM users WHERE email=?', (email,)):
        return jsonify({'error': 'An account with this email already exists'}), 409
    uid = execute('INSERT INTO users (email, password_hash, role) VALUES (?,?,?)',
                  (email, authn.hash_password(password), role))
    log_action('auth.registered', 'user', uid, {'email': email, 'role': role,
                                                'provider': authn.provider_mode()})
    return jsonify({'id': uid, 'email': email, 'role': role}), 201


@app.post('/api/auth/login')
@limiter.limit('30/minute')
def auth_login():
    import authn
    b = request.get_json() or {}
    email = (b.get('email') or '').strip().lower()
    row = one('SELECT * FROM users WHERE email=?', (email,))
    if not row or not authn.verify_password(b.get('password') or '', row['password_hash']):
        log_action('auth.login_failed', 'user', row['id'] if row else None, {'email': email})
        return jsonify({'error': 'Invalid email or password'}), 401
    token = authn.mint_token(get_db(), row['id'])
    log_action('auth.login', 'user', row['id'], {'email': email})
    return jsonify({'token': token,
                    'user': {'id': row['id'], 'email': row['email'], 'role': row['role']}})


@app.get('/api/auth/me')
def auth_me():
    user = getattr(g, 'auth_user', None)
    if not user:
        return jsonify({'error': 'A valid bearer token is required'}), 401
    return jsonify({'id': user['id'], 'email': user['email'], 'role': user['role']})


# ─────────────────────────────────────────────────────────
# Routes — Resource ACLs (R1S1E2)
# ─────────────────────────────────────────────────────────
_ACL_RESOURCE_TABLES = {'artifact': 'artifacts', 'session': 'sessions',
                        'connection': 'connections'}


@app.put('/api/acl/<rtype>/<rid>')
@require_role('admin')
def put_acl(rtype, rid):
    if rtype == 'explore':
        row = _latest_schema_row('default')
        cubes = {c.get('name') for c in json.loads(row['schema_json'])['cubes']} if row else set()
        if rid not in cubes:
            return jsonify({'error': 'Explore not found'}), 404
    else:
        table = _ACL_RESOURCE_TABLES.get(rtype)
        if not table:
            return jsonify({'error': f'Unknown resource type {rtype!r}'}), 400
        if not one(f'SELECT id FROM {table} WHERE id=?', (rid,)):
            return jsonify({'error': 'Resource not found'}), 404
    entries = (request.get_json() or {}).get('entries')
    if not isinstance(entries, list):
        return jsonify({'error': 'entries list required'}), 400
    for e in entries:
        if not isinstance(e, dict) or not e.get('principal') or e.get('role') not in ACL_ROLES:
            return jsonify({'error': f'each entry needs principal and role in {ACL_ROLES}'}), 400
    db = get_db()
    db.execute('DELETE FROM resource_acls WHERE resource_type=? AND resource_id=?',
               (rtype, str(rid)))
    for e in entries:
        db.execute('INSERT INTO resource_acls (resource_type, resource_id, principal, role) '
                   'VALUES (?,?,?,?)', (rtype, str(rid), e['principal'], e['role']))
    db.commit()
    log_action('acl.updated', rtype, rid,
               {'entries': [{'principal': e['principal'], 'role': e['role']} for e in entries]})
    return jsonify(acl_entries(rtype, rid))


@app.get('/api/acl/<rtype>/<rid>')
def get_acl(rtype, rid):
    if rtype not in _ACL_RESOURCE_TABLES and rtype != 'explore':
        return jsonify({'error': f'Unknown resource type {rtype!r}'}), 400
    return jsonify(acl_entries(rtype, rid))


# ─────────────────────────────────────────────────────────
# Routes — Connections
# ─────────────────────────────────────────────────────────
def _mask_connection(row):
    if not row:
        return row
    conn_type = row.get('type', 'snowflake')
    if conn_type == 'bigquery':
        row['username'] = None
    else:
        row['username'] = decrypt(row['username'])
    row['password'] = MASKED_PASSWORD if row.get('password') else None
    return row


@app.get('/api/connections')
def list_connections():
    return jsonify([_mask_connection(r) for r in many('SELECT * FROM connections ORDER BY created_at DESC')])

@app.get('/api/connections/<int:id>')
def get_connection(id):
    row = one('SELECT * FROM connections WHERE id=?', (id,))
    return jsonify(_mask_connection(row)) if row else (jsonify({'error': 'Not found'}), 404)

@app.post('/api/connections/test')
@limiter.limit('10/minute')
def test_connection():
    b = request.get_json() or {}
    conn_type = b.get('type', 'snowflake')
    try:
        simulated = False
        if conn_type == 'postgres':
            _test_postgres(b)
        elif conn_type == 'bigquery':
            _test_bigquery(b)
        elif conn_type in ('snowflake', 'mysql', 'redshift', 'databricks', 'duckdb',
                           'gsheet', 'rest_api', 'webhook', 'file'):
            errors = validate_connection_payload(b)
            if errors:
                return jsonify({'ok': False, 'error': 'Validation failed',
                                'fields': errors}), 422
            simulated = True  # offline contract — no live warehouse in this stack
        else:
            return jsonify({'ok': False, 'error': f'Unknown connector type: {conn_type}'}), 400
        return jsonify({'ok': True, 'message': 'Connection successful', 'simulated': simulated})
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)}), 422


def _map_connection_fields(b):
    conn_type = b.get('type', 'snowflake')
    if conn_type == 'postgres':
        return {
            'name': b.get('name', 'My PostgreSQL'),
            'type': 'postgres',
            'account': b.get('host'),
            'username': encrypt(b.get('username')),
            'password': encrypt(b.get('password')),
            'warehouse': b.get('port', '5432'),
            'database_name': b.get('database_name'),
            'schema_name': b.get('schema_name', 'public'),
            'owner_email': b.get('owner_email'),
        }
    if conn_type in ('mysql', 'redshift'):
        return {
            'name': b.get('name', f'My {conn_type.title()}'),
            'type': conn_type,
            'account': b.get('host'),
            'username': encrypt(b.get('username')),
            'password': encrypt(b.get('password')),
            'warehouse': b.get('port'),
            'database_name': b.get('database_name'),
            'schema_name': b.get('schema_name'),
            'owner_email': b.get('owner_email'),
        }
    if conn_type == 'databricks':
        return {
            'name': b.get('name', 'My Databricks'),
            'type': 'databricks',
            'account': b.get('host'),
            'username': None,
            'password': encrypt(b.get('access_token')),
            'warehouse': b.get('http_path'),
            'database_name': b.get('catalog'),
            'schema_name': b.get('schema_name'),
            'owner_email': b.get('owner_email'),
        }
    if conn_type == 'duckdb':
        return {
            'name': b.get('name', 'My DuckDB'),
            'type': 'duckdb',
            'account': b.get('database_path'),
            'username': None, 'password': None, 'warehouse': None,
            'database_name': b.get('database_path'), 'schema_name': None,
            'owner_email': b.get('owner_email'),
        }
    if conn_type == 'rest_api':
        return {
            'name': b.get('name', 'My API Source'),
            'type': 'rest_api',
            'account': b.get('endpoint_url'),
            'username': None,
            'password': encrypt(b.get('auth_header')) if b.get('auth_header') else None,
            'warehouse': None, 'database_name': None, 'schema_name': None,
            'owner_email': b.get('owner_email'),
        }
    if conn_type == 'webhook':
        return {
            'name': b.get('name', 'My Webhook Source'),
            'type': 'webhook',
            'account': None,  # set to token hash post-insert
            'username': None, 'password': None, 'warehouse': None,
            'database_name': None, 'schema_name': None,
            'owner_email': b.get('owner_email'),
        }
    if conn_type == 'gsheet':
        return {
            'name': b.get('name', 'My Google Sheet'),
            'type': 'gsheet',
            'account': b.get('sheet_url'),
            'username': None, 'password': None, 'warehouse': None,
            'database_name': None, 'schema_name': None,
            'owner_email': b.get('owner_email'),
        }
    if conn_type == 'bigquery':
        return {
            'name': b.get('name', 'My BigQuery'),
            'type': 'bigquery',
            'account': b.get('project_id'),
            'username': None,
            'password': encrypt(b.get('credentials_json')),
            'warehouse': None,
            'database_name': b.get('dataset'),
            'schema_name': None,
            'owner_email': b.get('owner_email'),
        }
    return {
        'name': b.get('name', 'My Snowflake'),
        'type': 'snowflake',
        'account': b.get('account'),
        'username': encrypt(b.get('username')),
        'password': encrypt(b.get('password')),
        'warehouse': b.get('warehouse'),
        'database_name': b.get('database_name'),
        'schema_name': b.get('schema_name'),
        'owner_email': b.get('owner_email'),
    }


@app.post('/api/connections')
@limiter.limit('10/minute')
@require_role('admin')
def create_connection():
    b = request.get_json() or {}
    errors = validate_connection_payload(b)
    if errors:
        log_action('connection.rejected', 'connection', None,
                   {'type': b.get('type', 'snowflake'), 'fields': sorted(errors), 'outcome': 'failure'})
        return jsonify({'error': 'Validation failed', 'fields': errors}), 400
    f = _map_connection_fields(b)
    lid = execute(
        'INSERT INTO connections (name,type,account,username,password,warehouse,database_name,schema_name,owner_email) '
        'VALUES (?,?,?,?,?,?,?,?,?)',
        (f['name'], f['type'], f['account'], f['username'], f['password'],
         f['warehouse'], f['database_name'], f['schema_name'], f['owner_email']),
    )
    log_action('connection.created', 'connection', lid, {'name': f['name'], 'type': f['type']})
    resp = _mask_connection(one('SELECT * FROM connections WHERE id=?', (lid,)))
    if f['type'] == 'rest_api':
        minutes = int(b.get('poll_interval_minutes') or 60)
        dbx = get_db()
        dbx.execute("UPDATE connections SET poll_interval_minutes=?, "
                    "next_poll_at=datetime('now', ?) WHERE id=?",
                    (minutes, f'+{minutes} minutes', lid))
        dbx.commit()
    if f['type'] == 'webhook':
        # R2S2E1: capability token — raw shown once, only its hash at rest
        import hashlib as _h
        import secrets as _s
        raw = _s.token_urlsafe(24)
        execute2 = get_db()
        execute2.execute('UPDATE connections SET account=? WHERE id=?',
                         (_h.sha256(raw.encode()).hexdigest(), lid))
        execute2.commit()
        resp['webhook_token'] = raw
        resp['webhook_url'] = f'/api/ingest/webhook/{raw}'
        resp['account'] = None
    return jsonify(resp), 201

@app.delete('/api/connections/<int:id>')
@require_role('admin')
def delete_connection(id):
    execute('DELETE FROM connections WHERE id=?', (id,))
    log_action('connection.deleted', 'connection', id)
    return '', 204


# ─────────────────────────────────────────────────────────
# REST API source polling (R2S2E2)
# ─────────────────────────────────────────────────────────
def _fixture_api_rows(connection_id):
    """Deterministic offline stand-in for the remote API response."""
    rand = seeded_rng(connection_id + 100)
    return [{'order_id': i + 1, 'amount': round(20 + rand() * 300, 2),
             'status': ['new', 'shipped'][int(rand() * 2) % 2]} for i in range(5)]


def _poll_connection(conn, connection_row):
    """Fetch (or simulate) the API and append rows; returns (count, mode)."""
    rows = _fixture_api_rows(connection_row['id'])   # sandbox: no outbound network
    mode = 'offline_fixture'
    for r in rows:
        conn.execute('INSERT INTO webhook_events (connection_id, payload_json) VALUES (?,?)',
                     (connection_row['id'], json.dumps(r)))
    minutes = connection_row['poll_interval_minutes'] or 60
    conn.execute("UPDATE connections SET next_poll_at=datetime('now', ?) WHERE id=?",
                 (f'+{minutes} minutes', connection_row['id']))
    conn.execute('INSERT INTO poll_runs (connection_id, status, rows_ingested, mode) '
                 "VALUES (?, 'done', ?, ?)", (connection_row['id'], len(rows), mode))
    conn.commit()
    return len(rows), mode


def _run_due_api_polls(conn):
    due = conn.execute(
        "SELECT * FROM connections WHERE type='rest_api' "
        "AND next_poll_at IS NOT NULL AND next_poll_at <= datetime('now')").fetchall()
    for row in due:
        try:
            _poll_connection(conn, dict(row))
        except Exception:
            pass
    return len(due)


@app.post('/api/connections/<int:id>/poll')
@limiter.limit('20/minute')
@require_role('admin', 'analyst')
def poll_rest_source(id):
    row = one('SELECT * FROM connections WHERE id=?', (id,))
    if not row:
        return jsonify({'error': 'Not found'}), 404
    if row['type'] != 'rest_api':
        return jsonify({'error': 'Only rest_api connections can be polled'}), 409
    count, mode = _poll_connection(get_db(), row)
    log_action('source.polled', 'connection', id, {'rows': count, 'mode': mode})
    return jsonify({'connection_id': id, 'rows_ingested': count, 'mode': mode}), 201


# ─────────────────────────────────────────────────────────
# Routes — Webhook ingest (R2S2E1)
# ─────────────────────────────────────────────────────────
@app.post('/api/ingest/webhook/<token>')
@limiter.limit('120/minute')
def webhook_ingest(token):
    import hashlib as _h
    row = one("SELECT id FROM connections WHERE type='webhook' AND account=?",
              (_h.sha256(token.encode()).hexdigest(),))
    if not row:
        return jsonify({'error': 'Unknown webhook token'}), 404
    payload = request.get_json(silent=True)
    if payload is None:
        return jsonify({'error': 'JSON body required'}), 400
    lid = execute('INSERT INTO webhook_events (connection_id, payload_json) VALUES (?,?)',
                  (row['id'], json.dumps(payload)))
    import events
    events.emit(get_db(), 'data_arrived', {'connection_id': row['id'], 'webhook_event_id': lid})
    return jsonify({'id': lid, 'connection_id': row['id']}), 201


@app.get('/api/connections/<int:id>/events')
def list_webhook_events(id):
    rows = many('SELECT * FROM webhook_events WHERE connection_id=? '
                'ORDER BY id DESC LIMIT 200', (id,))
    for r in rows:
        r['payload'] = json.loads(r.pop('payload_json'))
    return jsonify(rows)


# ─────────────────────────────────────────────────────────
# Routes — File upload ingestion (R2S1E1)
# ─────────────────────────────────────────────────────────
def _coerce_cell(v):
    if v is None or v == '':
        return None
    try:
        return int(v)
    except (ValueError, TypeError):
        pass
    try:
        return float(v)
    except (ValueError, TypeError):
        return v


def _rows_from_csv(data: bytes):
    import csv
    import io as _io
    text = data.decode('utf-8-sig', errors='replace')
    reader = csv.reader(_io.StringIO(text))
    rows = [r for r in reader if r]
    if len(rows) < 2:
        raise ValueError('CSV needs a header row and at least one data row')
    header = [c.strip() or f'col_{i}' for i, c in enumerate(rows[0])]
    return header, [tuple(_coerce_cell(c) for c in r) for r in rows[1:]]


def _rows_from_xlsx(data: bytes):
    import io as _io
    try:
        import openpyxl
    except ImportError:
        raise NotImplementedError('XLSX support requires openpyxl')
    wb = openpyxl.load_workbook(_io.BytesIO(data), read_only=True)
    ws = wb.active
    rows = [[c for c in row] for row in ws.iter_rows(values_only=True) if any(row)]
    if len(rows) < 2:
        raise ValueError('Sheet needs a header row and at least one data row')
    header = [str(c or f'col_{i}') for i, c in enumerate(rows[0])]
    return header, [tuple(r) for r in rows[1:]]


@app.post('/api/uploads')
@limiter.limit('10/minute')
@require_role('admin', 'analyst')
def upload_file():
    import re as _re
    import profiler
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'error': 'multipart field "file" is required'}), 400
    name = f.filename.lower()
    data = f.read()
    try:
        if name.endswith('.csv'):
            header, rows = _rows_from_csv(data)
        elif name.endswith('.xlsx'):
            header, rows = _rows_from_xlsx(data)
        else:
            return jsonify({'error': f'Unsupported file type for {f.filename!r} — '
                                     'CSV and XLSX are accepted'}), 415
    except NotImplementedError as e:
        return jsonify({'error': str(e)}), 415
    except Exception as e:
        return jsonify({'error': f'Could not parse file: {e}'}), 400

    slug = _re.sub(r'[^a-z0-9]+', '_', name.rsplit('.', 1)[0]).strip('_') or 'upload'
    table = f'src_upload_{slug}'
    db = get_db()
    col_sql = ', '.join(f'"{c}"' for c in header)
    db.execute(f'DROP TABLE IF EXISTS "{table}"')
    db.execute(f'CREATE TABLE "{table}" ({col_sql})')
    placeholders = ','.join(['?'] * len(header))
    db.executemany(f'INSERT INTO "{table}" VALUES ({placeholders})',
                   [r[:len(header)] + (None,) * (len(header) - len(r)) for r in rows])
    db.commit()

    prof = profiler.profile_table(db, table)
    cid = execute('INSERT INTO connections (name, type, database_name, owner_email) '
                  "VALUES (?, 'file', ?, ?)",
                  (f.filename, table, request.headers.get('X-User-Email')))
    execute("INSERT INTO ingestion_profiles (connection_id, table_name, status, row_count, "
            "sampled_rows, columns_json, completed_at) VALUES (?,?,?,?,?,?,datetime('now'))",
            (cid, table, 'done', prof['row_count'], prof['sampled_rows'],
             json.dumps(prof['columns'])))
    log_action('upload.ingested', 'connection', cid,
               {'file': f.filename, 'table': table, 'rows': prof['row_count']})
    return jsonify({'connection_id': cid, 'table': table,
                    'row_count': prof['row_count'], 'profile': prof}), 201


# ─────────────────────────────────────────────────────────
# Routes — Ingestion profiling (Sprint 1 / F-005)
# ─────────────────────────────────────────────────────────
def _demo_source_rows():
    """Deterministic synthetic source sample for simulated connectors."""
    rand = seeded_rng(7)
    statuses = ['new', 'shipped', 'returned', 'cancelled']
    rows = []
    for i in range(2000):
        d = date(2024, 1, 1) + timedelta(days=i % 120)
        rows.append((
            i + 1,                                    # order_id
            round(20 + rand() * 480, 2),              # amount
            statuses[int(rand() * len(statuses)) % len(statuses)],  # status
            d.isoformat(),                            # created_at
            1 if rand() > 0.85 else 0,                # is_gift
            None if rand() < 0.08 else f'note for order {i + 1} ' + 'x' * 45,  # note
        ))
    return ['order_id', 'amount', 'status', 'created_at', 'is_gift', 'note'], rows


@app.post('/api/connections/<int:id>/profile')
@limiter.limit('10/minute')
def run_ingestion_profile(id):
    import profiler
    conn_row = one('SELECT * FROM connections WHERE id=?', (id,))
    if not conn_row:
        return jsonify({'error': 'Not found'}), 404

    table_name = (request.get_json(silent=True) or {}).get('table', 'src_orders')
    columns, rows = _demo_source_rows()
    prof = profiler.profile_rows(table_name, columns, rows[:profiler.SAMPLE_LIMIT],
                                 row_count=len(rows))
    # R2S1E2: sheets are low-trust — flag noisy columns automatically
    trust, warnings = 'normal', []
    if conn_row['type'] == 'gsheet':
        trust = 'low'
        for c in prof['columns']:
            if (c.get('null_pct') or 0) > 5:
                warnings.append(f"column '{c['name']}' has {c['null_pct']}% nulls — "
                                f'low-trust source, verify before ML use')
    lid = execute(
        "INSERT INTO ingestion_profiles (connection_id,table_name,status,row_count,sampled_rows,"
        "columns_json,completed_at) VALUES (?,?,?,?,?,?,datetime('now'))",
        (id, prof['table'], 'done', prof['row_count'], prof['sampled_rows'],
         json.dumps(prof['columns'])),
    )
    log_action('profile.completed', 'ingestion_profile', lid,
               {'connection_id': id, 'table': prof['table'],
                'row_count': prof['row_count'], 'sampled_rows': prof['sampled_rows']})
    saved = one('SELECT * FROM ingestion_profiles WHERE id=?', (lid,))
    saved['columns'] = json.loads(saved.pop('columns_json'))
    saved['trust'] = trust
    saved['warnings'] = warnings
    return jsonify(saved), 201


@app.get('/api/connections/<int:id>/profiles')
def list_ingestion_profiles(id):
    rows = many('SELECT * FROM ingestion_profiles WHERE connection_id=? ORDER BY created_at DESC', (id,))
    for r in rows:
        r['columns'] = json.loads(r.pop('columns_json') or '[]')
    return jsonify(rows)


# ─────────────────────────────────────────────────────────
# Routes — Governance manifests (Sprint 2 / F-007, F-008)
# ─────────────────────────────────────────────────────────
def _manifest_row(connection_id, version=None):
    if version:
        return one('SELECT * FROM governance_manifests WHERE connection_id=? AND version=?',
                   (connection_id, version))
    return one('SELECT * FROM governance_manifests WHERE connection_id=? ORDER BY id DESC LIMIT 1',
               (connection_id,))


@app.get('/api/integrations/<int:id>/health_history')
def health_history_route(id):
    table = request.args.get('table')
    if table:
        return jsonify(many('SELECT * FROM health_history WHERE connection_id=? AND '
                            'table_name=? ORDER BY id', (id, table)))
    return jsonify(many('SELECT * FROM health_history WHERE connection_id=? ORDER BY id', (id,)))


@app.put('/api/governance/thresholds')
@require_role('admin', 'analyst')
def put_governance_threshold():
    b = request.get_json() or {}
    cid = b.get('connectionId')
    if not cid:
        return jsonify({'error': 'connectionId required'}), 400
    if not isinstance(b.get('min_health'), int) or not (0 <= b['min_health'] <= 100):
        return jsonify({'error': 'min_health must be an integer 0–100'}), 400
    db = get_db()
    db.execute('INSERT INTO governance_thresholds (connection_id, min_health) VALUES (?,?) '
               'ON CONFLICT(connection_id) DO UPDATE SET min_health=excluded.min_health, '
               "updated_at=datetime('now')", (cid, b['min_health']))
    db.commit()
    log_action('governance.threshold_set', 'connection', cid, {'min_health': b['min_health']})
    return jsonify({'connection_id': cid, 'min_health': b['min_health']})


@app.put('/api/contracts')
@require_role('admin', 'analyst')
def put_data_contract():
    b = request.get_json() or {}
    cid, table = b.get('connectionId'), b.get('table')
    if not cid or not table:
        return jsonify({'error': 'connectionId and table required'}), 400
    cols = b.get('required_columns', [])
    if not isinstance(cols, list):
        return jsonify({'error': 'required_columns must be a list'}), 400
    db = get_db()
    db.execute('INSERT INTO data_contracts (connection_id, table_name, required_columns_json, '
               'min_rows, max_age_hours) VALUES (?,?,?,?,?) '
               'ON CONFLICT(connection_id, table_name) DO UPDATE SET '
               'required_columns_json=excluded.required_columns_json, '
               'min_rows=excluded.min_rows, max_age_hours=excluded.max_age_hours, '
               "updated_at=datetime('now')",
               (cid, table, json.dumps(cols), b.get('min_rows'), b.get('max_age_hours')))
    db.commit()
    log_action('contract.set', 'connection', cid, {'table': table, 'required_columns': cols})
    return jsonify({'connection_id': cid, 'table': table, 'required_columns': cols})


@app.get('/api/contracts')
def list_data_contracts():
    cid = request.args.get('connection_id')
    rows = (many('SELECT * FROM data_contracts WHERE connection_id=? ORDER BY id', (cid,))
            if cid else many('SELECT * FROM data_contracts ORDER BY id'))
    for r in rows:
        r['required_columns'] = json.loads(r.pop('required_columns_json') or '[]')
    return jsonify(rows)


@app.delete('/api/contracts/<int:id>')
@require_role('admin', 'analyst')
def delete_data_contract(id):
    execute('DELETE FROM data_contracts WHERE id=?', (id,))
    return '', 204


@app.put('/api/tables/sla')
@require_role('admin', 'analyst')
def put_freshness_sla():
    b = request.get_json() or {}
    cid, table = b.get('connectionId'), b.get('table')
    if not cid or not table:
        return jsonify({'error': 'connectionId and table required'}), 400
    hours = b.get('max_age_hours')
    if not isinstance(hours, (int, float)) or hours <= 0:
        return jsonify({'error': 'max_age_hours must be a positive number'}), 400
    db = get_db()
    db.execute('INSERT INTO freshness_slas (connection_id, table_name, max_age_hours) '
               'VALUES (?,?,?) ON CONFLICT(connection_id, table_name) DO UPDATE SET '
               "max_age_hours=excluded.max_age_hours, updated_at=datetime('now')",
               (cid, table, hours))
    db.commit()
    log_action('freshness.sla_set', 'connection', cid, {'table': table, 'max_age_hours': hours})
    return jsonify({'connection_id': cid, 'table': table, 'max_age_hours': hours})


@app.get('/api/tables/sla')
def list_freshness_slas():
    cid = request.args.get('connection_id')
    if cid:
        return jsonify(many('SELECT * FROM freshness_slas WHERE connection_id=? ORDER BY id',
                            (cid,)))
    return jsonify(many('SELECT * FROM freshness_slas ORDER BY id'))


@app.get('/api/alerts')
def list_alerts():
    atype = request.args.get('type')
    cid = request.args.get('connection_id')
    clauses, params = [], []
    if atype:
        clauses.append('type=?'); params.append(atype)
    if cid:
        clauses.append('connection_id=?'); params.append(cid)
    where = ('WHERE ' + ' AND '.join(clauses)) if clauses else ''
    rows = many(f'SELECT * FROM alerts {where} ORDER BY id DESC LIMIT 200', tuple(params))
    for r in rows:
        r['detail'] = json.loads(r.pop('detail_json') or '{}')
    return jsonify(rows)


@app.get('/api/lineage/<int:id>')
def lineage_dag(id):
    row = _manifest_row(id)
    if not row:
        return jsonify({'error': 'No manifest for this connection'}), 404
    m = json.loads(row['manifest_json'])
    nodes = [{'id': t['name'], 'kind': 'table', 'label': t['name'],
              'health_score': t.get('health_score'), 'freshness': t.get('freshness'),
              'row_count': t.get('row_count'),
              'dq_gate_status': t.get('dq_gate_status')} for t in m.get('tables', [])]
    edges = [{'from': e['from'], 'to': e['to'], 'on': e.get('on'), 'kind': 'join'}
             for e in m.get('lineage_edges', [])]

    # R32S1E5: the connection is a source node feeding root tables (no
    # incoming join edge), and accepted metric definitions join the graph.
    conn_row = one('SELECT * FROM connections WHERE id=?', (id,))
    if conn_row:
        sid = f'source:{id}'
        nodes.append({'id': sid, 'kind': 'source',
                      'label': conn_row.get('account') or conn_row.get('type') or 'source'})
        targets = {e['to'] for e in edges}
        for t in m.get('tables', []):
            if t['name'] not in targets:
                edges.append({'from': sid, 'to': t['name'], 'kind': 'loads'})
    fact_anchor_m = next((t['name'] for t in m.get('tables', [])
                          if t['name'].startswith('fact')),
                         (m.get('tables') or [{}])[0].get('name'))
    for d in many("SELECT sd.* FROM semantic_definitions sd JOIN governance_runs gr "
                  "ON sd.run_id = gr.id WHERE gr.connection_id=? AND sd.status='accepted' "
                  "AND lower(sd.type)='metric' ORDER BY sd.id", (id,)):
        nid = f"metric:{d['id']}"
        nodes.append({'id': nid, 'kind': 'metric', 'label': d['name'],
                      'explore': d['explore']})
        if fact_anchor_m:
            edges.append({'from': fact_anchor_m, 'to': nid, 'kind': 'defines'})

    # downstream: gold tables + artifacts reachable through sessions of this connection
    sessions = many('SELECT id FROM sessions WHERE connection_id=?', (id,))
    fact_anchor = next((t['name'] for t in m.get('tables', [])
                        if t['name'].startswith('fact')), None)
    for s in sessions:
        for g in many('SELECT * FROM gold_tables WHERE session_id=?', (s['id'],)):
            nid = f"gold:{g['id']}"
            nodes.append({'id': nid, 'kind': 'gold_table', 'label': g['table_name'],
                          'status': g['status']})
            if fact_anchor:
                edges.append({'from': fact_anchor, 'to': nid, 'kind': 'materializes'})
        for mr in many('SELECT * FROM model_registry WHERE session_id=?', (s['id'],)):
            nid = f"model:{mr['id']}"
            nodes.append({'id': nid, 'kind': 'model',
                          'label': f"{mr['model_id']} v{mr['version']}",
                          'status': mr['status']})
            if fact_anchor:
                edges.append({'from': fact_anchor, 'to': nid, 'kind': 'trains'})
        for r_ in many('SELECT id FROM pipeline_runs WHERE session_id=?', (s['id'],)):
            for a in many('SELECT * FROM artifacts WHERE pipeline_run_id=?', (r_['id'],)):
                nid = f"artifact:{a['id']}"
                nodes.append({'id': nid, 'kind': 'artifact', 'label': a['title'],
                              'artifact_id': a['id']})
                if fact_anchor:
                    edges.append({'from': fact_anchor, 'to': nid, 'kind': 'feeds'})
    return jsonify({'connection_id': id, 'nodes': nodes, 'edges': edges,
                    'manifest_version': m.get('manifest_version')})


@app.get('/api/integrations/<int:id>/drift')
def list_drift_alerts(id):
    rows = many("SELECT * FROM alerts WHERE type='drift' AND connection_id=? "
                'ORDER BY id DESC', (id,))
    for r in rows:
        r['detail'] = json.loads(r.pop('detail_json') or '{}')
    return jsonify(rows)


@app.get('/api/integrations/<int:id>/manifest')
def get_governance_manifest(id):
    row = _manifest_row(id, request.args.get('version'))
    if not row:
        return jsonify({'error': 'No manifest found'}), 404
    m = json.loads(row['manifest_json'])
    # optional table pagination (Sprint 13 / F-054)
    if request.args.get('table_page') or request.args.get('table_per_page'):
        try:
            page = max(1, int(request.args.get('table_page', 1)))
            per_page = max(1, min(100, int(request.args.get('table_per_page', 20))))
        except ValueError:
            return jsonify({'error': 'table_page/table_per_page must be integers'}), 400
        tables = m.get('tables', [])
        m['table_total'] = len(tables)
        m['table_page'] = page
        m['table_per_page'] = per_page
        m['tables'] = tables[(page - 1) * per_page: page * per_page]
    return jsonify(m)


@app.get('/api/integrations/<int:id>/manifest/versions')
def list_manifest_versions(id):
    rows = many('SELECT id, version, run_id, created_at FROM governance_manifests '
                'WHERE connection_id=? ORDER BY id DESC', (id,))
    if not request.args.get('diffs'):
        return jsonify(rows)
    # R32S1E6: annotate with status + structural changes vs previous version.
    payloads = {r['id']: json.loads(one(
        'SELECT manifest_json FROM governance_manifests WHERE id=?',
        (r['id'],))['manifest_json']) for r in rows}
    for i, r in enumerate(rows):
        pending = one('SELECT COUNT(*) AS n FROM semantic_definitions '
                      "WHERE run_id=? AND status='pending' AND confidence < ?",
                      (r['run_id'], REVIEW_CONFIDENCE_THRESHOLD))['n']
        if i > 0:
            r['status'] = 'SUPERSEDED'
        elif pending:
            r['status'] = 'REVIEW REQUIRED'
        else:
            r['status'] = 'ACTIVE'
        cur = {t['name']: t for t in payloads[r['id']].get('tables', [])}
        prev_row = rows[i + 1] if i + 1 < len(rows) else None
        prev = ({t['name']: t for t in payloads[prev_row['id']].get('tables', [])}
                if prev_row else {})
        added = sorted(set(cur) - set(prev)) if prev_row else []
        removed = sorted(set(prev) - set(cur)) if prev_row else []
        modified = []
        if prev_row:
            for name in sorted(set(cur) & set(prev)):
                reasons = []
                if cur[name].get('health_score') != prev[name].get('health_score'):
                    reasons.append(f"health {prev[name].get('health_score')} → "
                                   f"{cur[name].get('health_score')}")
                if len(cur[name].get('columns') or []) != len(prev[name].get('columns') or []):
                    reasons.append('column set changed')
                if cur[name].get('dq_gate_status') != prev[name].get('dq_gate_status'):
                    reasons.append(f"gate {prev[name].get('dq_gate_status')} → "
                                   f"{cur[name].get('dq_gate_status')}")
                if reasons:
                    modified.append({'table': name, 'reason': ' · '.join(reasons)})
        r['changes'] = {'added': added, 'modified': modified, 'removed': removed}
        r['pending_reviews'] = pending
    return jsonify(rows)


@app.post('/api/integrations/<int:id>/manifest/approve_pii')
@require_role('admin')
def approve_manifest_pii(id):
    import manifest as manifest_mod
    b = request.get_json() or {}
    cols = b.get('columns') or []
    if not cols:
        return jsonify({'error': 'columns required — [{table, column}, ...]'}), 400
    row = _manifest_row(id)
    if not row:
        return jsonify({'error': 'No manifest found'}), 404

    m = json.loads(row['manifest_json'])
    targets = {(c.get('table'), c.get('column')) for c in cols}
    changed = []
    now = datetime.utcnow().isoformat() + 'Z'
    for t in m['tables']:
        for c in t['columns']:
            if (t['name'], c['name']) in targets and c.get('pii_flags'):
                c['allow_ml_use'] = True
                c['pii_approval'] = {
                    'approved_by': request.headers.get('X-User-Email', 'admin'),
                    'justification': b.get('justification'),
                    'approved_at': now,
                    'expires_at': b.get('expires_at'),
                }
                changed.append(f"{t['name']}.{c['name']}")
    if not changed:
        return jsonify({'error': 'No matching PII-flagged columns found'}), 400

    m['derived_from'] = m['manifest_version']
    saved = manifest_mod.save_manifest(get_db(), id, m)
    log_action('manifest.pii_approved', 'manifest', id,
               {'columns': changed, 'justification': b.get('justification'),
                'version': saved['manifest_version']})
    return jsonify(saved)


@app.post('/api/integrations/<int:id>/manifest/rollback')
@require_role('admin')
def rollback_manifest(id):
    import manifest as manifest_mod
    version = (request.get_json() or {}).get('version')
    if not version:
        return jsonify({'error': 'version required'}), 400
    row = _manifest_row(id, version)
    if not row:
        return jsonify({'error': 'Version not found'}), 404
    m = json.loads(row['manifest_json'])
    m['rolled_back_from'] = version
    saved = manifest_mod.save_manifest(get_db(), id, m)
    log_action('manifest.rolled_back', 'manifest', id,
               {'to_version': version, 'new_version': saved['manifest_version']})
    return jsonify(saved)


# ─────────────────────────────────────────────────────────
# Routes — Governance
# ─────────────────────────────────────────────────────────
@app.post('/api/governance/run')
@limiter.limit('5/minute')
def start_governance():
    b   = request.get_json() or {}
    cid = b.get('connectionId')
    if not cid:
        return jsonify({'error': 'connectionId required'}), 400
    tables_count = len(b.get('tables', [])) or b.get('tables_count', 0)
    if tables_count > 3:
        org_id = request.headers.get('X-Org-Id')
        current = get_org_plan(org_id)
        if _PLAN_RANK.get(current, 0) < _PLAN_RANK['pro']:
            return jsonify({
                'error': 'Upgrade required',
                'required_plan': 'pro',
                'current_plan': current,
            }), 403
    lid = execute('INSERT INTO governance_runs (connection_id,status,current_step) VALUES (?,?,?)',
                  (cid, 'running', 0))
    log_action('governance.started', 'governance_run', lid, {'connection_id': cid})
    simulate_governance(lid)
    return jsonify({'runId': lid}), 201

@app.get('/api/governance/stream/<int:run_id>')
def governance_stream(run_id):
    q = Queue()
    with _gov_lock:
        _gov_clients.setdefault(run_id, []).append(q)

    def generate():
        row = one('SELECT * FROM governance_runs WHERE id=?', (run_id,))
        if row:
            yield f"data: {json.dumps({'step': row['current_step'], 'status': row['status']})}\n\n"
        if row and row['status'] == 'done':
            return
        try:
            while True:
                try:
                    data = q.get(timeout=15)
                    if data is None:
                        return
                    yield f"data: {json.dumps(data)}\n\n"
                    if data.get('status') == 'done':
                        return
                except Empty:
                    yield ': ping\n\n'
        finally:
            with _gov_lock:
                lst = _gov_clients.get(run_id, [])
                if q in lst:
                    lst.remove(q)

    return sse_response(generate)

@app.get('/api/governance/<int:id>')
def get_governance(id):
    row = one('SELECT * FROM governance_runs WHERE id=?', (id,))
    return jsonify(row) if row else (jsonify({'error': 'Not found'}), 404)



# ─────────────────────────────────────────────────────────
# Custom DQ test authoring (R3S1E5) — safe expression subset only
# ─────────────────────────────────────────────────────────
import re as _dq_re

_IDENT = r'[A-Za-z_][A-Za-z0-9_]*'
_NUM = r'-?\d+(?:\.\d+)?'
_STR = r"'[^'\\;]*'"
_EXPR_RES = [
    _dq_re.compile(rf'^({_IDENT})\s*(=|!=|<>|<=|>=|<|>)\s*({_NUM}|{_STR})$'),
    _dq_re.compile(rf'^({_IDENT})\s+IS\s+(NOT\s+)?NULL$', _dq_re.I),
]


def compile_dq_expression(table: str, expression: str):
    """Compile `col op literal` / `col IS [NOT] NULL` into a violation-count
    query. Anything outside the subset is rejected — injection impossible."""
    expr = (expression or '').strip()
    if not _dq_re.match(rf'^{_IDENT}$', table or ''):
        return None, 'invalid table name'
    for rx in _EXPR_RES:
        m = rx.match(expr)
        if not m:
            continue
        col = m.group(1)
        if rx is _EXPR_RES[0]:
            op, lit = m.group(2), m.group(3)
            op = '!=' if op == '<>' else op
            predicate = f'"{col}" {op} {lit}'
        else:
            predicate = f'"{col}" IS {"NOT " if m.group(2) else ""}NULL'
        # NULLs are excluded from comparison checks (use IS NOT NULL for those);
        # SQL three-valued logic handles this: NOT (NULL op x) is NULL → unmatched.
        return f'SELECT COUNT(*) FROM "{table}" WHERE NOT ({predicate})', None
    return None, ('expression must be "<column> <op> <literal>" or '
                  '"<column> IS [NOT] NULL"')


RULE_THRESHOLDS = {
    'schema_fingerprint': 'recorded per run',
    'pk_uniqueness': '100% unique',
    'null_rate_for_key_columns': 'below key-column caps',
    'row_count_minimum': 'above table minimums',
    'freshness_sla': 'within table SLA',
    'pii_detection': 'no unapproved PII',
    'distribution_shift': 'within 50% of baseline',
}


def _rule_settings(cid):
    return {r['rule_id']: r for r in many(
        'SELECT rule_id, enabled, block_on_failure FROM dq_rule_settings '
        'WHERE connection_id=?', (cid,))}


@app.get('/api/dq/rules')
def list_dq_rules():
    """R32S1E4: merged rule catalog — system gate rules + custom tests,
    with per-connection enable / block-on-failure settings applied."""
    import dq as dq_mod
    cid = request.args.get('connection_id')
    if not cid:
        return jsonify({'error': 'connection_id required'}), 400
    settings = _rule_settings(cid)
    rules = []
    for rule_id, name, severity in dq_mod.MVP_RULES:
        s = settings.get(rule_id, {})
        block = s.get('block_on_failure')
        rules.append({
            'rule_id': rule_id, 'rule_name': name, 'severity': severity,
            'kind': 'system', 'threshold': RULE_THRESHOLDS.get(rule_id, '—'),
            'enabled': bool(s.get('enabled', 1)),
            'block_on_failure': bool(block if block is not None
                                     else severity == 'critical'),
        })
    for t in many('SELECT * FROM dq_custom_tests WHERE connection_id=? ORDER BY id', (cid,)):
        rid = f"custom:{t['id']}"
        s = settings.get(rid, {})
        block = s.get('block_on_failure')
        rules.append({
            'rule_id': rid, 'rule_name': t['expression'], 'severity': 'custom',
            'kind': 'custom', 'threshold': '0 violations',
            'table': t['table_name'], 'last_status': t['last_status'],
            'enabled': bool(t['enabled']),
            'block_on_failure': bool(block) if block is not None else False,
        })
    return jsonify({'rules': rules})


@app.put('/api/dq/rules/<rule_id>')
@require_role('admin', 'analyst')
def put_dq_rule(rule_id):
    import dq as dq_mod
    b = request.get_json() or {}
    cid = b.get('connectionId')
    if not cid:
        return jsonify({'error': 'connectionId required'}), 400
    known = {r[0] for r in dq_mod.MVP_RULES}
    if rule_id.startswith('custom:'):
        tid = rule_id.split(':', 1)[1]
        row = one('SELECT id FROM dq_custom_tests WHERE id=? AND connection_id=?', (tid, cid))
        if not row:
            return jsonify({'error': 'Custom test not found'}), 404
        if 'enabled' in b:
            execute('UPDATE dq_custom_tests SET enabled=? WHERE id=?',
                    (1 if b['enabled'] else 0, tid))
    elif rule_id not in known:
        return jsonify({'error': 'Unknown rule'}), 404
    prev = one('SELECT * FROM dq_rule_settings WHERE connection_id=? AND rule_id=?',
               (cid, rule_id)) or {}
    enabled = b.get('enabled', bool(prev.get('enabled', 1)))
    block = b.get('block_on_failure', prev.get('block_on_failure'))
    execute('INSERT INTO dq_rule_settings (connection_id, rule_id, enabled, block_on_failure) '
            'VALUES (?,?,?,?) ON CONFLICT(connection_id, rule_id) DO UPDATE SET '
            "enabled=excluded.enabled, block_on_failure=excluded.block_on_failure, "
            "updated_at=datetime('now')",
            (cid, rule_id, 1 if enabled else 0,
             None if block is None else (1 if block else 0)))
    log_action('dq.rule_updated', 'connection', cid,
               {'rule_id': rule_id, 'enabled': bool(enabled),
                'block_on_failure': None if block is None else bool(block)})
    return jsonify({'rule_id': rule_id, 'enabled': bool(enabled),
                    'block_on_failure': None if block is None else bool(block)})


@app.post('/api/dq/tests')
@require_role('admin', 'analyst')
def create_dq_test():
    b = request.get_json() or {}
    cid, table, expr = b.get('connectionId'), b.get('table'), b.get('expression')
    if not cid or not table:
        return jsonify({'error': 'connectionId and table required'}), 400
    sql, err = compile_dq_expression(table, expr)
    if err:
        return jsonify({'error': err}), 400
    lid = execute('INSERT INTO dq_custom_tests (connection_id, table_name, expression, '
                  'compiled_sql) VALUES (?,?,?,?)', (cid, table, expr.strip(), sql))
    log_action('dq.test_created', 'connection', cid, {'table': table, 'expression': expr})
    return jsonify({'id': lid, 'compiled_sql': sql, 'table': table,
                    'expression': expr.strip()}), 201


@app.get('/api/dq/tests')
def list_dq_tests():
    cid = request.args.get('connection_id')
    rows = (many('SELECT * FROM dq_custom_tests WHERE connection_id=? ORDER BY id', (cid,))
            if cid else many('SELECT * FROM dq_custom_tests ORDER BY id'))
    return jsonify(rows)


@app.post('/api/dq/tests/run')
def run_dq_tests():
    cid = request.args.get('connection_id') or (request.get_json(silent=True) or {}).get('connectionId')
    if not cid:
        return jsonify({'error': 'connection_id required'}), 400
    tests = many('SELECT * FROM dq_custom_tests WHERE connection_id=? AND enabled=1 '
                 'ORDER BY id', (cid,))
    db = get_db()
    results = []
    for t in tests:
        try:
            violations = db.execute(t['compiled_sql']).fetchone()[0]
            status = 'PASS' if violations == 0 else 'FAIL'
        except Exception as exc:
            violations, status = None, f'ERROR: {exc}'
        db.execute('UPDATE dq_custom_tests SET last_status=?, last_violations=?, '
                   "last_run_at=datetime('now') WHERE id=?",
                   (status if status in ('PASS', 'FAIL') else 'ERROR', violations, t['id']))
        results.append({'id': t['id'], 'table': t['table_name'],
                        'expression': t['expression'], 'status': status,
                        'violations': violations})
    db.commit()
    log_action('dq.tests_run', 'connection', cid,
               {'total': len(results),
                'failed': sum(1 for r in results if r['status'] == 'FAIL')})
    return jsonify({'connection_id': int(cid), 'results': results})


# ─────────────────────────────────────────────────────────
# Routes — DQ Gate engine (Sprint 13 / F-053)
# ─────────────────────────────────────────────────────────
@app.post('/api/dq/evaluate')
@limiter.limit('30/minute')
def dq_evaluate():
    import uuid
    import dq as dq_mod
    b = request.get_json() or {}
    manifest_dict, connection_id, baseline = None, None, None

    if b.get('manifest') is not None:
        manifest_dict = b['manifest']
    elif b.get('connectionId'):
        connection_id = b['connectionId']
        row = _manifest_row(connection_id, b.get('version'))
        if not row:
            return jsonify({'error': 'No manifest found for this connection'}), 404
        manifest_dict = json.loads(row['manifest_json'])
        prev = one('SELECT manifest_json FROM governance_manifests WHERE connection_id=? '
                   'AND id < (SELECT id FROM governance_manifests WHERE connection_id=? '
                   'AND version=?) ORDER BY id DESC LIMIT 1',
                   (connection_id, connection_id, manifest_dict.get('manifest_version')))
        baseline = json.loads(prev['manifest_json']) if prev else None
    else:
        return jsonify({'error': 'Provide either "manifest" or "connectionId"'}), 400

    errs = dq_mod.validate_manifest_input(manifest_dict)
    if errs:
        return jsonify({'error': 'Invalid governance manifest', 'errors': errs}), 400

    result = dq_mod.evaluate_manifest(
        manifest_dict, baseline=baseline,
        settings=_rule_settings(connection_id) if connection_id else None)
    result['trace_id'] = uuid.uuid4().hex

    execute('INSERT INTO dq_gate_results (connection_id, manifest_version, outcome, '
            'result_hash, trace_id, rules_json, evaluated_at) VALUES (?,?,?,?,?,?,?)',
            (connection_id, result['manifest_version'], result['outcome'],
             result['result_hash'], result['trace_id'], json.dumps(result['rules']),
             result['evaluated_at']))
    log_action('dq.evaluated', 'dq_gate_result', result['trace_id'],
               {'connection_id': connection_id, 'outcome': result['outcome'],
                'manifest_version': result['manifest_version'],
                'result_hash': result['result_hash']})

    # SSE dq_gate event for orchestrator consumers on the manifest's run
    run_id = manifest_dict.get('run_id')
    if run_id:
        broadcast_gov(run_id, {'type': 'dq_gate', 'run_id': run_id,
                               'outcome': result['outcome'],
                               'trace_id': result['trace_id']})

    status_code = 409 if result['outcome'] == 'BLOCK' else 200
    return jsonify({'dq_gate_result': result}), status_code


@app.get('/api/dq/results')
def dq_results():
    cid = request.args.get('connection_id')
    limit = min(int(request.args.get('limit', 50) or 50), 200)
    if cid:
        rows = many('SELECT * FROM dq_gate_results WHERE connection_id=? '
                    'ORDER BY id DESC LIMIT ?', (cid, limit))
    else:
        rows = many('SELECT * FROM dq_gate_results ORDER BY id DESC LIMIT ?', (limit,))
    for r in rows:
        r['rules'] = json.loads(r.pop('rules_json') or '[]')
    return jsonify(rows)


# ─────────────────────────────────────────────────────────
# Routes — dbt project import (R2S2E3)
# ─────────────────────────────────────────────────────────
_DBT_MEASURE_WORDS = ('amount', 'revenue', 'total', 'price', 'cost', 'value',
                      'quantity', 'count', 'spend', 'profit', 'margin')


def _dbt_column_kind(name):
    n = name.lower()
    if n == 'id' or n.endswith('_id'):
        return 'dimension', 'number'
    if any(w in n for w in ('date', '_at', 'timestamp', 'day', 'month')):
        return 'dimension', 'time'
    if any(w in n for w in _DBT_MEASURE_WORDS):
        return 'measure', None
    return 'dimension', 'string'


@app.post('/api/integrations/<int:id>/dbt_import')
@require_role('admin', 'analyst')
def dbt_import(id):
    import semantic_layer as sl
    if not one('SELECT id FROM connections WHERE id=?', (id,)):
        return jsonify({'error': 'Connection not found'}), 404
    manifest_body = request.get_json() or {}
    nodes = manifest_body.get('nodes')
    if not isinstance(nodes, dict) or not nodes:
        return jsonify({'error': 'A dbt manifest with a non-empty "nodes" map is required'}), 400

    # dbt tests → per-model quality notes
    notes_by_model = {}
    tests_mapped = 0
    for node in nodes.values():
        if node.get('resource_type') == 'test':
            meta = node.get('test_metadata') or {}
            kwargs = meta.get('kwargs') or {}
            model = kwargs.get('model') or ''
            col = kwargs.get('column_name') or ''
            if model:
                notes_by_model.setdefault(model, []).append(
                    f"dbt test {meta.get('name', 'unknown')} on {col}")
                tests_mapped += 1

    cubes, imported = [], []
    for node in nodes.values():
        if node.get('resource_type') != 'model':
            continue
        name = node.get('name')
        if not name:
            continue
        measures, dimensions = [], []
        primary_date = False
        for cname in (node.get('columns') or {}):
            kind, dtype = _dbt_column_kind(cname)
            if kind == 'measure':
                measures.append({'name': cname, 'sql': cname, 'aggregation': 'sum',
                                 'description': None, 'format': None,
                                 'allowed_filter_dimensions': None,
                                 'confidence': 'high', 'ml_allowed': True})
            else:
                dim = {'name': cname, 'sql': cname, 'type': dtype, 'confidence': 'high'}
                if dtype == 'time' and not primary_date:
                    dim['is_primary_date'] = True
                    primary_date = True
                dimensions.append(dim)
        cubes.append({
            'name': name,
            'sql_table': f"{node.get('schema') or 'public'}.{name}",
            'description': f'Imported from dbt project '
                           f"{(manifest_body.get('metadata') or {}).get('project_name', '')}".strip(),
            'dq_gate_status': 'PASS',
            'measures': measures, 'dimensions': dimensions, 'joins': [],
            'dq_notes': sorted(notes_by_model.get(name, [])),
            'source': 'dbt',
        })
        imported.append(name)
    if not imported:
        return jsonify({'error': 'Manifest contains no model nodes'}), 400

    row = _latest_schema_row('default')
    schema = json.loads(row['schema_json']) if row else {
        'schema_type': 'cube', 'generator': 'analytiq-semantic-builder',
        'workspace_id': 'default', 'cubes': [], 'notes': []}
    keep = [c for c in schema['cubes'] if c.get('name') not in set(imported)]
    schema['cubes'] = sorted(keep + cubes, key=lambda c: c['name'])
    schema['notes'] = sorted(set(schema.get('notes') or []) |
                             {f'dbt import: {len(imported)} model(s), '
                              f'{tests_mapped} test(s) mapped'})
    errs = sl.validate_cube_schema(schema)
    if errs:
        return jsonify({'error': 'Imported schema failed validation', 'errors': errs}), 422
    version = _persist_schema('default', schema, f'dbt import ({len(imported)} models)')
    log_action('dbt.imported', 'semantic_schema', id,
               {'models': imported, 'tests_mapped': tests_mapped, 'version': version})
    return jsonify({'version': version, 'imported_models': imported,
                    'tests_mapped': tests_mapped}), 201


# ─────────────────────────────────────────────────────────
# Routes — Semantic layer schema (Sprint 3 / F-011..F-013)
# ─────────────────────────────────────────────────────────
def _latest_schema_row(ws, version=None):
    if version:
        return one('SELECT * FROM semantic_schemas WHERE workspace_id=? AND version=?', (ws, version))
    return one('SELECT * FROM semantic_schemas WHERE workspace_id=? ORDER BY id DESC LIMIT 1', (ws,))


def _persist_schema(ws, new_schema, change_note):
    import semantic_layer as sl
    prev = _latest_schema_row(ws)
    prev_schema = json.loads(prev['schema_json']) if prev else None
    prev_version = prev['version'] if prev else None
    version = sl.next_schema_version(prev_schema, prev_version, new_schema)
    execute('INSERT INTO semantic_schemas (workspace_id, version, schema_json, change_note) '
            'VALUES (?,?,?,?)', (ws, version, json.dumps(new_schema), change_note))
    return version


@app.post('/api/semantic/<ws>/generate')
@require_role('admin', 'analyst')
def semantic_generate(ws):
    import semantic_layer as sl
    b = request.get_json() or {}
    cid = b.get('connectionId')
    if not cid:
        return jsonify({'error': 'connectionId required'}), 400
    row = _manifest_row(cid)
    if not row:
        return jsonify({'error': 'No governance manifest for this connection'}), 404
    m = json.loads(row['manifest_json'])
    schema = sl.build_cube_schema(m)
    errs = sl.validate_cube_schema(schema)
    if errs:
        return jsonify({'error': 'Generated schema failed validation', 'errors': errs}), 500
    version = _persist_schema(ws, schema, f"generated from manifest {m.get('manifest_version')}")
    log_action('semantic.schema_generated', 'semantic_schema', ws,
               {'version': version, 'manifest_version': m.get('manifest_version')})
    return jsonify({'version': version, 'schema': schema}), 201


CONF_NUM = {'high': 0.9, 'medium': 0.75, 'low': 0.55}


def _latest_manifest_health():
    row = one('SELECT manifest_json FROM governance_manifests ORDER BY id DESC LIMIT 1')
    if not row:
        return {}
    return {t['name']: t.get('health_score')
            for t in json.loads(row['manifest_json']).get('tables', [])}


@app.get('/api/semantic/<ws>/summary')
def semantic_summary(ws):
    """R32S2E1: overview KPIs for the semantic layer screens."""
    row = _latest_schema_row(ws)
    manifest = one('SELECT id, version, run_id FROM governance_manifests '
                   'ORDER BY id DESC LIMIT 1')
    m_status = None
    if manifest:
        pending = one('SELECT COUNT(*) AS n FROM semantic_definitions '
                      "WHERE run_id=? AND status='pending' AND confidence < ?",
                      (manifest['run_id'], REVIEW_CONFIDENCE_THRESHOLD))['n']
        m_status = 'REVIEW REQUIRED' if pending else 'ACTIVE'
    if not row:
        return jsonify({'exists': False, 'version': None, 'explores': 0,
                        'metrics': {'total': 0, 'governed': 0, 'draft': 0},
                        'dimensions': 0, 'join_paths': 0, 'conflicts': 0,
                        'pending_reviews': 0,
                        'manifest': ({'version': manifest['version'],
                                      'status': m_status} if manifest
                                     else {'version': None, 'status': None})})
    cubes = json.loads(row['schema_json']).get('cubes', [])
    measures = [ms for c in cubes for ms in c.get('measures', [])]
    governed = sum(1 for ms in measures if ms.get('confidence') == 'high')
    conflicts = one(
        'SELECT COUNT(DISTINCT p.name) AS n FROM semantic_definitions p '
        "WHERE p.status='pending' AND p.confidence < ? AND EXISTS ("
        "  SELECT 1 FROM semantic_definitions a WHERE a.name=p.name "
        "  AND a.status='accepted')", (REVIEW_CONFIDENCE_THRESHOLD,))['n']
    pending_total = one('SELECT COUNT(*) AS n FROM semantic_definitions '
                        "WHERE status='pending' AND confidence < ?",
                        (REVIEW_CONFIDENCE_THRESHOLD,))['n']
    return jsonify({
        'exists': True, 'version': row['version'],
        'explores': len(cubes),
        'metrics': {'total': len(measures), 'governed': governed,
                    'draft': len(measures) - governed},
        'dimensions': sum(len(c.get('dimensions', [])) for c in cubes),
        'join_paths': sum(len(c.get('joins', [])) for c in cubes),
        'conflicts': conflicts, 'pending_reviews': pending_total,
        'manifest': ({'version': manifest['version'], 'status': m_status}
                     if manifest else {'version': None, 'status': None}),
    })


# R32S2E3: per-dimension cardinality heuristics for the field picker.
DIM_CARDINALITY = {'week': 52, 'month': 12, 'quarter': 4, 'date': 365, 'day': 365,
                   'region': 4, 'state': 50, 'store': 42, 'name': 42, 'city': 60,
                   'segment': 5, 'tier': 3, 'channel': 4, 'category': 12,
                   'status': 5, 'id': 1000}


def _dim_card(name):
    for key, n in DIM_CARDINALITY.items():
        if key in name.lower():
            return n
    return 8


@app.post('/api/semantic/<ws>/preview')
def semantic_preview(ws):
    """R32S2E3 (DEP): bounded, read-only, deterministic field-picker preview.
    No warehouse round-trip — rows are seeded from the field selection so the
    same picks always preview identically. 100-row cap enforced."""
    t0 = time.time()
    b = request.get_json() or {}
    dims, measures = b.get('dimensions') or [], b.get('measures') or []
    if not dims and not measures:
        return jsonify({'error': 'Pick at least one dimension or measure'}), 400
    row = _latest_schema_row(ws)
    if not row:
        return jsonify({'error': 'No semantic schema found'}), 404
    cubes = json.loads(row['schema_json']).get('cubes', [])
    known_d = {d['name'] for c in cubes for d in c.get('dimensions', [])}
    known_m = {m['name'] for c in cubes for m in c.get('measures', [])}
    bad = [f for f in dims if f not in known_d] + [f for f in measures if f not in known_m]
    if bad:
        return jsonify({'error': f"Unknown fields: {', '.join(bad)}"}), 400

    series = 1
    for d in dims:
        series *= _dim_card(d)
    seed = sum(ord(ch) for ch in ''.join(dims + measures)) or 42
    rand = seeded_rng(seed)
    n_rows = min(100, max(4, series))
    rows = []
    for i in range(n_rows):
        r_ = []
        for d in dims:
            card = _dim_card(d)
            if any(k in d.lower() for k in ('week', 'month', 'date', 'day', 'quarter')):
                r_.append(f'2026-W{26 - (i % card) % 52:02d}')
            else:
                r_.append(f'{d.split("_")[0]}-{int(rand() * card) + 1}')
        for m in measures:
            base = 1000 + (sum(ord(ch) for ch in m) % 9000)
            r_.append(round(base * (0.6 + rand() * 0.8), 2))
        rows.append(r_)
    warning = None
    if series > 500:
        warning = (f'{series:,} series across {len(dims)} dimensions — '
                   'consider a Top-N before charting this.')
    return jsonify({'columns': dims + measures, 'rows': rows,
                    'row_count': len(rows), 'capped': True,
                    'elapsed_ms': round((time.time() - t0) * 1000, 1),
                    'series_estimate': series, 'warning': warning})


@app.get('/api/semantic/<ws>/conflicts')
def semantic_conflicts(ws):
    """R32S2E2: conflicted vocabulary — a pending low-confidence definition
    sharing a name with an accepted one; ids included for diff deep links."""
    rows = many(
        "SELECT p.id AS pending_id, p.name, p.confidence AS pending_confidence, "
        "p.explore, a.id AS accepted_id FROM semantic_definitions p "
        "JOIN semantic_definitions a ON a.name = p.name AND a.status='accepted' "
        "WHERE p.status='pending' AND p.confidence < ? "
        "GROUP BY p.name HAVING p.id = MAX(p.id)",
        (REVIEW_CONFIDENCE_THRESHOLD,))
    return jsonify({'conflicts': rows})


@app.get('/api/semantic/<ws>/explores')
def semantic_explores(ws):
    """R32S2E1: per-explore rows for the explores list + detail."""
    row = _latest_schema_row(ws)
    if not row:
        return jsonify({'explores': []})
    cubes = json.loads(row['schema_json']).get('cubes', [])
    health = _latest_manifest_health()
    out = []
    for c in cubes:
        own = (c.get('sql_table') or c['name']).split('.')[-1]
        tables = [own] + sorted({j.get('to') for j in c.get('joins', []) if j.get('to')})
        confs = [CONF_NUM.get(ms.get('confidence'), 0.75) for ms in c.get('measures', [])]
        confs += [CONF_NUM.get(d.get('confidence'), 0.75) for d in c.get('dimensions', [])]
        arts = _artifacts_using_explores([c['name']])
        out.append({
            'name': c['name'], 'tables': tables,
            'metrics': len(c.get('measures', [])),
            'dimensions': len(c.get('dimensions', [])),
            'joins': len(c.get('joins', [])),
            'health': health.get(own),
            'confidence': round(sum(confs) / len(confs), 2) if confs else 0.75,
            'dq_gate_status': c.get('dq_gate_status'),
            'used_by': len(arts),
            'used_by_artifacts': arts[:10],
        })
    return jsonify({'explores': out, 'version': row['version']})


@app.get('/api/semantic/<ws>/schema')
def semantic_schema(ws):
    row = _latest_schema_row(ws, request.args.get('version'))
    if not row:
        return jsonify({'error': 'No semantic schema found'}), 404
    return jsonify({'version': row['version'], 'schema': json.loads(row['schema_json']),
                    'change_note': row['change_note'], 'created_at': row['created_at']})


@app.get('/api/semantic/<ws>/schema/versions')
def semantic_schema_versions(ws):
    return jsonify(many(
        'SELECT id, version, change_note, created_at FROM semantic_schemas '
        'WHERE workspace_id=? ORDER BY id DESC', (ws,)))


@app.get('/api/semantic/<ws>/schema/diff')
def semantic_schema_diff(ws):
    import semantic_layer as sl
    v_from, v_to = request.args.get('from'), request.args.get('to')
    if not v_from or not v_to:
        return jsonify({'error': 'from and to versions required'}), 400
    r_from, r_to = _latest_schema_row(ws, v_from), _latest_schema_row(ws, v_to)
    if not r_from or not r_to:
        return jsonify({'error': 'Version not found'}), 404
    d = sl.diff_schemas(json.loads(r_from['schema_json']), json.loads(r_to['schema_json']))
    return jsonify({'from': v_from, 'to': v_to, **d})


@app.post('/api/semantic/<ws>/schema/rollback')
@require_role('admin', 'analyst')
def semantic_schema_rollback(ws):
    version = (request.get_json() or {}).get('version')
    if not version:
        return jsonify({'error': 'version required'}), 400
    row = _latest_schema_row(ws, version)
    if not row:
        return jsonify({'error': 'Version not found'}), 404
    schema = json.loads(row['schema_json'])
    new_version = _persist_schema(ws, schema, f'rollback to {version}')
    log_action('semantic.schema_rolled_back', 'semantic_schema', ws,
               {'to_version': version, 'new_version': new_version})
    return jsonify({'version': new_version, 'schema': schema})


@app.post('/api/semantic/<ws>/validate')
def semantic_validate(ws):
    import semantic_layer as sl
    errs = sl.validate_cube_schema(request.get_json() or {})
    return jsonify({'valid': not errs, 'errors': errs})


_PDT_SQL_RE = None


def _pdt_sql_ok(sql):
    s = (sql or '').strip().rstrip(';').strip()
    if not s or ';' in s:
        return None
    if not s.upper().startswith('SELECT'):
        return None
    banned = ('INSERT', 'UPDATE', 'DELETE', 'DROP', 'ALTER', 'CREATE', 'ATTACH', 'PRAGMA')
    tokens = set(w.upper() for w in s.replace('(', ' ').replace(')', ' ').split())
    if tokens & set(banned):
        return None
    return s


def _materialize_pdt(db, name, sql):
    table = f'pdt_{name}'
    db.execute(f'DROP TABLE IF EXISTS "{table}"')
    db.execute(f'CREATE TABLE "{table}" AS {sql}')
    n = db.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0]
    db.commit()
    return table, n


@app.post('/api/semantic/<ws>/pdts')
@require_role('admin')
def create_pdt(ws):
    import re as _re
    b = request.get_json() or {}
    name = b.get('name') or ''
    if not _re.match(r'^[a-z_][a-z0-9_]*$', name):
        return jsonify({'error': 'name must be a snake_case identifier'}), 400
    sql = _pdt_sql_ok(b.get('sql'))
    if not sql:
        return jsonify({'error': 'sql must be a single SELECT statement'}), 400
    if b.get('dry_run'):
        # R32S2E3: validate + count without persisting — temp materialization
        db = get_db()
        try:
            tmp = f'pdt_dryrun_{name}'
            db.execute(f'DROP TABLE IF EXISTS "{tmp}"')
            db.execute(f'CREATE TABLE "{tmp}" AS {sql}')
            n = db.execute(f'SELECT COUNT(*) FROM "{tmp}"').fetchone()[0]
            db.execute(f'DROP TABLE "{tmp}"')
            db.commit()
        except Exception as exc:
            return jsonify({'error': f'SQL failed: {exc}', 'valid': False}), 400
        return jsonify({'valid': True, 'row_count': n, 'dry_run': True})
    if one('SELECT id FROM pdts WHERE workspace_id=? AND name=?', (ws, name)):
        return jsonify({'error': f'PDT {name!r} already exists'}), 400
    db = get_db()
    try:
        table, n = _materialize_pdt(db, name, sql)
    except Exception as exc:
        return jsonify({'error': f'SQL failed: {exc}'}), 400
    execute('INSERT INTO pdts (workspace_id, name, sql, row_count, last_refreshed_at) '
            "VALUES (?,?,?,?,datetime('now'))", (ws, name, sql, n))
    log_action('pdt.created', 'semantic_schema', ws, {'name': name, 'rows': n})
    return jsonify({'name': name, 'table': table, 'row_count': n}), 201


@app.get('/api/semantic/<ws>/pdts')
def list_pdts(ws):
    return jsonify(many('SELECT * FROM pdts WHERE workspace_id=? ORDER BY name', (ws,)))


@app.post('/api/semantic/<ws>/pdts/<name>/refresh')
@require_role('admin', 'analyst')
def refresh_pdt(ws, name):
    row = one('SELECT * FROM pdts WHERE workspace_id=? AND name=?', (ws, name))
    if not row:
        return jsonify({'error': 'Not found'}), 404
    db = get_db()
    try:
        table, n = _materialize_pdt(db, name, row['sql'])
    except Exception as exc:
        return jsonify({'error': f'Refresh failed: {exc}'}), 400
    db.execute("UPDATE pdts SET row_count=?, last_refreshed_at=datetime('now') WHERE id=?",
               (n, row['id']))
    db.commit()
    log_action('pdt.refreshed', 'semantic_schema', ws, {'name': name, 'rows': n})
    return jsonify({'name': name, 'table': table, 'row_count': n})


@app.get('/api/semantic/<ws>/preagg_recommendations')
def preagg_recommendations(ws):
    """R3S2E6: recommend summary tables from observed query patterns."""
    rows = many("SELECT path, COUNT(*) AS hits, AVG(duration_ms) AS avg_ms "
                "FROM service_logs WHERE path LIKE '%/chart' OR path LIKE '/api/gold/%' "
                'GROUP BY path HAVING COUNT(*) >= 5 ORDER BY hits DESC LIMIT 10')
    recs = []
    for r in rows:
        slug = r['path'].strip('/').replace('/', '_').replace('api_', '')
        recs.append({
            'suggested_table': f'summary_{slug}',
            'source_path': r['path'],
            'hits': r['hits'],
            'avg_ms': round(r['avg_ms'] or 0, 2),
            'reason': (f"{r['hits']} requests observed — materializing a summary "
                       f'table would serve this from milliseconds'),
        })
    return jsonify(recs)


def record_activity(artifact_id, kind, detail=None):
    execute('INSERT INTO artifact_activity (artifact_id, kind, actor, detail) '
            'VALUES (?,?,?,?)',
            (artifact_id, kind, getattr(g, 'user_email', None) or current_role(), detail))


def _artifacts_using_explores(cube_names):
    """R3S2E5: artifacts whose session spec references any of these explores."""
    if not cube_names:
        return []
    out, seen = [], set()
    for art in many('SELECT * FROM artifacts WHERE pipeline_run_id IS NOT NULL'):
        run = one('SELECT session_id FROM pipeline_runs WHERE id=?', (art['pipeline_run_id'],))
        if not run or not run['session_id']:
            continue
        spec_row = one('SELECT spec_json FROM session_specs WHERE session_id=? '
                       'ORDER BY spec_version DESC LIMIT 1', (run['session_id'],))
        if not spec_row:
            continue
        used = set(json.loads(spec_row['spec_json']).get('explores_used') or [])
        if used & set(cube_names) and art['id'] not in seen:
            seen.add(art['id'])
            out.append({'id': art['id'], 'title': art['title']})
    return out


@app.get('/api/semantic/<ws>/impacts')
def semantic_impacts(ws):
    import semantic_layer as sl
    v_from, v_to = request.args.get('from'), request.args.get('to')
    if not v_from or not v_to:
        return jsonify({'error': 'from and to versions required'}), 400
    r_from, r_to = _latest_schema_row(ws, v_from), _latest_schema_row(ws, v_to)
    if not r_from or not r_to:
        return jsonify({'error': 'Version not found'}), 404
    d = sl.diff_schemas(json.loads(r_from['schema_json']), json.loads(r_to['schema_json']))
    changed = d['changed_cubes'] + d['removed_cubes'] + d['added_cubes']
    return jsonify({**d, 'from': v_from, 'to': v_to,
                    'impacted_artifacts': _artifacts_using_explores(changed)})


@app.post('/api/semantic/<ws>/metrics/calculated')
@require_role('admin', 'analyst')
def create_calculated_metric(ws):
    import re as _re
    import semantic_layer as sl
    from artifact_gen import METRIC_FORMATS
    b = request.get_json() or {}
    name, expr, fmt = b.get('name'), (b.get('expr') or '').strip(), b.get('format')
    if not name or not _re.match(r'^[a-z_][a-z0-9_]*$', name):
        return jsonify({'error': 'name must be a snake_case identifier'}), 400
    if fmt is not None and fmt not in METRIC_FORMATS:
        return jsonify({'error': f'format must be one of {METRIC_FORMATS}'}), 400
    row = _latest_schema_row(ws)
    if not row:
        return jsonify({'error': 'No semantic schema found'}), 404
    schema = json.loads(row['schema_json'])
    measures = {m['name']: c for c in schema['cubes'] for m in c.get('measures', [])}
    if name in measures:
        return jsonify({'error': f'metric {name!r} already exists'}), 400

    # safe arithmetic subset: existing measure names, numbers, + - * / ( )
    if not expr or not _re.match(r'^[\w\s+\-*/().0-9]+$', expr) or expr.rstrip()[-1] in '+-*/(':
        return jsonify({'error': 'expr must be arithmetic over existing metrics'}), 400
    idents = set(_re.findall(r'[a-zA-Z_][a-zA-Z0-9_]*', expr))
    unknown = sorted(i for i in idents if i not in measures)
    if unknown:
        return jsonify({'error': f'unknown metric(s) in expression: {", ".join(unknown)}'}), 400
    # R10S1E2: record metric derivation in the knowledge graph
    import knowledge_graph as kg
    for source in sorted(idents):
        kg.add_edge(get_db(), 'metric_derived_from_metric',
                    f'metric:{name}', f'metric:{source}')
    if not idents:
        return jsonify({'error': 'expression must reference at least one metric'}), 400

    host_cube = measures[sorted(idents)[0]]
    host_cube['measures'].append({
        'name': name, 'sql': expr, 'aggregation': 'derived', 'description': b.get('description'),
        'format': fmt, 'allowed_filter_dimensions': None,
        'confidence': 'high', 'ml_allowed': False, 'calculated': True,
        'inputs': sorted(idents),
    })
    errs = sl.validate_cube_schema(schema)
    if errs:
        return jsonify({'error': 'Validation failed', 'errors': errs}), 400
    version = _persist_schema(ws, schema, f'calculated metric: {name}')
    log_action('semantic.calculated_metric', 'semantic_schema', ws,
               {'name': name, 'expr': expr, 'version': version})
    return jsonify({'version': version, 'metric': name, 'cube': host_cube['name']}), 201


@app.post('/api/semantic/<ws>/explores')
@require_role('admin', 'analyst')
def create_explore(ws):
    import semantic_layer as sl
    cube = request.get_json() or {}
    row = _latest_schema_row(ws)
    schema = json.loads(row['schema_json']) if row else {
        'schema_type': 'cube', 'generator': 'analytiq-semantic-builder',
        'workspace_id': ws, 'cubes': [], 'notes': []}
    candidate = json.loads(json.dumps(schema))
    candidate['cubes'].append(cube)
    errs = sl.validate_cube_schema(candidate)
    if errs:
        return jsonify({'error': 'Validation failed', 'errors': errs}), 400
    version = _persist_schema(ws, candidate, f"explore created: {cube.get('name')}")
    log_action('semantic.explore_created', 'semantic_schema', ws,
               {'explore': cube.get('name'), 'version': version})
    return jsonify({'version': version, 'schema': candidate}), 201


@app.patch('/api/semantic/<ws>/explores/<name>')
@require_role('admin', 'analyst')
def edit_explore(ws, name):
    import semantic_layer as sl
    denied = acl_allows('explore', name, 'write')
    if denied:
        return denied
    b = request.get_json() or {}
    dry_run = bool(b.pop('dry_run', False))
    row = _latest_schema_row(ws)
    if not row:
        return jsonify({'error': 'No semantic schema found'}), 404
    schema = json.loads(row['schema_json'])
    cube = next((c for c in schema['cubes'] if c.get('name') == name), None)
    if cube is None:
        return jsonify({'error': f'Explore {name!r} not found'}), 404
    cube.update({k: v for k, v in b.items() if k != 'name'})
    errs = sl.validate_cube_schema(schema)
    if dry_run:
        return jsonify({'valid': not errs, 'errors': errs, 'preview': cube})
    if errs:
        return jsonify({'error': 'Validation failed', 'errors': errs}), 400
    version = _persist_schema(ws, schema, f'explore edited: {name}')
    impacted = _artifacts_using_explores([name])
    log_action('semantic.explore_edited', 'semantic_schema', ws,
               {'explore': name, 'version': version,
                'impacted_artifacts': [a['id'] for a in impacted]})
    return jsonify({'version': version, 'schema': schema,
                    'impacted_artifacts': impacted})


# ─────────────────────────────────────────────────────────
# Routes — Human review queue (Sprint 3 / F-014)
# ─────────────────────────────────────────────────────────
REVIEW_CONFIDENCE_THRESHOLD = 0.70


@app.get('/api/governance/summary')
def governance_summary():
    """R32S1E1-US1: the governance overview's KPI aggregate (ch15 §1) —
    composed from the real substrate; each metric degrades to 0/None rather
    than failing the card grid."""
    def safe(fn, default=0):
        try:
            return fn()
        except Exception:
            return default
    pend = "SELECT COUNT(*) AS n FROM semantic_definitions WHERE status='pending'"
    d = {
        'awaiting_review': safe(lambda: one(pend)['n']),
        'review_high': safe(lambda: one(
            "SELECT COUNT(*) AS n FROM semantic_definitions "
            "WHERE status='pending' AND COALESCE(confidence,1) < 0.7")['n']),
        'pii_flags': safe(lambda: one(
            "SELECT COUNT(*) AS n FROM semantic_definitions "
            "WHERE status='pending' AND LOWER(COALESCE(type,'')) = 'pii'")['n']),
        'tables_blocked': safe(lambda: one(
            "SELECT COUNT(DISTINCT subject) AS n FROM alerts "
            "WHERE LOWER(type) LIKE '%contract%' OR LOWER(type) LIKE '%block%'")['n']),
        'freshness_breaches': safe(lambda: one(
            "SELECT COUNT(*) AS n FROM alerts WHERE LOWER(type) LIKE '%fresh%' "
            "OR LOWER(type) LIKE '%sla%'")['n']),
        'schema_drift': safe(lambda: one(
            "SELECT COUNT(*) AS n FROM alerts WHERE LOWER(type) LIKE '%drift%'")['n']),
        'contract_failures_7d': safe(lambda: one(
            "SELECT COUNT(*) AS n FROM dq_gate_results WHERE LOWER(status) NOT IN "
            "('pass','passed') AND created_at >= datetime('now','-7 days')")['n']),
        'health_score': safe(lambda: (lambda r: int(r['s']) if r and r['s'] is not None else None)(
            one("SELECT AVG(health_score) AS s FROM health_history "
                "WHERE run_id = (SELECT MAX(run_id) FROM health_history)")), None),
        'health_trend': safe(lambda: [int(r['s']) for r in many(
            "SELECT AVG(health_score) AS s FROM health_history "
            "GROUP BY run_id ORDER BY run_id DESC LIMIT 12")][::-1], []),
    }
    return jsonify(d)


@app.get('/api/governance/latest')
def latest_governance_run():
    """R10S2E6: newest governance run — lets review surfaces deep-link."""
    row = one('SELECT id, connection_id FROM governance_runs ORDER BY id DESC LIMIT 1')
    if not row:
        return jsonify({'error': 'No governance runs yet'}), 404
    # R32S1E4: connection_id lets rule surfaces scope settings without a wizard context
    return jsonify({'run_id': row['id'], 'connection_id': row['connection_id']})


@app.get('/api/reviews/<int:run_id>')
def review_queue(run_id):
    items = many(
        "SELECT * FROM semantic_definitions WHERE run_id=? AND status='pending' "
        "AND confidence < ? ORDER BY confidence ASC",
        (run_id, REVIEW_CONFIDENCE_THRESHOLD))
    # R10S2E6: evidence-ranked triage (?ranked=1). Legacy confidence-ASC
    # order is the default; triage reorders and annotates, never authorizes.
    if request.args.get('ranked'):
        import governance_review as gr
        items = gr.annotate_and_rank(get_db(), items)
    return jsonify(items)


@app.get('/api/reviews/items/<int:def_id>')
def review_item_detail(def_id):
    """R32S1E3-US1: definition-review diff data — the pending item, the
    accepted CURRENT counterpart (same name, most recent), and the built
    dashboards the approval will re-validate."""
    item = one('SELECT * FROM semantic_definitions WHERE id=?', (def_id,))
    if not item:
        return jsonify({'error': 'Review item not found'}), 404
    current = one(
        "SELECT * FROM semantic_definitions WHERE name=? AND status='accepted' "
        'AND id != ? ORDER BY id DESC LIMIT 1', (item['name'], def_id))
    affected = many(
        'SELECT a.id, a.title FROM artifacts a ORDER BY a.id DESC LIMIT 12')
    return jsonify({'item': item, 'current': current,
                    'affected_count': len(affected), 'affected': affected})


@app.post('/api/reviews/items/<int:def_id>')
@require_role('admin', 'analyst')
def review_action(def_id):
    b = request.get_json() or {}
    action = b.get('action')
    if action not in ('accept', 'edit', 'reject'):
        return jsonify({'error': "action must be 'accept', 'edit', or 'reject'"}), 400
    row = one('SELECT * FROM semantic_definitions WHERE id=?', (def_id,))
    if not row:
        return jsonify({'error': 'Not found'}), 404
    db = get_db()
    if action == 'accept':
        db.execute("UPDATE semantic_definitions SET status='accepted', "
                   "confidence=MAX(COALESCE(confidence,0), 0.95) WHERE id=?", (def_id,))
        log_action('semantic.accepted', 'semantic_definition', def_id)
    elif action == 'edit':
        if not b.get('definition'):
            return jsonify({'error': 'definition text required for edit'}), 400
        db.execute("UPDATE semantic_definitions SET definition=?, confidence=0.96, "
                   "status='accepted' WHERE id=?", (b['definition'], def_id))
        log_action('semantic.edited', 'semantic_definition', def_id)
    else:
        db.execute("UPDATE semantic_definitions SET status='rejected' WHERE id=?", (def_id,))
        log_action('semantic.rejected', 'semantic_definition', def_id)
    db.commit()
    return jsonify(one('SELECT * FROM semantic_definitions WHERE id=?', (def_id,)))


# ─────────────────────────────────────────────────────────
# Routes — Tables
# ─────────────────────────────────────────────────────────
@app.get('/api/tables/<int:run_id>')
def get_tables(run_id):
    return jsonify(many('SELECT * FROM cataloged_tables WHERE run_id=? ORDER BY health_score DESC', (run_id,)))


# ─────────────────────────────────────────────────────────
# Routes — Semantic
# ─────────────────────────────────────────────────────────
@app.get('/api/semantic/<int:run_id>')
def get_semantic(run_id):
    return jsonify(many('SELECT * FROM semantic_definitions WHERE run_id=? ORDER BY confidence ASC', (run_id,)))

@app.patch('/api/semantic/<int:id>')
@require_role('admin', 'analyst')
def update_semantic(id):
    b      = request.get_json() or {}
    status = b.get('status')
    defn   = b.get('definition')
    db = get_db()
    if defn is not None:
        db.execute('UPDATE semantic_definitions SET definition=?,confidence=?,status=? WHERE id=?',
                   (defn, 0.96, 'accepted', id))
        log_action('semantic.edited', 'semantic_definition', id)
    elif status:
        db.execute('UPDATE semantic_definitions SET status=? WHERE id=?', (status, id))
        action = {'accepted': 'semantic.accepted', 'rejected': 'semantic.rejected'}.get(status, f'semantic.{status}')
        log_action(action, 'semantic_definition', id)
    db.commit()
    return jsonify(one('SELECT * FROM semantic_definitions WHERE id=?', (id,)))


# ─────────────────────────────────────────────────────────
# Routes — Sessions
# ─────────────────────────────────────────────────────────
@app.get('/api/sessions/<int:id>')
def get_session(id):
    row = one('SELECT * FROM sessions WHERE id=?', (id,))
    return jsonify(row) if row else (jsonify({'error': 'Not found'}), 404)

def _spec_for_session(session_id):
    row = one('SELECT * FROM session_specs WHERE session_id=? '
              'ORDER BY spec_version DESC LIMIT 1', (session_id,))
    return json.loads(row['spec_json']) if row else None


def _create_session_with_spec(base_spec, overrides, parent_id=None):
    import planner
    spec = {**base_spec, **{k: v for k, v in (overrides or {}).items() if k != 'name'}}
    errs = planner.validate_session_spec(spec)
    if errs:
        return None, errs
    sid = execute('INSERT INTO sessions (parent_session_id, metric, grain, horizon, '
                  "training_start, training_end, status) VALUES (?,?,?,?,?,?, 'pending')",
                  (parent_id, spec.get('target_metric'), spec.get('grain'),
                   spec.get('prediction_horizon') or 14,
                   (spec.get('date_range') or {}).get('start', '2023-01-01'),
                   (spec.get('date_range') or {}).get('end', '2023-12-31')))
    import hashlib as _h
    execute('INSERT INTO session_specs (session_id, spec_version, payload_hash, spec_json) '
            'VALUES (?,?,?,?)',
            (sid, 1, _h.sha256(json.dumps(spec, sort_keys=True).encode()).hexdigest(),
             json.dumps(spec)))
    return sid, None


@app.get('/api/sessions')
def list_sessions():
    """R4S1E1: session history with latest spec + downstream artifacts."""
    out = []
    for s in many('SELECT * FROM sessions ORDER BY id DESC LIMIT 100'):
        arts = many('SELECT a.id, a.title FROM artifacts a JOIN pipeline_runs pr '
                    'ON pr.id = a.pipeline_run_id WHERE pr.session_id=?', (s['id'],))
        out.append({**s, 'spec': _spec_for_session(s['id']), 'artifacts': arts})
    return jsonify(out)


@app.post('/api/sessions/<int:id>/fork')
@require_role('admin', 'analyst')
def fork_session(id):
    if not one('SELECT id FROM sessions WHERE id=?', (id,)):
        return jsonify({'error': 'Session not found'}), 404
    base = _spec_for_session(id)
    if not base:
        return jsonify({'error': 'Session has no confirmed spec to fork'}), 409
    sid, errs = _create_session_with_spec(base, request.get_json() or {}, parent_id=id)
    if errs:
        return jsonify({'error': 'Fork produced an invalid spec', 'errors': errs}), 400
    log_action('session.forked', 'session', sid, {'parent': id})
    return jsonify(one('SELECT * FROM sessions WHERE id=?', (sid,))), 201


@app.post('/api/templates')
@require_role('admin', 'analyst')
def create_template():
    import planner
    b = request.get_json() or {}
    if not b.get('name') or not isinstance(b.get('spec'), dict):
        return jsonify({'error': 'name and spec required'}), 400
    errs = planner.validate_session_spec(b['spec'])
    if errs:
        return jsonify({'error': 'Invalid template spec', 'errors': errs}), 400
    lid = execute('INSERT INTO session_templates (name, spec_json, created_by) VALUES (?,?,?)',
                  (b['name'], json.dumps(b['spec']), getattr(g, 'user_email', None)))
    log_action('template.created', 'session_template', lid, {'name': b['name']})
    return jsonify({'id': lid, 'name': b['name']}), 201


@app.get('/api/templates')
def list_templates():
    rows = many('SELECT * FROM session_templates ORDER BY id DESC')
    for r in rows:
        r['spec'] = json.loads(r.pop('spec_json'))
    return jsonify(rows)


@app.post('/api/sessions/from_template/<int:tid>')
@require_role('admin', 'analyst')
def session_from_template(tid):
    row = one('SELECT * FROM session_templates WHERE id=?', (tid,))
    if not row:
        return jsonify({'error': 'Template not found'}), 404
    sid, errs = _create_session_with_spec(json.loads(row['spec_json']),
                                          request.get_json() or {})
    if errs:
        return jsonify({'error': 'Template overrides produced an invalid spec',
                        'errors': errs}), 400
    log_action('session.from_template', 'session', sid, {'template_id': tid})
    return jsonify(one('SELECT * FROM sessions WHERE id=?', (sid,))), 201


@app.post('/api/sessions/<int:id>/message')
@limiter.limit('30/minute')
def session_message(id):
    """R4S1E2: PRD streaming protocol — planning → agent_start →
    (agent_complete | human_required | error), persisted for replay."""
    import planner
    if not one('SELECT id FROM sessions WHERE id=?', (id,)):
        return jsonify({'error': 'Session not found'}), 404
    b = request.get_json() or {}
    message = (b.get('message') or '').strip()
    if not message:
        return jsonify({'error': 'message required'}), 400

    schema_row = _latest_schema_row(b.get('workspaceId', 'default'))
    semantic_schema = json.loads(schema_row['schema_json']) if schema_row else None
    if semantic_schema:
        semantic_schema['cubes'] = [c for c in semantic_schema['cubes']
                                    if acl_permits('explore', c.get('name'), 'read')]
    manifest_dict = None
    if b.get('connectionId'):
        mrow = _manifest_row(b['connectionId'])
        manifest_dict = json.loads(mrow['manifest_json']) if mrow else None

    # the generator outlives the request-scoped connection → dedicated conn
    stream_conn = _new_conn()

    def emit(event_type, payload):
        stream_conn.execute('INSERT INTO session_events (session_id, type, payload_json) '
                            'VALUES (?,?,?)', (id, event_type, json.dumps(payload)))
        stream_conn.commit()
        return f'data: {json.dumps({"type": event_type, "payload": payload})}\n\n'

    def generate():
        try:
            yield from _generate_inner()
        finally:
            try:
                stream_conn.close()
            except Exception:
                pass

    def _generate_inner():
        yield emit('planning', {'message': message})
        yield emit('agent_start', {'agent': 'session_planner'})
        try:
            result = planner.plan_session(
                message, semantic_schema=semantic_schema, manifest=manifest_dict,
                schema_version=schema_row['version'] if schema_row else None)
        except Exception as exc:
            yield emit('error', {'error': str(exc),
                                 'remediation': 'Rephrase the question and retry.'})
            return
        if result.get('needs_clarification'):
            yield emit('human_required', {'question': result['question'],
                                          'options': result['options'],
                                          'intent_guess': result.get('intent_guess')})
            return
        errs = planner.validate_session_spec(result)
        if errs:
            yield emit('error', {'error': 'Planner produced an invalid spec',
                                 'errors': errs})
            return
        yield emit('agent_complete', {'agent': 'session_planner', **result})

    return sse_response(generate)


@app.get('/api/sessions/<int:id>/events')
def list_session_events(id):
    rows = many('SELECT * FROM session_events WHERE session_id=? ORDER BY id', (id,))
    for r in rows:
        r['payload'] = json.loads(r.pop('payload_json') or '{}')
    return jsonify(rows)


def _adaptive_planner_threshold(conn, user):
    """R10S2E4: expertise-conditioned clarification threshold.
    novice (<3 recorded steps): base + 0.05 — clarify more.
    expert (≥10 steps, ≥60% intent consistency): base - 0.10 — skip
    redundant clarification, surface assumptions inline instead."""
    row = conn.execute("SELECT value FROM platform_settings WHERE key='planner_threshold'").fetchone()
    try:
        base = float(row['value']) if row else 0.85
    except (TypeError, ValueError):
        base = 0.85
    rows = conn.execute('SELECT detail_json FROM intent_history WHERE user_id=?',
                        (user,)).fetchall()
    n = len(rows)
    intents = {}
    for r in rows:
        i = json.loads(r['detail_json'] or '{}').get('intent')
        if i:
            intents[i] = intents.get(i, 0) + 1
    top_share = (max(intents.values()) / sum(intents.values())) if intents else 0.0
    if n < 3:
        return round(min(0.95, base + 0.05), 2), 'novice', base
    if n >= 10 and top_share >= 0.6:
        return round(max(0.5, base - 0.10), 2), 'expert', base
    return base, 'standard', base


@app.get('/api/planner/threshold')
def planner_threshold():
    user = getattr(g, 'user_email', None) or 'default'
    eff, mode, base = _adaptive_planner_threshold(get_db(), user)
    return jsonify({'base': base, 'effective': eff, 'mode': mode, 'user': user})


@app.put('/api/platform/planner_threshold')
@require_role('admin')
def put_planner_threshold():
    b = request.get_json() or {}
    try:
        thr = float(b.get('threshold'))
        assert 0.5 <= thr <= 0.99
    except (TypeError, ValueError, AssertionError):
        return jsonify({'error': 'threshold must be a float in [0.5, 0.99]'}), 400
    execute("INSERT INTO platform_settings (key, value) VALUES ('planner_threshold', ?) "
            "ON CONFLICT(key) DO UPDATE SET value=excluded.value, updated_at=datetime('now')",
            (str(thr),))
    log_action('platform.planner_threshold_set', 'platform', None, {'threshold': thr})
    return jsonify({'threshold': thr})


@app.get('/api/reuse_candidates')
def reuse_candidates():
    """R10S2E7: prior validated plans as candidate starting points."""
    import knowledge_reuse as kr
    metric = request.args.get('metric')
    if not metric:
        return jsonify({'error': 'metric required'}), 400
    return jsonify({'metric': metric,
                    'candidates': kr.find_candidates(get_db(), metric,
                                                     grain=request.args.get('grain'))})


@app.post('/api/sessions/<int:id>/reuse/<plan_uid>')
def reuse_plan(id, plan_uid):
    """R10S2E7: apply a prior plan to this session — through the FULL spec
    validation gate set for the new context (never a governance bypass)."""
    import knowledge_reuse as kr
    import planner
    import uas
    sess = one('SELECT * FROM sessions WHERE id=?', (id,))
    if not sess:
        return jsonify({'error': 'Session not found'}), 404
    node = uas.get_by_uid(get_db(), plan_uid)
    if not node:
        return jsonify({'error': 'Plan not found'}), 404
    if node['artifact_type'] != 'dashboard_plan':
        return jsonify({'error': f"Artifact is a {node['artifact_type']}, not a dashboard plan"}), 422
    payload = json.loads(node['payload_json'])
    spec = kr.build_spec_from_plan(payload, sess)
    errs = planner.validate_session_spec(spec)
    if errs:
        return jsonify({'error': 'Reused plan failed validation for this context',
                        'validation': {'errors': errs}}), 422
    prev = one('SELECT MAX(spec_version) AS v FROM session_specs WHERE session_id=?', (id,))
    version = (prev['v'] or 0) + 1
    import hashlib as _hl
    execute('INSERT INTO session_specs (session_id, spec_version, payload_hash, spec_json) '
            'VALUES (?,?,?,?)',
            (id, version, _hl.sha256(json.dumps(spec, sort_keys=True).encode()).hexdigest(),
             json.dumps(spec)))
    log_action('plan.reused', 'session', id,
               {'plan_uid': plan_uid, 'similarity_source': 'uas+kg', 'spec_version': version})
    return jsonify({'session_id': id, 'spec_version': version, 'spec': spec,
                    'validation': {'errors': []}}), 201


@app.get('/api/sessions/warm_start')
def sessions_warm_start():
    """R10S1E3: likely intent categories + recent metrics for this user —
    context, never a pre-committed plan."""
    import intent_history as ih
    return jsonify(ih.warm_start(get_db(), getattr(g, 'user_email', None) or 'default'))


@app.post('/api/sessions/plan')
@limiter.limit('30/minute')
def plan_session_route():
    import planner
    b = request.get_json() or {}
    message = (b.get('message') or '').strip()
    if not message:
        return jsonify({'error': 'message required'}), 400

    schema_row, schema_version = None, None
    ws = b.get('workspaceId', 'default')
    schema_row = _latest_schema_row(ws)
    semantic_schema = json.loads(schema_row['schema_json']) if schema_row else None
    schema_version = schema_row['version'] if schema_row else None
    # R3S2E4: drop explores the caller isn't permitted to use
    if semantic_schema:
        semantic_schema['cubes'] = [c for c in semantic_schema['cubes']
                                    if acl_permits('explore', c.get('name'), 'read')]

    manifest_dict = None
    cid = b.get('connectionId')
    if cid:
        mrow = _manifest_row(cid)
        manifest_dict = json.loads(mrow['manifest_json']) if mrow else None

    # R9S1E1: resolve through the cost ladder (cache → template → small → frontier)
    import orchestrator
    import re as _re
    norm = _re.sub(r'\s+', ' ', message.lower()).strip()
    pattern = [t for t in _re.split(r'[^a-z]+', norm) if t and not t.isdigit()]
    orchestrator.dispatch(get_db(), 'intent_classification', ['intent', norm],
                          lambda: {'classified': True}, workspace=ws)
    # Cache identity must include the caller's effective explore visibility —
    # permission-scoped results can never leak across differently-permitted
    # callers (caught by R3S2E4 regression).
    import hashlib as _hl
    visible = sorted(c.get('name', '') for c in (semantic_schema or {}).get('cubes', []))
    acl_fp = _hl.sha256('|'.join(visible).encode()).hexdigest()[:16]
    # R10S2E4: expertise-conditioned threshold; the cache signature includes
    # the planner mode so expert plans never serve novice callers.
    _user = getattr(g, 'user_email', None) or 'default'
    _thr, _mode, _base = _adaptive_planner_threshold(get_db(), _user)
    disp = orchestrator.dispatch(
        get_db(), 'session_planning',
        ['plan', norm, schema_version or 'no-schema', acl_fp, f'mode:{_mode}'],
        lambda: planner.plan_session(message, semantic_schema=semantic_schema,
                                     manifest=manifest_dict, schema_version=schema_version,
                                     clarify_threshold=_thr),
        pattern=pattern, workspace=ws)
    result = disp['result']
    result['planner_mode'] = _mode
    if _mode == 'expert' and not result.get('needs_clarification'):
        result['assumptions'] = [
            f"grain assumed {result.get('grain', 'Location · Day')} (say e.g. 'weekly by store' to change)",
            f"horizon assumed {result.get('prediction_horizon') or 14} days",
            'date range assumed trailing 12 months',
        ]
    # R10S1E1: stored preference is a prior, never a constraint — apply the
    # remembered grain only when the current turn is silent about grain.
    if not result.get('needs_clarification'):
        import agent_memory as am
        explicit_grain = bool(_re.search(
            r'\b(daily|weekly|monthly|per (day|week|month)|by (day|week|month|location|customer|region|store|channel))\b',
            norm))
        if not explicit_grain:
            prior = am.recall(get_db(), 'planner', category='filter_pattern', workspace=ws)
            grain_prior = next((p for p in prior if p['mem_key'].startswith('grain')), None)
            if grain_prior:
                result['grain'] = grain_prior['value']
                result['memory_applied'] = True
        # R16S1E1: disclose masked PII columns so the plan card can render
        # its ACCESS row (§7.5) — the user sees what the plan cannot touch.
        import pii as _pii
        masked = []
        for t in (manifest_dict or {}).get('tables', []):
            for col in (t.get('columns') or []):
                flag = col.get('pii_flags')
                if flag and flag.get('confidence', 0) >= _pii.BLOCK_CONFIDENCE:
                    masked.append(f"{t['name']}.{col['name']}")
        result['access_limitations'] = {
            'masked_columns': sorted(set(masked)),
            'note': (f'{len(set(masked))} PII column(s) excluded (masked)' if masked
                     else 'No PII restrictions apply to this plan'),
        }
        # R10S1E3: record the question step of the investigation chain
        import intent_history as ih
        ih.record(get_db(), getattr(g, 'user_email', None) or 'default', 'question',
                  detail={'intent': result.get('intent'),
                          'metric': result.get('target_metric'),
                          'message': norm[:160]})
    if not result.get('needs_clarification'):
        errs = planner.validate_session_spec(result)
        if errs:
            log_action('session.plan_invalid', 'session_spec', None, {'errors': [e['code'] for e in errs]})
            return jsonify({'error': 'Planner produced an invalid spec', 'errors': errs}), 500
        log_action('session.planned', 'session_spec', None,
                   {'intent': result['intent'], 'target_metric': result['target_metric']})
    return jsonify(result)


@app.post('/api/sessions/<int:id>/spec')
@require_role('admin', 'analyst')
def persist_session_spec(id):
    import hashlib
    import planner
    if not one('SELECT id FROM sessions WHERE id=?', (id,)):
        return jsonify({'error': 'Session not found'}), 404
    spec = request.get_json() or {}
    errs = planner.validate_session_spec(spec)
    if errs:
        return jsonify({'error': 'Validation failed', 'errors': errs}), 400

    payload_hash = hashlib.sha256(json.dumps(spec, sort_keys=True).encode()).hexdigest()
    idem_key = request.headers.get('Idempotency-Key')
    if idem_key:
        existing = one('SELECT * FROM session_specs WHERE session_id=? AND idempotency_key=? '
                       'AND payload_hash=?', (id, idem_key, payload_hash))
        if existing:
            return jsonify({'session_id': id, 'spec_version': existing['spec_version'],
                            'spec': json.loads(existing['spec_json']), 'replayed': True}), 200

    prev = one('SELECT MAX(spec_version) AS v FROM session_specs WHERE session_id=?', (id,))
    version = (prev['v'] or 0) + 1
    execute('INSERT INTO session_specs (session_id, spec_version, idempotency_key, payload_hash, spec_json) '
            'VALUES (?,?,?,?,?)', (id, version, idem_key, payload_hash, json.dumps(spec)))
    # R10S1E1: a confirmed spec is a real preference signal — remember the grain
    if spec.get('grain'):
        try:
            import agent_memory as am
            am.remember(get_db(), 'planner', 'filter_pattern', 'grain_pref', spec['grain'])
        except ValueError:
            pass  # PII-gated values are never stored
    # R10S1E2: co-analysis edges from the confirmed spec
    import knowledge_graph as kg
    kg.ingest_spec(get_db(), id, spec)
    # R10S1E3: spec step of the investigation chain
    import intent_history as ih
    ih.record(get_db(), getattr(g, 'user_email', None) or 'default', 'spec',
              ref_id=id, session_id=id,
              detail={'intent': spec.get('intent'), 'target_metric': spec.get('target_metric')})
    log_action('session.spec_confirmed', 'session_spec', id,
               {'spec_version': version, 'target_metric': spec.get('target_metric'),
                'semantic_layer_version': spec.get('semantic_layer_version'),
                'governance_manifest_version': spec.get('governance_manifest_version')})
    return jsonify({'session_id': id, 'spec_version': version, 'spec': spec}), 201


@app.get('/api/sessions/<int:id>/spec')
def get_session_spec(id):
    version = request.args.get('version')
    if version:
        row = one('SELECT * FROM session_specs WHERE session_id=? AND spec_version=?', (id, version))
    else:
        row = one('SELECT * FROM session_specs WHERE session_id=? ORDER BY spec_version DESC LIMIT 1', (id,))
    if not row:
        return jsonify({'error': 'No confirmed spec for this session'}), 404
    return jsonify({'session_id': id, 'spec_version': row['spec_version'],
                    'spec': json.loads(row['spec_json']), 'created_at': row['created_at']})


@app.post('/api/sessions')
def create_session():
    b   = request.get_json() or {}
    lid = execute(
        'INSERT INTO sessions (connection_id,run_id,metric,grain,horizon,training_start,training_end,status,is_sandbox) '
        'VALUES (?,?,?,?,?,?,?,?,?)',
        (b.get('connectionId'), b.get('runId'), b.get('metric', 'Net Revenue'),
         b.get('grain', 'Location · Day'), b.get('horizon', 14),
         b.get('training_start', '2023-01-01'), b.get('training_end', '2023-12-31'), 'pending',
         1 if b.get('sandbox') else 0),
    )
    return jsonify(one('SELECT * FROM sessions WHERE id=?', (lid,))), 201



# ─────────────────────────────────────────────────────────
# Routes — Data Modeler (Sprint 5 / F-021..F-025)
# ─────────────────────────────────────────────────────────
def _materialize_gold(conn, physical, columns, grain_keys, date_range):
    """Write deterministic synthetic rows at the requested grain (demo world:
    the INSERT-SELECT would run in the warehouse; here we materialize locally)."""
    from datetime import date as _date, timedelta as _td
    import warehouse as _wh
    d = _wh.get_dialect('sqlite')
    rand = seeded_rng(11)
    start = _date.fromisoformat(date_range['start'])
    end = _date.fromisoformat(date_range['end'])
    days = (end - start).days + 1
    locations = list(range(1, 9))
    rows = []
    for li in locations:
        for i in range(days):
            day = (start + _td(days=i)).isoformat()
            row = []
            for c in columns:
                n, t = c['name'], c['type']
                if n == 'day':
                    row.append(day)
                elif n.endswith('_id'):
                    row.append(li)
                elif t == 'measure':
                    row.append(round(400 + rand() * 400, 2))
                elif t == 'flag':
                    row.append(1 if rand() > 0.5 else 0)
                elif t == 'date':
                    row.append(day)
                else:
                    row.append(['gold', 'silver', 'bronze'][li % 3])
            rows.append(tuple(row))
    placeholders = ','.join(['?'] * len(columns))
    conn.executemany(f'INSERT INTO {d.quote_identifier(physical)} VALUES ({placeholders})', rows)
    conn.commit()
    return len(rows)


@app.post('/api/modeler/generate')
@limiter.limit('10/minute')
@require_role('admin', 'analyst')
def modeler_generate():
    import modeler
    b = request.get_json() or {}
    sid = b.get('sessionId')
    if not sid:
        return jsonify({'error': 'sessionId required'}), 400
    mode = b.get('mode', 'dry_run')
    if mode not in ('dry_run', 'execute'):
        return jsonify({'error': "mode must be 'dry_run' or 'execute'"}), 400

    spec_row = one('SELECT * FROM session_specs WHERE session_id=? '
                   'ORDER BY spec_version DESC LIMIT 1', (sid,))
    if not spec_row:
        return jsonify({'error': 'No confirmed session spec — confirm the spec first',
                        'remediation': f'POST /api/sessions/{sid}/spec'}), 409
    spec = json.loads(spec_row['spec_json'])

    ws = b.get('workspaceId', 'default')
    schema_row = _latest_schema_row(ws)
    if not schema_row:
        return jsonify({'error': 'No semantic schema — generate it first',
                        'remediation': f'POST /api/semantic/{ws}/generate'}), 409
    cube_schema = json.loads(schema_row['schema_json'])

    prev = one('SELECT MAX(version) AS v FROM gold_tables WHERE session_id=?', (sid,))
    version = (prev['v'] or 0) + 1
    try:
        out = modeler.generate_gold_sql(spec, cube_schema, workspace_id=ws,
                                        session_id=sid, version=version)
    except modeler.GoldGenerationError as e:
        log_action('modeler.generation_failed', 'gold_table', sid, {'error': str(e)})
        return jsonify({'error': str(e), 'dq_gate': 'BLOCK'}), 422

    date_range = spec.get('date_range') or {'start': '2023-01-01', 'end': '2023-12-31'}
    fanout = modeler.detect_fanout(cube_schema, spec.get('explores_used') or [])
    dq = {'leakage': out['leakage'], 'fanout': fanout}

    if mode == 'dry_run':
        import splits as splits_mod
        dq['split'] = splits_mod.validate_split(out['split_config'], row_count=12847,
                                                horizon=spec.get('prediction_horizon'))
        log_action('modeler.dry_run', 'gold_table', sid,
                   {'table': out['table_name'], 'output_hash': out['output_hash'],
                    'dq': {k: (v.get('status') if isinstance(v, dict) else None)
                           for k, v in dq.items()}})
        return jsonify({'mode': 'dry_run', 'table_name': out['table_name'],
                        'ddl': out['ddl'], 'insert_sql': out['insert_sql'],
                        'split_config': out['split_config'], 'output_hash': out['output_hash'],
                        'dq_gates': dq})

    # execute: materialize under a physical (SQLite-safe) name
    import splits as splits_mod
    import warehouse as _wh
    d = _wh.get_dialect('sqlite')
    # R33S1E2 fix: physical gold tables are session-scoped — identical specs
    # across sessions used to share one table, so the second write duplicated
    # grain keys and blocked itself. Re-executes for the same session start
    # from a clean slate (delete-then-insert keeps the write idempotent).
    physical = f"{out['table_name'].replace('.', '__')}__s{sid}"
    db = get_db()
    ddl_physical = d.compile_create_table({'table': physical, 'if_not_exists': True,
                                           'columns': out['columns']})
    db.execute(ddl_physical)
    db.execute(f'DELETE FROM {d.quote_identifier(physical)}')
    db.commit()
    expected = 8 * (( __import__('datetime').date.fromisoformat(date_range['end'])
                    - __import__('datetime').date.fromisoformat(date_range['start'])).days + 1)
    produced = _materialize_gold(db, physical, out['columns'], out['grain_keys'], date_range)

    grain = modeler.validate_grain(db, physical, out['grain_keys'])
    disc = abs(expected - produced) / max(expected, 1) * 100
    row_tol = {'status': 'PASS' if disc <= 2 else 'WARN' if disc <= 5 else 'BLOCK',
               'expected': expected, 'produced': produced,
               'discrepancy_pct': round(disc, 2)}
    dq['grain'] = grain
    dq['row_tolerance'] = row_tol
    dq['split'] = splits_mod.validate_split(out['split_config'], row_count=produced,
                                            horizon=spec.get('prediction_horizon'))
    status = 'written' if grain['status'] == 'PASS' and row_tol['status'] != 'BLOCK' else 'blocked'

    # Sprint 6 / F-026: immutable feature manifest per gold write
    import feature_manifest as fm_mod
    fm = fm_mod.build_feature_manifest(out, spec, ws, sid, cube_schema)
    fm = fm_mod.save_feature_manifest(get_db(), fm)
    log_action('feature_manifest.created', 'feature_manifest', fm['id'],
               {'manifest_version': fm['manifest_version'], 'session_id': sid,
                'gold_table_name': fm['gold_table_name'],
                'feature_count': len(fm['feature_list'])})

    lid = execute(
        'INSERT INTO gold_tables (session_id, table_name, physical_table, version, ddl, '
        'insert_sql, output_hash, row_count, status, manifest_version, split_config_json, dq_json) '
        'VALUES (?,?,?,?,?,?,?,?,?,?,?,?)',
        (sid, out['table_name'], physical, version, out['ddl'], out['insert_sql'],
         out['output_hash'], produced, status, fm['manifest_version'],
         json.dumps(out['split_config']), json.dumps(dq)))
    log_action('modeler.executed', 'gold_table', lid,
               {'table': out['table_name'], 'rows': produced, 'status': status,
                'output_hash': out['output_hash']})
    return jsonify({'mode': 'execute', 'id': lid, 'table_name': out['table_name'],
                    'physical_table': physical, 'row_count': produced,
                    'output_hash': out['output_hash'], 'status': status,
                    'split_config': out['split_config'], 'dq_gates': dq}), 201


def _latest_gold(sid):
    return one("SELECT * FROM gold_tables WHERE session_id=? AND status='written' "
               'ORDER BY version DESC LIMIT 1', (sid,))


@app.post('/api/modeler/custom_features')
@require_role('admin', 'analyst')
def create_custom_feature():
    import re as _re
    import modeler
    b = request.get_json() or {}
    sid, name, expr = b.get('sessionId'), b.get('name'), (b.get('expr') or '').strip()
    if not sid or not name or not _re.match(r'^[a-z_][a-z0-9_]*$', name or ''):
        return jsonify({'error': 'sessionId and snake_case name required'}), 400
    gold = _latest_gold(sid)
    if not gold:
        return jsonify({'error': 'No gold table for this session'}), 409
    db = get_db()
    cols = {c[1] for c in db.execute(
        f'PRAGMA table_info("{gold["physical_table"]}")').fetchall()}
    if not expr or not _re.match(r'^[\w\s+\-*/().0-9]+$', expr) or expr.rstrip()[-1] in '+-*/(':
        return jsonify({'error': 'expr must be arithmetic over existing gold columns'}), 400
    idents = set(_re.findall(r'[a-zA-Z_][a-zA-Z0-9_]*', expr))
    unknown = sorted(i for i in idents if i not in cols)
    if unknown or not idents:
        return jsonify({'error': f'unknown column(s): {", ".join(unknown) or "none referenced"}'}), 400

    spec = _spec_for_session(sid) or {}
    leak = modeler.scan_leakage([name], spec.get('target_metric') or 'Net Revenue',
                                spec.get('prediction_horizon'))['features'][0]
    lid = execute('INSERT INTO custom_features (session_id, name, expr, leakage_json) '
                  'VALUES (?,?,?,?)', (sid, name, expr, json.dumps(leak)))
    log_action('custom_feature.created', 'gold_table', sid,
               {'name': name, 'expr': expr, 'leakage': leak['action']})
    return jsonify({'id': lid, 'name': name, 'status': 'pending_review',
                    'leakage': leak}), 201


@app.post('/api/modeler/custom_features/<int:id>/approve')
@require_role('admin', 'analyst')
def approve_custom_feature(id):
    row = one('SELECT * FROM custom_features WHERE id=?', (id,))
    if not row:
        return jsonify({'error': 'Not found'}), 404
    db = get_db()
    db.execute("UPDATE custom_features SET status='approved' WHERE id=?", (id,))
    db.commit()
    log_action('custom_feature.approved', 'gold_table', row['session_id'],
               {'name': row['name']})
    return jsonify(one('SELECT * FROM custom_features WHERE id=?', (id,)))


@app.post('/api/modeler/custom_features/<int:id>/apply')
@require_role('admin', 'analyst')
def apply_custom_feature(id):
    row = one('SELECT * FROM custom_features WHERE id=?', (id,))
    if not row:
        return jsonify({'error': 'Not found'}), 404
    if row['status'] != 'approved':
        return jsonify({'error': 'Feature must be approved before application'}), 409
    gold = _latest_gold(row['session_id'])
    if not gold:
        return jsonify({'error': 'No gold table'}), 409
    db = get_db()
    physical = gold['physical_table']
    cols = {c[1] for c in db.execute(f'PRAGMA table_info("{physical}")').fetchall()}
    if row['name'] not in cols:
        db.execute(f'ALTER TABLE "{physical}" ADD COLUMN "{row["name"]}" REAL')
    # expr is validated arithmetic over quoted-safe identifiers
    db.execute(f'UPDATE "{physical}" SET "{row["name"]}" = {row["expr"]}')
    db.execute('UPDATE custom_features SET applied=1 WHERE id=?', (id,))
    db.commit()
    log_action('custom_feature.applied', 'gold_table', row['session_id'],
               {'name': row['name']})
    return jsonify({'id': id, 'applied': True, 'physical_table': physical})


@app.post('/api/modeler/leakage/confirm')
@require_role('admin', 'analyst')
def confirm_leakage():
    b = request.get_json() or {}
    sid, feats = b.get('sessionId'), b.get('features') or []
    if not sid or not feats:
        return jsonify({'error': 'sessionId and features required'}), 400
    db = get_db()
    for f in feats:
        db.execute('INSERT INTO leakage_confirmations (session_id, feature, justification, '
                   'confirmed_by) VALUES (?,?,?,?) '
                   'ON CONFLICT(session_id, feature) DO UPDATE SET '
                   'justification=excluded.justification',
                   (sid, f, b.get('justification'), getattr(g, 'user_email', None)))
    db.commit()
    log_action('leakage.confirmed', 'gold_table', sid,
               {'features': feats, 'justification': b.get('justification')})
    return jsonify({'session_id': sid, 'confirmed': feats})


@app.post('/api/calendars')
@require_role('admin', 'analyst')
def create_calendar():
    import re as _re
    b = request.get_json() or {}
    dates = b.get('dates')
    if not b.get('name') or not isinstance(dates, list) or not dates:
        return jsonify({'error': 'name and a non-empty dates list required'}), 400
    for d in dates:
        if not _re.match(r'^\d{4}-\d{2}-\d{2}$', str(d)):
            return jsonify({'error': f'dates must be ISO YYYY-MM-DD (got {d!r})'}), 400
    lid = execute('INSERT INTO holiday_calendars (name, dates_json) VALUES (?,?)',
                  (b['name'], json.dumps(sorted(dates))))
    log_action('calendar.created', 'holiday_calendar', lid, {'name': b['name'],
                                                             'days': len(dates)})
    return jsonify({'id': lid, 'name': b['name'], 'dates': sorted(dates)}), 201


@app.get('/api/calendars')
def list_calendars():
    rows = many('SELECT * FROM holiday_calendars ORDER BY id')
    for r in rows:
        r['dates'] = json.loads(r.pop('dates_json'))
    return jsonify(rows)


@app.post('/api/modeler/enrich')
@limiter.limit('10/minute')
@require_role('admin', 'analyst')
def modeler_enrich():
    """R4S2E1: Stage-3 feature engineering over the latest gold table."""
    import feature_engineering as fe
    b = request.get_json() or {}
    sid = b.get('sessionId')
    if not sid:
        return jsonify({'error': 'sessionId required'}), 400
    gold = one("SELECT * FROM gold_tables WHERE session_id=? AND status='written' "
               'ORDER BY version DESC LIMIT 1', (sid,))
    if not gold:
        return jsonify({'error': 'No gold table — run the modeler (execute) first'}), 409
    physical = gold['physical_table']
    db = get_db()

    cols = [c[1] for c in db.execute(f'PRAGMA table_info("{physical}")').fetchall()]
    target = next((c for c in cols if c.startswith('target_')), None)
    if not target or 'day' not in cols or 'location_id' not in cols:
        return jsonify({'error': 'gold table missing target/day/location_id'}), 422

    extra_holidays = set()
    for cal in many('SELECT dates_json FROM holiday_calendars'):
        extra_holidays |= set(json.loads(cal['dates_json']))

    # R6S2E2: annotations on this session's artifacts feed a binary feature
    _annotated_days = {r_['timestamp'] for r_ in many(
        'SELECT ann.timestamp FROM artifact_annotations ann '
        'JOIN artifacts a ON a.id = ann.artifact_id '
        'JOIN pipeline_runs pr ON pr.id = a.pipeline_run_id '
        'WHERE pr.session_id=? AND ann.timestamp IS NOT NULL', (sid,))}

    # compute per-location temporal features
    locations = [r_[0] for r_ in db.execute(
        f'SELECT DISTINCT location_id FROM "{physical}" ORDER BY location_id').fetchall()]
    feature_names = None
    updates = {}   # (location, day) → {feature: value}
    for loc in locations:
        rows = db.execute(f'SELECT day, "{target}" FROM "{physical}" '
                          f'WHERE location_id=? ORDER BY day', (loc,)).fetchall()
        days = [r_[0] for r_ in rows]
        series = [r_[1] for r_ in rows]
        feats = fe.temporal_features(series)
        feats['is_holiday'] = [1 if fe.is_holiday(d, extra=extra_holidays) else 0
                               for d in days]
        feats['is_annotated_event'] = [1 if d in _annotated_days else 0 for d in days]
        feature_names = list(feats)
        for i, d in enumerate(days):
            updates[(loc, d)] = {name: vals[i] for name, vals in feats.items()}

    added = []
    for name in feature_names or []:
        if name not in cols:
            db.execute(f'ALTER TABLE "{physical}" ADD COLUMN "{name}" REAL')
        added.append(name)
    for (loc, d), vals in updates.items():
        sets = ', '.join(f'"{n}"=?' for n in vals)
        db.execute(f'UPDATE "{physical}" SET {sets} WHERE location_id=? AND day=?',
                   (*vals.values(), loc, d))
    db.commit()

    # R4S2E2 — imputation (rolling median) for numeric feature columns with NULLs
    decisions = {'encodings': {}, 'imputation': {}, 'collinearity': {}}
    cols = [c[1] for c in db.execute(f'PRAGMA table_info("{physical}")').fetchall()]
    numeric_cols = [c for c in cols
                    if c not in ('day', 'location_id') and not c.startswith('target_')
                    and db.execute(f'SELECT typeof("{c}") FROM "{physical}" '
                                   f'WHERE "{c}" IS NOT NULL LIMIT 1').fetchone()
                    and db.execute(f'SELECT typeof("{c}") FROM "{physical}" '
                                   f'WHERE "{c}" IS NOT NULL LIMIT 1').fetchone()[0]
                    in ('integer', 'real')]
    for c in numeric_cols:
        nulls = db.execute(f'SELECT COUNT(*) FROM "{physical}" WHERE "{c}" IS NULL').fetchone()[0]
        if not nulls:
            continue
        total_imputed = 0
        for loc in locations:
            rows = db.execute(f'SELECT rowid, "{c}" FROM "{physical}" '
                              f'WHERE location_id=? ORDER BY day', (loc,)).fetchall()
            vals = [r_[1] for r_ in rows]
            filled, n_imp = fe.impute_measure(vals)
            if n_imp:
                for (rowid, old_v), new_v in zip(rows, filled):
                    if old_v is None:
                        db.execute(f'UPDATE "{physical}" SET "{c}"=? WHERE rowid=?',
                                   (new_v, rowid))
                total_imputed += n_imp
        if total_imputed:
            decisions['imputation'][c] = total_imputed
    db.commit()

    # encodings for string dimensions: one-hot (low card) / frequency (high card)
    text_cols = [c for c in cols
                 if c not in ('day',) and not c.startswith('target_')
                 and (db.execute(f'SELECT typeof("{c}") FROM "{physical}" '
                                 f'WHERE "{c}" IS NOT NULL LIMIT 1').fetchone() or ['null'])[0]
                 == 'text']
    for c in text_cols:
        card = db.execute(f'SELECT COUNT(DISTINCT "{c}") FROM "{physical}"').fetchone()[0]
        rows = db.execute(f'SELECT rowid, "{c}" FROM "{physical}" ORDER BY rowid').fetchall()
        vals = [r_[1] for r_ in rows]
        if card <= fe.ONE_HOT_MAX_CARDINALITY:
            decisions['encodings'][c] = 'one_hot'
            encoded = fe.one_hot(vals, c)
        else:
            decisions['encodings'][c] = 'frequency'
            encoded = {f'{c}_freq': fe.frequency_encode(vals)}
        for ename, evals in encoded.items():
            if ename not in cols and ename not in added:
                db.execute(f'ALTER TABLE "{physical}" ADD COLUMN "{ename}" REAL')
                added.append(ename)
            for (rowid, _), v in zip(rows, evals):
                db.execute(f'UPDATE "{physical}" SET "{ename}"=? WHERE rowid=?', (v, rowid))
    db.commit()

    # collinearity pruning across engineered numeric features (cap 200)
    sample = db.execute(f'SELECT * FROM "{physical}" WHERE location_id=? ORDER BY day',
                        (locations[0],)).fetchall() if locations else []
    if sample:
        colnames = sample[0].keys()
        target_series = [r_[target] for r_ in sample]
        candidates = {c: [r_[c] for r_ in sample] for c in colnames
                      if c in added and c != 'is_holiday'}
        dropped, report = fe.collinear_drops(candidates, target_series)
        decisions['collinearity'] = {'dropped': dropped, 'report': report}
        added = [a for a in added if a not in set(dropped)]
    else:
        decisions['collinearity'] = {'dropped': [], 'report': {}}
    added = added[:fe.MAX_FEATURES]
    decisions['feature_count'] = len(added)

    # enriched feature manifest (minor bump)
    import feature_manifest as fm_mod
    prev = one('SELECT * FROM feature_manifests WHERE session_id=? ORDER BY id DESC LIMIT 1',
               (sid,))
    prev_feats = json.loads(prev['feature_list_json']) if prev else []
    new_feats = prev_feats + [
        {'name': n, 'dtype': 'measure' if n != 'is_holiday' else 'flag',
         'source': 'feature_engineering',
         'transformations': ['temporal' if n != 'is_holiday' else 'holiday_calendar']}
        for n in added if n not in {f_['name'] for f_ in prev_feats}]
    manifest = {
        'workspace_id': prev['workspace_id'] if prev else 'default',
        'session_id': sid,
        'gold_table_name': gold['table_name'],
        'gold_output_hash': gold['output_hash'],
        'generated_at': datetime.utcnow().isoformat() + 'Z',
        'feature_list': new_feats,
        'enrichment_status': 'enriched',
        'manifest_version': None,
    }
    fm = fm_mod.save_feature_manifest(db, manifest)
    db.execute('UPDATE gold_tables SET manifest_version=? WHERE id=?',
               (fm['manifest_version'], gold['id']))
    db.commit()
    log_action('modeler.enriched', 'gold_table', gold['id'],
               {'added_features': added, 'manifest_version': fm['manifest_version'],
                'decisions': {k: v for k, v in decisions.items() if k != 'collinearity'}})
    return jsonify({'session_id': sid, 'physical_table': physical,
                    'added_features': added,
                    'engineering_decisions': decisions,
                    'manifest_version': fm['manifest_version']}), 201


@app.get('/api/modeler/gold/<int:session_id>')
def list_gold_tables(session_id):
    rows = many('SELECT * FROM gold_tables WHERE session_id=? ORDER BY version DESC', (session_id,))
    for r in rows:
        r['split_config'] = json.loads(r.pop('split_config_json') or '{}')
        r['dq_gates'] = json.loads(r.pop('dq_json') or '{}')
    return jsonify(rows)


_GOLD_CACHE: dict = {}
GOLD_CACHE_TTL = 300  # 5 minutes


@app.get('/api/gold/<ws>/<table>')
def gold_query(ws, table):
    """R6S1E1: paginated gold-layer query API (member or embed-token access)."""
    import re as _re
    if not _re.match(r'^(gold_|analytics_)[A-Za-z0-9_.]*$', table):
        return jsonify({'error': 'Not a gold-layer table'}), 404
    # R7S1E2: embed-token access with server-side allowed_origins enforcement
    embed_token = request.args.get('embed_token')
    if embed_token:
        import embed_tokens as et
        payload = et.verify(embed_token, et.workspace_secret(ws))
        if not payload:
            return jsonify({'error': 'Invalid or expired embed token'}), 401
        origin = request.headers.get('Origin', '')
        allowed = payload.get('allowed_origins') or []
        if '*' not in allowed and origin not in allowed:
            return jsonify({'error': f'Origin {origin!r} is not permitted for this '
                                     'embed token'}), 403
    db = get_db()
    exists = db.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",
                        (table,)).fetchone()
    if not exists:
        return jsonify({'error': 'Not a gold-layer table'}), 404
    try:
        page = max(1, int(request.args.get('page', 1)))
        per_page = max(1, min(500, int(request.args.get('per_page', 50))))
    except ValueError:
        return jsonify({'error': 'page/per_page must be integers'}), 400

    fcol, fval = request.args.get('filter_col'), request.args.get('filter_val')
    where, params = '', []
    if fcol:
        cols = {c[1] for c in db.execute(f'PRAGMA table_info("{table}")').fetchall()}
        if fcol not in cols:
            return jsonify({'error': f'unknown filter column {fcol!r}'}), 400
        where = f' WHERE "{fcol}" = ?'
        params = [fval]

    # R20S1E1: apply active row-level-security policy (safe subset WHERE)
    pol = db.execute("SELECT * FROM rls_policies WHERE table_name=? AND status='on' "
                     'ORDER BY id DESC LIMIT 1', (table,)).fetchone()
    if pol:
        where = (where + ' AND ' if where else ' WHERE ') + f"({pol['expression']})"

    # R8S1E2: query results flow through the caching hierarchy's query layer,
    # keyed by the current governed context (gov manifest + semantic versions).
    import cache_hier
    gov_v, sem_v = _uas_context_versions(db, None)
    key_parts = [table, f'p{page}', f'pp{per_page}', f'f{fcol}', f'v{fval}']
    hit = cache_hier.get(db, 'query', ws, key_parts, gov_version=gov_v, sem_version=sem_v)
    if hit is not None:
        return jsonify({**hit, 'cached': True})

    total = db.execute(f'SELECT COUNT(*) FROM "{table}"{where}', params).fetchone()[0]
    rows = [dict(r) for r in db.execute(
        f'SELECT * FROM "{table}"{where} ORDER BY 1 LIMIT ? OFFSET ?',
        params + [per_page, (page - 1) * per_page]).fetchall()]
    payload = {'workspace': ws, 'table': table, 'rows': rows, 'total': total,
               'page': page, 'per_page': per_page, 'cached': False}
    cache_hier.put(db, 'query', ws, key_parts, payload,
                   gov_version=gov_v, sem_version=sem_v, ttl=GOLD_CACHE_TTL)
    return jsonify(payload)


@app.get('/api/feature_manifests/<int:id>')
def get_feature_manifest(id):
    row = one('SELECT * FROM feature_manifests WHERE id=?', (id,))
    if not row:
        return jsonify({'error': 'Not found'}), 404
    row['feature_list'] = json.loads(row.pop('feature_list_json') or '[]')
    return jsonify(row)


@app.get('/api/feature_manifests')
def list_feature_manifests():
    sid = request.args.get('session_id')
    if sid:
        rows = many('SELECT * FROM feature_manifests WHERE session_id=? ORDER BY id DESC', (sid,))
    else:
        rows = many('SELECT * FROM feature_manifests ORDER BY id DESC LIMIT 100')
    for r in rows:
        r['feature_list'] = json.loads(r.pop('feature_list_json') or '[]')
    return jsonify(rows)


@app.patch('/api/feature_manifests/<int:id>')
@app.put('/api/feature_manifests/<int:id>')
def mutate_feature_manifest(id):
    return jsonify({'error': 'Feature manifests are immutable — new gold writes '
                             'create new manifest versions'}), 409


# ─────────────────────────────────────────────────────────
# Routes — Training orchestrator (Sprint 7 / F-027..F-031)
# ─────────────────────────────────────────────────────────
def _execute_training_job(job_id: int, session_id: int, gold: dict, horizon: int, spec=None):
    def _run():
        import training as training_mod
        conn = thread_db()
        try:
            conn.execute("UPDATE training_jobs SET status='running', "
                         "started_at=datetime('now') WHERE id=?", (job_id,))
            conn.commit()
            try:
                card = training_mod.train_session(conn, session_id, gold, horizon=horizon, spec=spec)
            except Exception as exc:
                conn.execute("UPDATE training_jobs SET status='failed', error=?, "
                             "completed_at=datetime('now') WHERE id=?", (str(exc), job_id))
                conn.commit()
                log_action_bg(conn, 'training.failed', 'training_job', job_id,
                              {'error': str(exc)})
                return

            for t in card['trials']:
                conn.execute('INSERT INTO model_trials (job_id, session_id, params_json, '
                             'mape, folds_json) VALUES (?,?,?,?,?)',
                             (job_id, session_id, json.dumps(t['params']), t['mape'],
                              json.dumps(t['folds'])))
            # R5S1E3: explainability over the gold table (PII excluded)
            import explainability as ex
            top_features, concentration = [], {'status': 'FAIL', 'top10_share': 0.0}
            try:
                target_col = card.get('target_column')
                ranked = ex.explain_gold(conn, gold['physical_table'], target_col)
                top_features = ranked
                concentration = ex.concentration_gate(ranked)
            except Exception as exc:
                log.warning('explainability failed: %s', exc)
            card['gates']['concentration_gate'] = concentration

            cur = conn.execute(
                'INSERT INTO model_cards (session_id, job_id, algorithm, gold_table_name, '
                'gold_output_hash, feature_manifest_version, hyperparams_json, metrics_json, '
                'gates_json, target_type, lineage_json, status, trained_at) '
                'VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)',
                (session_id, job_id, card['algorithm'], card['gold_table_name'],
                 card['gold_output_hash'], card['feature_manifest_version'],
                 json.dumps(card['hyperparams']), json.dumps(card['metrics']),
                 json.dumps(card['gates']), card.get('target_type', 'regression'),
                 json.dumps(card.get('lineage') or {}), 'candidate', card['trained_at']))
            card_id = cur.lastrowid
            for rank, tf in enumerate(top_features, 1):
                conn.execute('INSERT INTO gold_model_insights (session_id, model_card_id, '
                             'feature, importance, shap_mean, rank) VALUES (?,?,?,?,?,?)',
                             (session_id, card_id, tf['name'], tf['importance'],
                              tf['shap_mean'], rank))
            conn.commit()
            conn.execute("UPDATE training_jobs SET status='done', model_card_id=?, "
                         "completed_at=datetime('now') WHERE id=?", (card_id, job_id))
            conn.commit()
            log_action_bg(conn, 'training.completed', 'training_job', job_id,
                          {'model_card_id': card_id,
                           'val_mape': card['metrics']['val_mape'],
                           'gates': {k: v['status'] for k, v in card['gates'].items()}})
            send_email(to='workspace-admins@analytiq.dev',
                       subject=f'Training job {job_id} complete',
                       html=(f"<p>Model card {card_id} — validation MAPE "
                             f"{card['metrics']['val_mape']}%.</p>"))
        finally:
            put_db(conn)

    threading.Thread(target=_run, daemon=True).start()


@app.post('/api/training/run')
@limiter.limit('10/minute')
@require_role('admin', 'analyst')
def start_training():
    b = request.get_json() or {}
    sid = b.get('sessionId')
    if not sid:
        return jsonify({'error': 'sessionId required'}), 400
    gold = one("SELECT * FROM gold_tables WHERE session_id=? AND status='written' "
               "ORDER BY version DESC LIMIT 1", (sid,))
    if not gold:
        return jsonify({'error': 'No gold table for this session — run the data modeler first',
                        'remediation': 'POST /api/modeler/generate {mode: execute}'}), 409

    # R4S2E3: applied HOLD-risk features need explicit human confirmation
    held = []
    for cf in many("SELECT * FROM custom_features WHERE session_id=? AND applied=1", (sid,)):
        leak = json.loads(cf['leakage_json'] or '{}')
        if leak.get('action') == 'HOLD' and not one(
                'SELECT 1 FROM leakage_confirmations WHERE session_id=? AND feature=?',
                (sid, cf['name'])):
            held.append(cf['name'])
    if held:
        log_action('training.blocked', 'training_job', None,
                   {'session_id': sid, 'reason': 'leakage_confirmation_required',
                    'features': held})
        return jsonify({'error': 'leakage_confirmation_required',
                        'held_features': held,
                        'remediation': 'POST /api/modeler/leakage/confirm'}), 409

    spec_row = one('SELECT spec_json FROM session_specs WHERE session_id=? '
                   'ORDER BY spec_version DESC LIMIT 1', (sid,))
    horizon = 14
    if spec_row:
        horizon = json.loads(spec_row['spec_json']).get('prediction_horizon') or 14

    job_id = execute('INSERT INTO training_jobs (session_id, gold_table_id, status) '
                     "VALUES (?,?,'queued')", (sid, gold['id']))
    log_action('training.queued', 'training_job', job_id,
               {'session_id': sid, 'gold_table': gold['table_name']})
    _execute_training_job(job_id, sid, gold, horizon, spec=_spec_for_session(sid))
    return jsonify({'jobId': job_id, 'status': 'queued'}), 201


@app.get('/api/training/jobs/<int:id>')
def get_training_job(id):
    row = one('SELECT * FROM training_jobs WHERE id=?', (id,))
    return jsonify(row) if row else (jsonify({'error': 'Not found'}), 404)


@app.get('/api/training/jobs')
def list_training_jobs():
    sid = request.args.get('session_id')
    if sid:
        return jsonify(many('SELECT * FROM training_jobs WHERE session_id=? ORDER BY id DESC', (sid,)))
    return jsonify(many('SELECT * FROM training_jobs ORDER BY id DESC LIMIT 100'))


@app.get('/api/training/jobs/<int:id>/trials')
def list_job_trials(id):
    rows = many('SELECT * FROM model_trials WHERE job_id=? ORDER BY mape ASC', (id,))
    for r in rows:
        r['params'] = json.loads(r.pop('params_json'))
        r['folds'] = json.loads(r.pop('folds_json') or '[]')
    return jsonify(rows)


@app.post('/api/training/jobs/<int:id>/promote')
@require_role('admin', 'analyst')
def promote_training_job(id):
    import training as training_mod
    job = one('SELECT * FROM training_jobs WHERE id=?', (id,))
    if not job:
        return jsonify({'error': 'Not found'}), 404
    if job['status'] != 'done' or not job['model_card_id']:
        return jsonify({'error': f"job is {job['status']} — promotion requires a completed job"}), 409
    card_id = job['model_card_id']
    card = one('SELECT * FROM model_cards WHERE id=?', (card_id,))
    gates = json.loads(card['gates_json'] or '{}')
    metrics = json.loads(card['metrics_json'] or '{}')
    hyper = json.loads(card['hyperparams_json'] or '{}')
    sid = job['session_id']
    gold = one('SELECT * FROM gold_tables WHERE id=?', (job['gold_table_id'],))
    horizon = (metrics.get('holdout_days') or 14)

    def _failed(g):
        return [k for k, v in g.items() if v.get('status') != 'PASS']

    # Repair loop: up to 3 cycles with alternate hyperparameter grids
    repair_cycles = 0
    while _failed(gates) and repair_cycles < 3 and gold:
        grid = training_mod.REPAIR_GRIDS[repair_cycles % len(training_mod.REPAIR_GRIDS)]
        repair_cycles += 1
        try:
            repaired = training_mod.train_session(get_db(), sid, gold,
                                                  horizon=horizon, grid=grid)
        except Exception as exc:
            log_action('model.repair_error', 'model_card', card_id, {'error': str(exc)})
            break
        gates, metrics, hyper = repaired['gates'], repaired['metrics'], repaired['hyperparams']
        db = get_db()
        db.execute('UPDATE model_cards SET hyperparams_json=?, metrics_json=?, gates_json=? '
                   'WHERE id=?', (json.dumps(hyper), json.dumps(metrics),
                                  json.dumps(gates), card_id))
        db.commit()
        log_action('model.repair_attempt', 'model_card', card_id,
                   {'cycle': repair_cycles, 'gates': {k: v['status'] for k, v in gates.items()}})

    if _failed(gates):
        db = get_db()
        db.execute("UPDATE model_cards SET status='rejected' WHERE id=?", (card_id,))
        db.commit()
        log_action('model.repair_exhausted', 'model_card', card_id,
                   {'repair_cycles': repair_cycles, 'human_required': True,
                    'failed_gates': _failed(gates)})
        return jsonify({'status': 'failed', 'model_card_id': card_id,
                        'repair_cycles': repair_cycles, 'human_required': True,
                        'gates': gates,
                        'remediation': 'Adjust features/date range or accept the model manually.'})

    prev_count = one('SELECT COUNT(*) AS c FROM model_registry WHERE session_id=?', (sid,))['c']
    version = prev_count + 1
    model_id = f'stm-lite-s{sid}-v{version}'
    db = get_db()
    db.execute("UPDATE model_registry SET status='archived' WHERE session_id=? AND status='active'",
               (sid,))
    db.commit()
    reg_id = execute('INSERT INTO model_registry (session_id, model_card_id, model_id, version, '
                     "artifact_uri, status) VALUES (?,?,?,?,?,'active')",
                     (sid, card_id, model_id, version, f'local://model_cards/{card_id}'))
    db.execute("UPDATE model_cards SET status='promoted' WHERE id=?", (card_id,))
    db.commit()
    log_action('model.promoted', 'model_registry', reg_id,
               {'model_id': model_id, 'model_card_id': card_id, 'version': version,
                'repair_cycles': repair_cycles})
    return jsonify({'status': 'promoted', 'model_id': model_id, 'registry_id': reg_id,
                    'model_card_id': card_id, 'model_version': version,
                    'repair_cycles': repair_cycles, 'gates': gates})


@app.get('/api/training/result/<int:session_id>')
def training_result(session_id):
    job = one('SELECT * FROM training_jobs WHERE session_id=? ORDER BY id DESC LIMIT 1',
              (session_id,))
    if not job:
        return jsonify({'error': 'No training runs for this session'}), 404
    metrics, links = {}, {'trials': f"/api/training/jobs/{job['id']}/trials"}
    if job['model_card_id']:
        card = one('SELECT * FROM model_cards WHERE id=?', (job['model_card_id'],))
        metrics = json.loads(card['metrics_json'] or '{}')
        links['model_card'] = f"/api/model_cards/{job['model_card_id']}"
        reg = one("SELECT * FROM model_registry WHERE model_card_id=? AND status='active'",
                  (job['model_card_id'],))
        links['registry_model'] = f"/api/registry/models/{reg['id']}" if reg else None
    return jsonify({'session_id': session_id, 'job_id': job['id'], 'status': job['status'],
                    'error': job['error'], 'model_card_id': job['model_card_id'],
                    'metrics': metrics, 'links': links})


CHALLENGER_WIN_PCT = 5.0


@app.post('/api/registry/challenger')
@require_role('admin', 'analyst')
def register_challenger():
    b = request.get_json() or {}
    sid, card_id = b.get('sessionId'), b.get('modelCardId')
    if not sid or not card_id:
        return jsonify({'error': 'sessionId and modelCardId required'}), 400
    card = one('SELECT * FROM model_cards WHERE id=?', (card_id,))
    if not card:
        return jsonify({'error': 'Model card not found'}), 404
    champion = one("SELECT * FROM model_registry WHERE session_id=? AND status='active'",
                   (sid,))
    if not champion:
        return jsonify({'error': 'No active champion for this session — promote first'}), 409
    version = one('SELECT COUNT(*) c FROM model_registry WHERE session_id=?', (sid,))['c'] + 1
    reg_id = execute('INSERT INTO model_registry (session_id, model_card_id, model_id, '
                     "version, artifact_uri, status) VALUES (?,?,?,?,?,'challenger')",
                     (sid, card_id, f'stm-lite-s{sid}-v{version}', version,
                      f'local://model_cards/{card_id}'))
    log_action('model.challenger_registered', 'model_registry', reg_id,
               {'session_id': sid, 'model_card_id': card_id,
                'champion_registry_id': champion['id']})
    return jsonify({'registry_id': reg_id, 'status': 'challenger',
                    'champion_registry_id': champion['id']}), 201


@app.post('/api/registry/challenger/<int:id>/evaluate')
@require_role('admin', 'analyst')
def evaluate_challenger(id):
    ch = one("SELECT * FROM model_registry WHERE id=? AND status='challenger'", (id,))
    if not ch:
        return jsonify({'error': 'Challenger not found'}), 404
    champion = one("SELECT * FROM model_registry WHERE session_id=? AND status='active'",
                   (ch['session_id'],))
    if not champion:
        return jsonify({'error': 'No active champion to compare against'}), 409

    def _mape_of(card_id):
        card = one('SELECT metrics_json FROM model_cards WHERE id=?', (card_id,))
        return (json.loads(card['metrics_json'] or '{}') or {}).get('val_mape')

    champ_mape = _mape_of(champion['model_card_id'])
    chall_mape = _mape_of(ch['model_card_id'])
    if champ_mape is None or chall_mape is None or champ_mape == 0:
        return jsonify({'error': 'Metrics unavailable for comparison'}), 422
    improvement = (champ_mape - chall_mape) / champ_mape * 100

    if improvement > CHALLENGER_WIN_PCT:
        db = get_db()
        db.execute("UPDATE model_registry SET status='archived' WHERE id=?", (champion['id'],))
        db.execute("UPDATE model_registry SET status='active' WHERE id=?", (id,))
        db.execute("UPDATE model_cards SET status='promoted' WHERE id=?",
                   (ch['model_card_id'],))
        db.commit()
        log_action('model.challenger_promoted', 'model_registry', id,
                   {'improvement_pct': round(improvement, 2),
                    'previous_champion': champion['id']})
        return jsonify({'outcome': 'challenger_promoted',
                        'improvement_pct': round(improvement, 2),
                        'new_champion_registry_id': id})
    log_action('model.challenger_retained', 'model_registry', id,
               {'improvement_pct': round(improvement, 2)})
    return jsonify({'outcome': 'champion_retained',
                    'improvement_pct': round(improvement, 2)})


@app.get('/api/models/overview')
def models_overview():
    """R33S1E1: models pillar overview — 6 KPIs + a typed row per model.
    Registry entries surface as CHAMPION / CHALLENGER / ARCHIVED; training
    jobs without a registry entry surface as TRAINING or RUN FAILED."""
    regs = many('SELECT * FROM model_registry ORDER BY id DESC LIMIT 100')
    jobs = many('SELECT * FROM training_jobs ORDER BY id DESC LIMIT 200')
    reg_jobs = set()
    rows = []

    def _accuracy(card):
        # training metrics store MAPE values already in percent (e.g. 5.24)
        m = json.loads((card or {}).get('metrics_json') or '{}')
        mape = m.get('val_mape', m.get('test_mape', m.get('mape')))
        if mape is not None:
            return {'label': 'MAPE', 'value': f"{round(mape, 1)}%"}
        if m.get('auc') is not None:
            return {'label': 'AUC', 'value': f"{round(m['auc'], 2)}"}
        if m.get('mae') is not None:
            return {'label': 'MAE', 'value': f"{round(m['mae'], 1)}"}
        return {'label': 'MAPE', 'value': None}

    for r_ in regs:
        card = one('SELECT * FROM model_cards WHERE id=?', (r_['model_card_id'],)) \
            if r_['model_card_id'] else None
        sess = one('SELECT * FROM sessions WHERE id=?', (r_['session_id'],)) \
            if r_['session_id'] else None
        job = one('SELECT * FROM training_jobs WHERE model_card_id=?',
                  (r_['model_card_id'],)) if r_['model_card_id'] else None
        if job:
            reg_jobs.add(job['id'])
        status = {'active': 'CHAMPION', 'challenger': 'CHALLENGER',
                  'archived': 'ARCHIVED'}.get(r_['status'], 'CHAMPION')
        rows.append({
            'registry_id': r_['id'], 'model_id': r_['model_id'],
            'version': r_['version'], 'status': status,
            'purpose': f"{sess['metric']} forecast" if sess else 'Forecast model',
            'algorithm': (card or {}).get('algorithm'),
            'grain': sess['grain'] if sess else None,
            'last_trained': (job or {}).get('completed_at') or r_['created_at'],
            'accuracy': _accuracy(card),
            'session_id': r_['session_id'], 'card_id': r_['model_card_id'],
            'job_id': (job or {}).get('id'),
        })
    for j in jobs:
        if j['id'] in reg_jobs or j['status'] == 'done':
            continue
        sess = one('SELECT * FROM sessions WHERE id=?', (j['session_id'],)) \
            if j['session_id'] else None
        rows.append({
            'registry_id': None, 'model_id': f"job_{j['id']}", 'version': None,
            'status': 'RUN FAILED' if j['status'] == 'failed' else 'TRAINING',
            'purpose': f"{sess['metric']} forecast" if sess else 'Training run',
            'algorithm': None, 'grain': sess['grain'] if sess else None,
            'last_trained': j['completed_at'] or j['created_at'],
            'accuracy': {'label': 'MAPE', 'value': None},
            'session_id': j['session_id'], 'card_id': None, 'job_id': j['id'],
            'error': j['error'],
        })

    kpis = {
        'promoted': one("SELECT COUNT(*) AS n FROM model_registry WHERE status='active'")['n'],
        'runs_30d': one("SELECT COUNT(*) AS n FROM training_jobs "
                        "WHERE created_at >= datetime('now', '-30 days')")['n'],
        'failed': one("SELECT COUNT(*) AS n FROM training_jobs WHERE status='failed'")['n'],
        'retrain_due': one("SELECT COUNT(*) AS n FROM alerts WHERE type='drift'")['n'],
        'champ_challenger': one("SELECT COUNT(*) AS n FROM model_registry "
                                "WHERE status='challenger'")['n'],
        'prediction_tables': one('SELECT COUNT(*) AS n FROM gold_tables')['n'],
    }
    return jsonify({'kpis': kpis, 'models': rows})


@app.get('/api/models/retrain_queue')
def models_retrain_queue():
    """R33S1E4: retrain center — live drift check per champion (real
    model_monitor run), failed training jobs, and healthy rows."""
    import model_monitor as mm
    rows = []
    for r_ in many("SELECT * FROM model_registry WHERE status='active' "
                   'ORDER BY id DESC LIMIT 50'):
        sid = r_['session_id']
        check = mm.check(get_db(), sid) if sid else None
        triggers = (check or {}).get('triggers') or []
        if triggers:
            det = []
            imp = (check or {}).get('importance_drift') or {}
            inp = (check or {}).get('input_drift') or {}
            if imp.get('drifted'):
                det.append(f"importance drift · tau {imp.get('kendall_tau')}")
            if inp.get('drifted'):
                det.append(f"input drift · PSI {inp.get('psi')}")
            rows.append({'kind': 'drift', 'model_id': r_['model_id'],
                         'session_id': sid,
                         'reason': 'drift-triggered · ' + ' · '.join(det),
                         'action': 'retrain'})
        else:
            rows.append({'kind': 'healthy', 'model_id': r_['model_id'],
                         'session_id': sid,
                         'reason': 'monitored · no drift on the latest check',
                         'action': 'retrain'})
    for j in many("SELECT * FROM training_jobs WHERE status='failed' "
                  'ORDER BY id DESC LIMIT 50'):
        sess = one('SELECT * FROM sessions WHERE id=?', (j['session_id'],)) \
            if j['session_id'] else None
        rows.append({'kind': 'failed', 'model_id': f"job_{j['id']}",
                     'session_id': j['session_id'], 'job_id': j['id'],
                     'reason': f"failed · {j['error'] or 'training error'}",
                     'action': 'logs'})
    counts = {'all': len(rows),
              'drift': sum(1 for r_ in rows if r_['kind'] == 'drift'),
              'failed': sum(1 for r_ in rows if r_['kind'] == 'failed'),
              'scheduled': 0,   # model-level schedules ship with R36S1
              'healthy': sum(1 for r_ in rows if r_['kind'] == 'healthy')}
    return jsonify({'rows': rows, 'counts': counts})


@app.get('/api/registry/models')
def list_registry_models():
    sid = request.args.get('session_id')
    if sid:
        return jsonify(many('SELECT * FROM model_registry WHERE session_id=? ORDER BY version DESC',
                            (sid,)))
    return jsonify(many('SELECT * FROM model_registry ORDER BY id DESC LIMIT 100'))


@app.get('/api/registry/models/<int:id>')
def get_registry_model(id):
    row = one('SELECT * FROM model_registry WHERE id=?', (id,))
    return jsonify(row) if row else (jsonify({'error': 'Not found'}), 404)


@app.post('/api/registry/models/<int:id>/archive')
@require_role('admin', 'analyst')
def archive_registry_model(id):
    if not one('SELECT id FROM model_registry WHERE id=?', (id,)):
        return jsonify({'error': 'Not found'}), 404
    db = get_db()
    db.execute("UPDATE model_registry SET status='archived' WHERE id=?", (id,))
    db.commit()
    log_action('model.archived', 'model_registry', id)
    return jsonify(one('SELECT * FROM model_registry WHERE id=?', (id,)))


@app.get('/api/model_cards/<int:id>')
def get_model_card(id):
    row = one('SELECT * FROM model_cards WHERE id=?', (id,))
    if not row:
        return jsonify({'error': 'Not found'}), 404
    row['hyperparams'] = json.loads(row.pop('hyperparams_json') or '{}')
    row['metrics'] = json.loads(row.pop('metrics_json') or '{}')
    row['gates'] = json.loads(row.pop('gates_json') or '{}')
    row['top_features'] = many('SELECT feature AS name, importance, shap_mean, rank '
                               'FROM gold_model_insights WHERE model_card_id=? '
                               'ORDER BY rank', (id,))
    row['lineage'] = json.loads(row.pop('lineage_json') or '{}')
    # R33S1E3: registry identity + linked artifacts for the model-card page
    row['registry'] = one('SELECT id, model_id, version, status FROM model_registry '
                          'WHERE model_card_id=? ORDER BY id DESC LIMIT 1', (id,))
    row['linked_artifacts'] = many(
        'SELECT a.id, a.title FROM artifacts a JOIN pipeline_runs pr '
        'ON a.pipeline_run_id = pr.id WHERE pr.session_id=? ORDER BY a.id DESC LIMIT 10',
        (row['session_id'],)) if row.get('session_id') else []
    return jsonify(row)


@app.get('/api/models/<int:card_id>/insights')
def model_insights(card_id):
    return jsonify(many('SELECT * FROM gold_model_insights WHERE model_card_id=? '
                        'ORDER BY rank', (card_id,)))


@app.get('/api/model_cards')
def list_model_cards():
    sid = request.args.get('session_id')
    rows = (many('SELECT * FROM model_cards WHERE session_id=? ORDER BY id DESC', (sid,))
            if sid else many('SELECT * FROM model_cards ORDER BY id DESC LIMIT 100'))
    for row in rows:
        row['hyperparams'] = json.loads(row.pop('hyperparams_json') or '{}')
        row['metrics'] = json.loads(row.pop('metrics_json') or '{}')
        row['gates'] = json.loads(row.pop('gates_json') or '{}')
    return jsonify(rows)


@app.post('/api/splits/preview')
def splits_preview():
    import splits as splits_mod
    b = request.get_json() or {}
    dr = b.get('date_range') or {'start': '2023-01-01', 'end': '2023-12-31'}
    row_count = int(b.get('row_count', 0))
    horizon = b.get('horizon')
    cfg = splits_mod.compute_split_config(dr, row_count=row_count, horizon=horizon)
    return jsonify({'split_config': cfg,
                    'validation': splits_mod.validate_split(cfg, row_count, horizon)})


@app.post('/api/sessions/<int:id>/save_artifact')
@require_role('admin', 'analyst')
def save_artifact_from_session(id):
    import hashlib
    import artifact_gen as ag
    if not one('SELECT id FROM sessions WHERE id=?', (id,)):
        return jsonify({'error': 'Session not found'}), 404
    run = one("SELECT * FROM pipeline_runs WHERE session_id=? AND status='done' "
              'ORDER BY id DESC LIMIT 1', (id,))
    if not run:
        return jsonify({'error': 'No completed pipeline run for this session — '
                                 'run the pipeline first'}), 409
    b = request.get_json() or {}
    sess = one('SELECT * FROM sessions WHERE id=?', (id,))
    title = b.get('title') or f"{sess['metric']} Forecast"
    aid = execute('INSERT INTO artifacts (title,type,mape,owner,dq_status,pipeline_run_id,is_sandbox) '
                  'VALUES (?,?,?,?,?,?,?)',
                  (title, b.get('type', 'Predictive'), run['mape'],
                   b.get('owner', 'analyst@acme.com'), 'pass', run['id'],
                   1 if sess.get('is_sandbox') else 0))

    # auto-render the self-contained file (F-040: record → store)
    art = one('SELECT * FROM artifacts WHERE id=?', (aid,))
    rows = many('SELECT * FROM chart_data WHERE pipeline_run_id=? ORDER BY day_index', (run['id'],))
    file_info = None
    if rows:
        html = ag.generate_artifact_html(art, rows, compute_kpis(rows))
        # R11S2E3: assembly goes through the repair loop; attempts retained
        _rep_attempts = []
        html, _rep_cycles, validation = ag.validate_and_repair(html, max_cycles=2,
                                                               attempt_log=_rep_attempts)
        _persist_repair_attempts(get_db(), 'artifact_render', aid, run['id'],
                                 _rep_attempts, validation['status'] == 'PASS')
        fid, sha, _uri = _store_artifact_file(aid, 1, html, validation)
        file_info = {'file_id': fid, 'version': 1, 'sha256': sha,
                     'size_bytes': validation['size_bytes'],
                     'validation_status': validation['status']}
    # R11S1E2: propagate stage confidences onto the assembled artifact
    import confidence as cf
    prop = cf.propagate(get_db(), id, run['id'])
    execute('UPDATE artifacts SET confidence=?, confidence_json=? WHERE id=?',
            (prop['confidence'], json.dumps(prop), aid))
    art = one('SELECT * FROM artifacts WHERE id=?', (aid,))
    # R16S2E4: initialize the sections layout from the run's viz specs
    ns_l = _uas_ns(sess)
    import uas as _uas
    spec_node = _uas.latest_by_logical(get_db(), f"{ns_l}:vega_lite_specs:s{id}")
    panels = (json.loads(spec_node['payload_json']).get('specs') if spec_node else None) or \
             [{'panel': 'timeseries_ci', 'mark': 'line'}, {'panel': 'forecast', 'mark': 'area'}]
    layout = {'sections': [{'id': p['panel'],
                            'title': p['panel'].replace('_', ' ').title(),
                            'mark': p.get('mark', 'line'), 'top_n': None, 'position': i}
                           for i, p in enumerate(panels)]}
    execute('UPDATE artifacts SET layout_json=? WHERE id=?', (json.dumps(layout), aid))
    _register_artifact_html_uas(get_db(), aid, run['id'], title, file_info)
    import search as search_mod
    if not sess.get('is_sandbox'):      # R9S2E6: sandbox stays out of prod search
        search_mod.index_artifact(get_db(), art)
    if not sess.get('is_sandbox'):      # R10S1E2: incremental KG ingest
        import knowledge_graph as kg
        kg.ingest_artifact(get_db(), aid)
        # R12S1E1: post-assembly opportunity evaluation (same insight engine)
        import opportunity
        _ins, _err = _compute_artifact_insights(aid)
        opportunity.evaluate(get_db(), aid, _ins or [])
    import intent_history as ih
    ih.record(get_db(), getattr(g, 'user_email', None) or 'default', 'artifact',
              ref_id=aid, session_id=id, detail={'metric': sess.get('metric'), 'title': title})
    log_action('artifact.saved', 'artifact', aid,
               {'session_id': id, 'pipeline_run_id': run['id'], 'title': title,
                'rendered': bool(file_info)})
    return jsonify({**art, 'file': file_info}), 201


# ─────────────────────────────────────────────────────────
# Routes — Pipeline
# ─────────────────────────────────────────────────────────
@app.post('/api/pipeline/run')
@limiter.limit('5/minute')
@plan_gate('pro')
def start_pipeline():
    b   = request.get_json() or {}
    sid = b.get('sessionId')
    if not sid:
        return jsonify({'error': 'sessionId required'}), 400
    # Sprint 3 / F-014: block training while low-confidence definitions are unreviewed
    sess = one('SELECT * FROM sessions WHERE id=?', (sid,))
    if sess and sess.get('run_id'):
        pending = many(
            "SELECT id, name, type, confidence FROM semantic_definitions "
            "WHERE run_id=? AND status='pending' AND confidence < 0.70",
            (sess['run_id'],))
        if pending:
            log_action('pipeline.blocked', 'pipeline_run', None,
                       {'session_id': sid, 'reason': 'human_review_required',
                        'pending_ids': [p['id'] for p in pending]})
            return jsonify({
                'error': 'human_review_required',
                'message': 'Low-confidence semantic definitions must be reviewed before training.',
                'pending_reviews': pending,
                'review_url': f"/api/reviews/{sess['run_id']}",
            }), 409
    # R3S1E4: block on active data-contract violations
    if sess and sess.get('connection_id'):
        mrow = _manifest_row(sess['connection_id'])
        if mrow:
            mdoc = json.loads(mrow['manifest_json'])
            violations = {t['name']: t['contract_violations'] for t in mdoc.get('tables', [])
                          if t.get('contract_violations')}
            if violations:
                log_action('pipeline.blocked', 'pipeline_run', None,
                           {'session_id': sid, 'reason': 'contract_violation',
                            'tables': sorted(violations)})
                return jsonify({'error': 'contract_violation',
                                'message': 'Data contracts are violated — resolve before training.',
                                'violations': violations}), 409
    lid = execute('INSERT INTO pipeline_runs (session_id,status,current_step,log_entries) VALUES (?,?,?,?)',
                  (sid, 'running', 0, '[]'))
    log_action('pipeline.started', 'pipeline_run', lid, {'session_id': sid})
    simulate_pipeline(lid)
    return jsonify({'runId': lid}), 201

@app.get('/api/pipeline/stream/<int:run_id>')
def pipeline_stream(run_id):
    q = Queue()
    with _pipe_lock:
        _pipe_clients.setdefault(run_id, []).append(q)

    def generate():
        row = one('SELECT * FROM pipeline_runs WHERE id=?', (run_id,))
        if row:
            logs = json.loads(row['log_entries'] or '[]')
            yield f"data: {json.dumps({'step': row['current_step'], 'status': row['status'], 'log': logs})}\n\n"
        if row and row['status'] == 'done':
            return
        try:
            while True:
                try:
                    data = q.get(timeout=15)
                    if data is None:
                        return
                    yield f"data: {json.dumps(data)}\n\n"
                    if data.get('status') == 'done':
                        return
                except Empty:
                    yield ': ping\n\n'
        finally:
            with _pipe_lock:
                lst = _pipe_clients.get(run_id, [])
                if q in lst:
                    lst.remove(q)

    return sse_response(generate)

@app.get('/api/pipeline/<int:run_id>/steps')
def pipeline_steps(run_id):
    rows = many('SELECT * FROM pipeline_steps WHERE run_id=? ORDER BY step', (run_id,))
    for r in rows:
        r['input_schema'] = json.loads(r.pop('input_schema_json') or '[]')
        r['output_schema'] = json.loads(r.pop('output_schema_json') or '[]')
    return jsonify(rows)


@app.post('/api/pipeline/<int:run_id>/steps/<int:step>/flag')
def flag_pipeline_step(run_id, step):
    row = one('SELECT id FROM pipeline_steps WHERE run_id=? AND step=?', (run_id, step))
    if not row:
        return jsonify({'error': 'Step not found'}), 404
    reason = (request.get_json() or {}).get('reason')
    db = get_db()
    db.execute('UPDATE pipeline_steps SET flagged=1, flag_reason=? WHERE id=?',
               (reason, row['id']))
    db.commit()
    log_action('pipeline.step_flagged', 'pipeline_run', run_id, {'step': step, 'reason': reason})
    return jsonify({'run_id': run_id, 'step': step, 'flagged': True})


@app.get('/api/sessions/<int:id>/suggestions')
def session_suggestions(id):
    """R4S1E3: 3–5 follow-up questions from the current spec + sibling sessions."""
    if not one('SELECT id FROM sessions WHERE id=?', (id,)):
        return jsonify({'error': 'Session not found'}), 404
    spec = _spec_for_session(id) or {}
    metric = spec.get('target_metric') or one(
        'SELECT metric FROM sessions WHERE id=?', (id,))['metric'] or 'Net Revenue'
    grain = spec.get('grain') or 'Location · Day'
    suggestions = [
        {'question': f'Why did {metric} change most recently?', 'intent': 'diagnostic'},
        {'question': f'Predict {metric} for the next 28 days by '
                     f"{grain.split('·')[0].strip().lower()}", 'intent': 'predictive'},
        {'question': f'Summarize {metric} performance over the training window',
         'intent': 'descriptive'},
    ]
    # patterns from sibling sessions (other confirmed metrics in the workspace)
    for s in many('SELECT DISTINCT metric FROM sessions WHERE id != ? '
                  'AND metric IS NOT NULL LIMIT 2', (id,)):
        if s['metric'] and s['metric'] != metric:
            suggestions.append({'question': f"Forecast {s['metric']} at the same grain",
                                'intent': 'predictive'})
    return jsonify(suggestions[:5])


@app.get('/api/pipeline/<int:id>')
def get_pipeline(id):
    row = one('SELECT * FROM pipeline_runs WHERE id=?', (id,))
    if not row:
        return jsonify({'error': 'Not found'}), 404
    row['log_entries'] = json.loads(row['log_entries'] or '[]')
    return jsonify(row)


# ─────────────────────────────────────────────────────────
# Routes — Artifacts
# ─────────────────────────────────────────────────────────
@app.get('/api/artifacts')
def list_artifacts():
    q         = request.args.get('q', '').strip()
    art_type  = request.args.get('type', '').strip()
    dq_status = request.args.get('dq_status', '').strip()
    try:
        page = max(1, int(request.args.get('page', 1)))
    except ValueError:
        return jsonify({'error': 'page must be an integer'}), 400
    try:
        per_page = max(1, min(100, int(request.args.get('per_page', 20))))
    except ValueError:
        return jsonify({'error': 'per_page must be an integer'}), 400

    where, params = [], []
    # R9S2E6: sandbox artifacts are excluded from production surfaces
    if request.args.get('sandbox'):
        where.append('a.is_sandbox = 1')
    else:
        where.append('a.is_sandbox = 0')
    if q:
        where.append('a.title LIKE ?')
        params.append(f'%{q}%')
    if art_type:
        where.append('a.type = ?')
        params.append(art_type)
    if dq_status:
        where.append('a.dq_status = ?')
        params.append(dq_status)
    if request.args.get('favorite'):
        where.append('a.favorite = 1')
    tag = request.args.get('tag')
    if tag:
        where.append("a.tags_json LIKE ?")
        params.append(f'%"{tag}"%')

    where_sql = ('WHERE ' + ' AND '.join(where)) if where else ''

    total = one(
        f'SELECT COUNT(*) AS cnt FROM artifacts a {where_sql}', params
    )['cnt']

    offset = (page - 1) * per_page
    items = many(
        f'SELECT a.*, COUNT(s.id) AS share_count, '
        f'sched.cron_expr AS schedule_cron, sched.enabled AS schedule_enabled '
        f'FROM artifacts a LEFT JOIN artifact_shares s ON s.artifact_id=a.id '
        f'LEFT JOIN artifact_schedules sched ON sched.artifact_id=a.id '
        f'{where_sql} '
        f'GROUP BY a.id, sched.cron_expr, sched.enabled ORDER BY a.created_at DESC '
        f'LIMIT ? OFFSET ?',
        params + [per_page, offset]
    )

    for it in items:
        it['tags'] = json.loads(it.pop('tags_json', None) or '[]')
        if it.get('pipeline_run_id'):
            it['thumbnail_url'] = f"/api/artifacts/{it['id']}/thumbnail"
    return jsonify({'items': items, 'total': total, 'page': page, 'per_page': per_page})

@app.get('/api/artifacts/<int:id>')
def get_artifact(id):
    denied = acl_allows('artifact', id, 'read')
    if denied:
        return denied
    row = one('SELECT * FROM artifacts WHERE id=?', (id,))
    if row:
        record_activity(id, 'viewed')
        row['tags'] = json.loads(row.pop('tags_json', None) or '[]')
    import confidence as cf
    if row.get('confidence') is not None:
        row['confidence_level'] = 'low' if row['confidence'] < cf.LOW_THRESHOLD else 'normal'
    return jsonify(row) if row else (jsonify({'error': 'Not found'}), 404)

@app.post('/api/artifacts')
@require_role('admin', 'analyst')
def create_artifact():
    b   = request.get_json() or {}
    lid = execute(
        'INSERT INTO artifacts (title,type,mape,owner,dq_status,pipeline_run_id) VALUES (?,?,?,?,?,?)',
        (b.get('title'), b.get('type', 'Predictive'), b.get('mape'),
         b.get('owner', 'analyst@acme.com'), b.get('dq_status', 'pass'), b.get('pipeline_run_id')),
    )
    log_action('artifact.created', 'artifact', lid, {'title': b.get('title'), 'type': b.get('type', 'Predictive')})
    import search as search_mod
    art = one('SELECT * FROM artifacts WHERE id=?', (lid,))
    search_mod.index_artifact(get_db(), art)
    return jsonify(art), 201

@app.patch('/api/artifacts/<int:id>')
@require_role('admin', 'analyst')
def rename_artifact(id):
    """R30S1E4-US1: the detail page's editable title persists renames
    (audited + search re-indexed). Title is the only mutable field here."""
    art = one('SELECT * FROM artifacts WHERE id=?', (id,))
    if not art:
        return jsonify({'error': 'Not found'}), 404
    b = request.get_json() or {}
    title = (b.get('title') or '').strip()
    if not title:
        return jsonify({'error': 'title is required'}), 400
    execute('UPDATE artifacts SET title=? WHERE id=?', (title, id))
    row = one('SELECT * FROM artifacts WHERE id=?', (id,))
    import search as search_mod
    search_mod.index_artifact(get_db(), row)
    log_action('artifact.renamed', 'artifact', id,
               {'from': art['title'], 'to': title})
    row['tags'] = json.loads(row.pop('tags_json', None) or '[]')
    return jsonify(row)


@app.delete('/api/artifacts/<int:id>')
@require_role('admin', 'analyst')
def delete_artifact(id):
    denied = acl_allows('artifact', id, 'write')
    if denied:
        return denied
    for child in ('artifact_shares', 'artifact_activity', 'artifact_annotations',
                  'metric_subscriptions', 'share_links', 'artifact_files'):
        execute(f'DELETE FROM {child} WHERE artifact_id=?', (id,))
    execute('DELETE FROM artifacts WHERE id=?', (id,))
    import search as search_mod
    search_mod.remove_artifact(get_db(), id)
    log_action('artifact.deleted', 'artifact', id)
    return '', 204

@app.get('/api/artifacts/<int:id>/chart')
def get_chart(id):
    art = one('SELECT * FROM artifacts WHERE id=?', (id,))
    if not art:
        return jsonify({'error': 'Not found'}), 404
    if not art['pipeline_run_id']:
        return jsonify({'rows': [], 'kpis': {'avgActual': 0, 'mape': 0, 'forecast14Avg': 0}})
    rows = many('SELECT * FROM chart_data WHERE pipeline_run_id=? ORDER BY day_index', (art['pipeline_run_id'],))
    return jsonify({'rows': rows, 'kpis': compute_kpis(rows)})

@app.get('/api/artifacts/<int:id>/export')
def export_artifact(id):
    fmt = request.args.get('format', 'json')
    if fmt not in ('csv', 'json'):
        return jsonify({'error': 'format must be csv or json'}), 400

    art = one('SELECT * FROM artifacts WHERE id=?', (id,))
    if not art:
        return jsonify({'error': 'Not found'}), 404

    rows = many(
        'SELECT day_index, date, actual, predicted, ci_low, ci_high, is_forecast '
        'FROM chart_data WHERE pipeline_run_id=? ORDER BY day_index',
        (art['pipeline_run_id'],),
    ) if art['pipeline_run_id'] else []

    if fmt == 'csv':
        import io, csv
        buf = io.StringIO()
        writer = csv.writer(buf)
        cols = ['day_index', 'date', 'actual', 'predicted', 'ci_low', 'ci_high', 'is_forecast']
        writer.writerow(cols)
        for r in rows:
            writer.writerow([r[c] for c in cols])
        resp = Response(buf.getvalue(), mimetype='text/csv')
        resp.headers['Content-Disposition'] = f'attachment; filename=artifact-{id}.csv'
        return resp

    kpis = compute_kpis(rows) if rows else {'avgActual': 0, 'mape': 0, 'forecast14Avg': 0}
    payload = {**art, 'kpis': kpis, 'chart_data': rows}
    for k in ('created_at',):
        if k in payload and payload[k] is not None:
            payload[k] = str(payload[k])
    resp = Response(json.dumps(payload, default=str), mimetype='application/json')
    resp.headers['Content-Disposition'] = f'attachment; filename=artifact-{id}.json'
    return resp


def _store_artifact_file(artifact_id, version, html, validation):
    """R1S2E5: write-through the storage interface; keep DB copy for compat."""
    import hashlib as _h
    import storage as storage_mod
    key = f'artifacts/{artifact_id}/v{version}.html'
    stored = storage_mod.put(key, html)
    sha = _h.sha256(html.encode()).hexdigest()
    fid = execute('INSERT INTO artifact_files (artifact_id, version, html, size_bytes, '
                  'sha256, storage_uri, validator_json) VALUES (?,?,?,?,?,?,?)',
                  (artifact_id, version, html, validation['size_bytes'], sha,
                   stored['uri'], json.dumps(validation)))
    return fid, sha, stored['uri']


@app.post('/api/artifacts/<int:id>/render')
@limiter.limit('20/minute')
@require_role('admin', 'analyst')
def render_artifact(id):
    import hashlib
    import artifact_gen as ag
    art = one('SELECT * FROM artifacts WHERE id=?', (id,))
    if not art:
        return jsonify({'error': 'Not found'}), 404
    if not art['pipeline_run_id']:
        return jsonify({'error': 'Artifact has no pipeline run — nothing to render'}), 409
    rows = many('SELECT * FROM chart_data WHERE pipeline_run_id=? ORDER BY day_index',
                (art['pipeline_run_id'],))
    if not rows:
        return jsonify({'error': 'No chart data for this artifact yet'}), 409

    kpis = compute_kpis(rows)
    card = one('SELECT mc.* FROM model_cards mc JOIN pipeline_runs pr '
               'ON pr.session_id = mc.session_id WHERE pr.id=? '
               'ORDER BY mc.id DESC LIMIT 1', (art['pipeline_run_id'],))
    trials, top_features = [], []
    if card:
        card = dict(card)
        card['lineage'] = json.loads(card.get('lineage_json') or '{}')
        top_features = many('SELECT feature AS name, importance, shap_mean FROM '
                            'gold_model_insights WHERE model_card_id=? ORDER BY rank',
                            (card['id'],))
        trial_rows = many('SELECT * FROM model_trials WHERE job_id=? ORDER BY mape',
                          (card['job_id'],))
        trials = [{'params': json.loads(t['params_json'] or '{}'), 'mape': t['mape']}
                  for t in trial_rows]
    annotations = many('SELECT * FROM artifact_annotations WHERE artifact_id=? ORDER BY id',
                       (id,))
    versions_hist = many('SELECT version, created_at FROM artifact_files '
                         'WHERE artifact_id=? ORDER BY version DESC', (id,))
    branding = one("SELECT * FROM workspace_branding WHERE workspace_id='default'")
    html = ag.generate_artifact_html(art, rows, kpis, model_card=card,
                                     trials=trials, top_features=top_features,
                                     annotations=annotations, versions=versions_hist,
                                     branding=branding)
    # R6S2E3: deterministic repair loop (max 2 cycles) before failing hard
    _attempts = []
    html, repair_cycles, validation = ag.validate_and_repair(html, max_cycles=2,
                                                             attempt_log=_attempts)
    _persist_repair_attempts(get_db(), 'artifact_render', id,
                             art.get('pipeline_run_id'), _attempts,
                             validation['status'] == 'PASS')
    if validation['status'] != 'PASS':
        log_action('artifact.render_failed', 'artifact', id,
                   {'repair_cycles': repair_cycles,
                    'checks': [c for c in validation['checks'] if not c['ok']]})
        return jsonify({'error': 'Artifact failed validation after repair',
                        'repair_cycles': repair_cycles, 'validation': validation}), 422

    prev = one('SELECT MAX(version) AS v FROM artifact_files WHERE artifact_id=?', (id,))
    version = (prev['v'] or 0) + 1
    fid, sha, _uri = _store_artifact_file(id, version, html, validation)
    log_action('artifact.rendered', 'artifact', id,
               {'file_id': fid, 'version': version, 'sha256': sha,
                'size_bytes': validation['size_bytes']})
    return jsonify({'file_id': fid, 'version': version, 'sha256': sha,
                    'size_bytes': validation['size_bytes'],
                    'artifact_uri': f'local://artifact_files/{fid}',
                    'validation': validation}), 201


@app.get('/api/artifacts/<int:id>/html')
def get_artifact_html(id):
    # R30S3E5: optional ?version= fetches a specific history entry (Compare)
    v = request.args.get('version')
    if v:
        row = one('SELECT * FROM artifact_files WHERE artifact_id=? AND version=?', (id, v))
    else:
        row = one('SELECT * FROM artifact_files WHERE artifact_id=? ORDER BY version DESC LIMIT 1',
                  (id,))
    if not row:
        return jsonify({'error': 'Artifact not rendered yet'}), 404
    return Response(row['html'], mimetype='text/html')


@app.get('/api/artifacts/<int:id>/versions')
def list_artifact_versions(id):
    """R30S3E5-US1: version history for the topbar Versions panel."""
    if not one('SELECT id FROM artifacts WHERE id=?', (id,)):
        return jsonify({'error': 'Not found'}), 404
    return jsonify(many('SELECT version, created_at FROM artifact_files '
                        'WHERE artifact_id=? ORDER BY version DESC', (id,)))


@app.post('/api/artifacts/<int:id>/versions/<int:v>/restore')
@require_role('admin', 'analyst')
def restore_artifact_version(id, v):
    """R30S3E5-US1: append-only restore — vN's html is re-stored as a NEW
    top version (history is never rewritten); audited."""
    old = one('SELECT * FROM artifact_files WHERE artifact_id=? AND version=?', (id, v))
    if not old:
        return jsonify({'error': 'Version not found'}), 404
    top = one('SELECT MAX(version) AS m FROM artifact_files WHERE artifact_id=?', (id,))['m']
    _store_artifact_file(id, top + 1, old['html'],
                         json.loads(old['validator_json'] or '{}'))
    log_action('artifact.version_restored', 'artifact', id,
               {'restored_from': v, 'new_version': top + 1})
    return jsonify({'version': top + 1, 'restored_from': v})


@app.put('/api/branding')
@require_role('admin')
def put_branding():
    import re as _re
    b = request.get_json() or {}
    color = b.get('primary_color')
    if color and not _re.match(r'^#[0-9a-fA-F]{6}$', color):
        return jsonify({'error': 'primary_color must be a #rrggbb hex value'}), 400
    db = get_db()
    db.execute('INSERT INTO workspace_branding (workspace_id, primary_color, logo_text, '
               "font_family) VALUES ('default',?,?,?) "
               'ON CONFLICT(workspace_id) DO UPDATE SET primary_color=excluded.primary_color, '
               'logo_text=excluded.logo_text, font_family=excluded.font_family, '
               "updated_at=datetime('now')",
               (color, b.get('logo_text'), b.get('font_family')))
    db.commit()
    log_action('branding.updated', 'workspace', 'default', {'primary_color': color})
    return jsonify(one("SELECT * FROM workspace_branding WHERE workspace_id='default'"))


@app.get('/api/branding')
def get_branding():
    return jsonify(one("SELECT * FROM workspace_branding WHERE workspace_id='default'") or {})


@app.post('/api/artifacts/<int:id>/annotations')
def add_annotation(id):
    if not one('SELECT id FROM artifacts WHERE id=?', (id,)):
        return jsonify({'error': 'Not found'}), 404
    b = request.get_json() or {}
    if not b.get('text'):
        return jsonify({'error': 'text required'}), 400
    lid = execute('INSERT INTO artifact_annotations (artifact_id, grain_value, timestamp, '
                  'text, author) VALUES (?,?,?,?,?)',
                  (id, b.get('grain_value'), b.get('timestamp'), b['text'],
                   getattr(g, 'user_email', None)))
    log_action('artifact.annotated', 'artifact', id, {'annotation_id': lid})
    record_activity(id, 'annotated', b['text'][:80])
    return jsonify(one('SELECT * FROM artifact_annotations WHERE id=?', (lid,))), 201


@app.get('/api/artifacts/<int:id>/annotations')
def list_annotations(id):
    return jsonify(many('SELECT * FROM artifact_annotations WHERE artifact_id=? '
                        'ORDER BY id', (id,)))


@app.post('/api/artifacts/<int:id>/subscriptions')
def add_metric_subscription(id):
    if not one('SELECT id FROM artifacts WHERE id=?', (id,)):
        return jsonify({'error': 'Not found'}), 404
    b = request.get_json() or {}
    if b.get('metric') not in ('mape', 'avg_actual', 'forecast_avg'):
        return jsonify({'error': "metric must be 'mape', 'avg_actual', or 'forecast_avg'"}), 400
    if not isinstance(b.get('threshold'), (int, float)):
        return jsonify({'error': 'numeric threshold required'}), 400
    lid = execute('INSERT INTO metric_subscriptions (artifact_id, metric, threshold, '
                  'direction, subscriber) VALUES (?,?,?,?,?)',
                  (id, b['metric'], b['threshold'], b.get('direction', 'above'),
                   getattr(g, 'user_email', None) or 'viewer@analytiq.dev'))
    log_action('subscription.created', 'artifact', id,
               {'metric': b['metric'], 'threshold': b['threshold']})
    record_activity(id, 'subscribed', b['metric'])
    return jsonify(one('SELECT * FROM metric_subscriptions WHERE id=?', (lid,))), 201


@app.get('/api/artifacts/<int:id>/preview')
def artifact_preview(id):
    art = one('SELECT * FROM artifacts WHERE id=?', (id,))
    if not art:
        return jsonify({'error': 'Not found'}), 404
    f = one('SELECT * FROM artifact_files WHERE artifact_id=? ORDER BY version DESC LIMIT 1', (id,))
    if not f:
        return jsonify({'error': 'Artifact not rendered yet'}), 404
    rows = many('SELECT * FROM chart_data WHERE pipeline_run_id=? ORDER BY day_index',
                (art['pipeline_run_id'],)) if art['pipeline_run_id'] else []
    return jsonify({'artifact': art, 'html': f['html'],
                    'kpis': compute_kpis(rows) if rows else {},
                    'version': f['version'], 'sha256': f['sha256']})


@app.get('/api/artifacts/<int:id>/files')
def list_artifact_files(id):
    rows = many('SELECT id, artifact_id, version, size_bytes, sha256, validator_json, created_at '
                'FROM artifact_files WHERE artifact_id=? ORDER BY version DESC', (id,))
    for r in rows:
        r['validation'] = json.loads(r.pop('validator_json') or '{}')
    return jsonify(rows)


def _compute_artifact_insights(id):
    """R7S2E3 insight engine, extracted (R12S1E1) so the Opportunity Engine
    reuses the same detector instead of duplicating it. Returns
    (insight_rows, error_response)."""
    art = one('SELECT * FROM artifacts WHERE id=?', (id,))
    if not art or not art['pipeline_run_id']:
        return None, (jsonify({'error': 'No chart data to scan'}), 404)
    rows = many('SELECT * FROM chart_data WHERE pipeline_run_id=? ORDER BY day_index',
                (art['pipeline_run_id'],))
    actuals = [(r_['day_index'], r_['actual'], r_['date']) for r_ in rows
               if r_['actual'] is not None]
    if len(actuals) < 14:
        return None, (jsonify({'error': 'Not enough history to scan'}), 409)
    vals = [a[1] for a in actuals]
    mean = sum(vals) / len(vals)
    import math as _m
    std = _m.sqrt(sum((v - mean) ** 2 for v in vals) / (len(vals) - 1)) or 1

    insights = []
    # anomalies beyond 3σ
    for idx, v, d in actuals:
        if abs(v - mean) > 3 * std:
            insights.append(('anomaly',
                             f'{d}: value {v:,.0f} is more than 3σ from the mean',
                             f'Why did the metric spike on {d}?'))
    # trend: first vs last quartile
    q = max(len(vals) // 4, 1)
    first_avg, last_avg = sum(vals[:q]) / q, sum(vals[-q:]) / q
    delta = (last_avg - first_avg) / (first_avg or 1) * 100
    if abs(delta) > 3:
        direction = 'upward' if delta > 0 else 'downward'
        insights.append(('trend',
                         f'{direction} trend: {delta:+.1f}% from the first to the '
                         f'last quarter of the window',
                         f'What is driving the {direction} trend in this metric?'))
    # weekday pattern
    byday = {}
    for idx, v, _d in actuals:
        byday.setdefault(idx % 7, []).append(v)
    day_means = {d_: sum(v_) / len(v_) for d_, v_ in byday.items()}
    hi_d, lo_d = max(day_means, key=day_means.get), min(day_means, key=day_means.get)
    spread = (day_means[hi_d] - day_means[lo_d]) / (mean or 1) * 100
    if spread > 10:
        names = ['Mon', 'Tue', 'Wed', 'Thu', 'Fri', 'Sat', 'Sun']
        insights.append(('weekday_pattern',
                         f'{names[hi_d]} runs {spread:.0f}% above {names[lo_d]} on average',
                         f'Summarize performance by weekday — why is {names[lo_d]} weakest?'))

    db = get_db()
    db.execute('DELETE FROM artifact_insights WHERE artifact_id=? AND dismissed=0', (id,))
    out = []
    for kind, summary, question in insights:
        lid = execute('INSERT INTO artifact_insights (artifact_id, kind, summary, '
                      'drill_question) VALUES (?,?,?,?)', (id, kind, summary, question))
        out.append({'id': lid, 'kind': kind, 'summary': summary,
                    'drill_question': question})
    log_action('insights.scanned', 'artifact', id, {'found': len(out)})
    return out, None


@app.post('/api/artifacts/<int:id>/insights/scan')
@limiter.limit('20/minute')
def scan_insights(id):
    """R7S2E3: background insight agent — anomalies, trends, weekday patterns."""
    out, err = _compute_artifact_insights(id)
    if err:
        return err
    return jsonify({'artifact_id': id, 'insights': out}), 201


@app.get('/api/artifacts/<int:id>/insights')
def list_insights(id):
    return jsonify(many('SELECT * FROM artifact_insights WHERE artifact_id=? AND dismissed=0 '
                        'ORDER BY id', (id,)))


@app.post('/api/insights/<int:id>/dismiss')
def dismiss_insight(id):
    row_i = one('SELECT * FROM artifact_insights WHERE id=?', (id,))
    if not row_i:
        return jsonify({'error': 'Not found'}), 404
    # R10S1E1: dismissals are a first-class preference signal (§17.3.1)
    try:
        import agent_memory as am
        am.remember(get_db(), 'insight_agent', 'dismissed_insight',
                    f"{row_i['kind']}:{row_i['artifact_id']}", row_i['summary'][:120])
    except ValueError:
        pass  # PII-gated values are never stored
    db = get_db()
    db.execute('UPDATE artifact_insights SET dismissed=1 WHERE id=?', (id,))
    db.commit()
    return jsonify({'id': id, 'dismissed': True})


@app.post('/api/insights/<int:id>/drill')
def drill_insight(id):
    """One-click drill-in: new session planned from the insight's question."""
    import planner
    row = one('SELECT * FROM artifact_insights WHERE id=?', (id,))
    if not row:
        return jsonify({'error': 'Not found'}), 404
    sid = execute("INSERT INTO sessions (metric, grain, status) "
                  "VALUES ('Net Revenue', 'Location · Day', 'pending')")
    schema_row = _latest_schema_row('default')
    plan = planner.plan_session(
        row['drill_question'],
        semantic_schema=json.loads(schema_row['schema_json']) if schema_row else None,
        schema_version=schema_row['version'] if schema_row else None)
    log_action('insight.drilled', 'artifact', row['artifact_id'],
               {'insight_id': id, 'session_id': sid})
    return jsonify({'session_id': sid, 'plan': plan,
                    'question': row['drill_question']}), 201


@app.get('/api/artifacts/<int:id>/roi')
def artifact_roi_route(id):
    """R12S2E5: adoption vs cost for one artifact."""
    import roi
    out = roi.artifact_roi(get_db(), id)
    if out is None:
        return jsonify({'error': 'Artifact not found'}), 404
    return jsonify(out)


@app.post('/api/workspace/roi_report')
@require_role('admin', 'analyst')
def workspace_roi_report():
    """R12S2E5: the ROI report is itself an AnalytIQ artifact — the platform
    demonstrates its own value with its own pipeline."""
    import artifact_gen as ag
    import roi
    rows = roi.workspace_report(get_db())
    # Native-artifact pattern (R7S2E3 precedent): the platform's own eight-
    # panel template fed with the ROI series, plus an embedded report table.
    chart_rows = [{'day_index': i, 'date': r['title'][:14],
                   'actual': r['adoption_score'], 'predicted': r['adoption_score'],
                   'ci_low': max(0, r['adoption_score'] - 1),
                   'ci_high': r['adoption_score'] + 1, 'is_forecast': 0}
                  for i, r in enumerate(rows)] or [
        {'day_index': 0, 'date': 'no artifacts', 'actual': 0, 'predicted': 0,
         'ci_low': 0, 'ci_high': 0, 'is_forecast': 0}]
    aid = execute('INSERT INTO artifacts (title, type, owner, dq_status, is_sandbox) '
                  "VALUES ('Workspace ROI Report', 'Report', ?, 'pass', 0)",
                  (getattr(g, 'user_email', None) or 'admin@acme.com',))
    art_row = one('SELECT * FROM artifacts WHERE id=?', (aid,))
    html = ag.generate_artifact_html(art_row, chart_rows, compute_kpis(chart_rows))
    cells = []
    for r in rows:
        cells.append('<tr><td>' + r['title'] + '</td><td>' + str(r['adoption_score'])
                     + '</td><td>$' + str(r['est_cost']) + '</td><td>'
                     + str(r['roi_ratio']) + '</td><td>' + str(r['signals']['views'])
                     + 'v ' + str(r['signals']['shares']) + 's '
                     + str(r['signals']['subscriptions']) + 'sub</td></tr>')
    body_rows = ''.join(cells) or '<tr><td colspan="5">No production artifacts yet</td></tr>'
    table = ('<section id="roi-report"><h2>Adoption vs cost</h2>'
             '<table role="table"><thead><tr><th>Artifact</th><th>Adoption score</th>'
             '<th>Est. cost</th><th>ROI ratio</th><th>Signals</th></tr></thead>'
             '<tbody>' + body_rows + '</tbody></table></section>')
    html = html.replace('</body>', table + '</body>', 1)
    _attempts_roi = []
    html, _cycles, validation = ag.validate_and_repair(html, attempt_log=_attempts_roi)
    _persist_repair_attempts(get_db(), 'artifact_render', aid, None, _attempts_roi,
                             validation['status'] == 'PASS')
    _store_artifact_file(aid, 1, html, validation)
    log_action('workspace.roi_report', 'artifact', aid, {'rows': len(rows)})
    return jsonify(one('SELECT * FROM artifacts WHERE id=?', (aid,))), 201


@app.post('/api/workspace/health_dashboard')
@require_role('admin', 'analyst')
def workspace_health_dashboard():
    """R7S2E3: the workspace health dashboard is itself an AnalytIQ artifact."""
    import hashlib as _h
    import artifact_gen as ag
    history = many('SELECT * FROM health_history ORDER BY id')
    runs = many('SELECT id FROM governance_runs ORDER BY id')
    per_run = {}
    for h in history:
        per_run.setdefault(h['run_id'], []).append(h['health_score'] or 0)
    series = [(rid, sum(v) / len(v)) for rid, v in sorted(per_run.items())]
    rows = [{'day_index': i, 'date': f'run {rid}', 'actual': round(avg, 1),
             'predicted': round(avg, 1), 'ci_low': round(avg - 5, 1),
             'ci_high': round(avg + 5, 1), 'is_forecast': 0}
            for i, (rid, avg) in enumerate(series)] or [
        {'day_index': 0, 'date': 'no runs', 'actual': 0, 'predicted': 0,
         'ci_low': 0, 'ci_high': 0, 'is_forecast': 0}]
    avg_health = round(sum(r_['actual'] for r_ in rows) / len(rows), 1)
    refresh_ok = one("SELECT COUNT(*) c FROM pipeline_runs WHERE status='done'")['c']
    kpis = {'avgActual': avg_health, 'mape': 0,
            'forecast14Avg': refresh_ok}

    aid = execute("INSERT INTO artifacts (title, type, dq_status) "
                  "VALUES ('Workspace health', 'Workspace Health', 'pass')")
    art = one('SELECT * FROM artifacts WHERE id=?', (aid,))
    html = ag.generate_artifact_html(art, rows, kpis)
    _attempts2 = []
    html, cycles, validation = ag.validate_and_repair(html, attempt_log=_attempts2)
    _persist_repair_attempts(get_db(), 'artifact_render', aid, None,
                             _attempts2, validation['status'] == 'PASS')
    fid, sha, _uri = _store_artifact_file(aid, 1, html, validation)
    log_action('workspace.health_dashboard', 'artifact', aid,
               {'avg_health': avg_health, 'runs': len(series)})
    return jsonify({**art, 'tags': [],
                    'file': {'file_id': fid, 'version': 1, 'sha256': sha,
                             'validation_status': validation['status']}}), 201


@app.get('/api/artifacts/<int:id>/thumbnail')
def artifact_thumbnail(id):
    """R7S2E2: lightweight SVG sparkline thumbnail from the chart data."""
    art = one('SELECT * FROM artifacts WHERE id=?', (id,))
    if not art or not art['pipeline_run_id']:
        return jsonify({'error': 'No chart data for a thumbnail'}), 404
    rows = many('SELECT predicted FROM chart_data WHERE pipeline_run_id=? '
                'ORDER BY day_index', (art['pipeline_run_id'],))
    if not rows:
        return jsonify({'error': 'No chart data for a thumbnail'}), 404
    vals = [r_['predicted'] for r_ in rows if r_['predicted'] is not None]
    lo, hi = min(vals), max(vals)
    span = (hi - lo) or 1
    W, H = 160, 48
    pts = ' '.join(f'{round(i * W / max(len(vals) - 1, 1), 1)},'
                   f'{round(H - 6 - (v - lo) * (H - 12) / span, 1)}'
                   for i, v in enumerate(vals))
    svg = (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {W} {H}" '
           f'width="{W}" height="{H}" role="img" aria-label="Forecast sparkline">'
           f'<rect width="{W}" height="{H}" fill="#f7f8fa" rx="6"/>'
           f'<polyline points="{pts}" fill="none" stroke="#4f7cff" stroke-width="1.6"/>'
           f'</svg>')
    return Response(svg, mimetype='image/svg+xml')


@app.post('/api/artifacts/<int:id>/favorite')
def toggle_favorite(id):
    row = one('SELECT favorite FROM artifacts WHERE id=?', (id,))
    if not row:
        return jsonify({'error': 'Not found'}), 404
    new = 0 if row['favorite'] else 1
    db = get_db()
    db.execute('UPDATE artifacts SET favorite=? WHERE id=?', (new, id))
    db.commit()
    if new:
        record_activity(id, 'favorited')
    return jsonify({'artifact_id': id, 'favorite': new})


@app.put('/api/artifacts/<int:id>/tags')
def put_tags(id):
    if not one('SELECT id FROM artifacts WHERE id=?', (id,)):
        return jsonify({'error': 'Not found'}), 404
    tags = (request.get_json() or {}).get('tags')
    if not isinstance(tags, list):
        return jsonify({'error': 'tags must be a list'}), 400
    db = get_db()
    db.execute('UPDATE artifacts SET tags_json=? WHERE id=?',
               (json.dumps(sorted(set(map(str, tags)))), id))
    db.commit()
    return jsonify({'artifact_id': id, 'tags': sorted(set(map(str, tags)))})


@app.get('/api/artifacts/<int:id>/activity')
def artifact_activity(id):
    if not one('SELECT id FROM artifacts WHERE id=?', (id,)):
        return jsonify({'error': 'Not found'}), 404
    return jsonify(many('SELECT * FROM artifact_activity WHERE artifact_id=? '
                        'ORDER BY id DESC LIMIT 200', (id,)))


@app.post('/api/artifacts/<int:id>/embed_tokens')
@require_role('admin', 'analyst')
def create_embed_token(id):
    import embed_tokens as et
    if not one('SELECT id FROM artifacts WHERE id=?', (id,)):
        return jsonify({'error': 'Not found'}), 404
    b = request.get_json() or {}
    scope = b.get('scope', 'read_only')
    if scope not in ('read_only', 'interactive'):
        return jsonify({'error': "scope must be 'read_only' or 'interactive'"}), 400
    hours = min(int(b.get('expires_in_hours', 24)), 24 * 365)
    token = et.sign({'artifact_id': id, 'workspace_id': 'default', 'scope': scope,
                     'allowed_origins': b.get('allowed_origins') or ['*']},
                    et.workspace_secret('default'), expires_in=hours * 3600)
    log_action('embed_token.created', 'artifact', id,
               {'scope': scope, 'expires_in_hours': hours,
                'allowed_origins': b.get('allowed_origins')})
    return jsonify({'token': token, 'scope': scope, 'expires_in_hours': hours}), 201


@app.post('/api/artifacts/<int:id>/share_links')
@require_role('admin', 'analyst')
def create_share_link(id):
    # R20S1E1: plan entitlement gate (starter has no public links)
    if not _entitled('public_links'):
        return jsonify({'error': 'Public links are not included in the '
                                 f'{_current_plan()} plan — upgrade to enable.'}), 403
    """R7S1E1: password-optional public link serving an hourly-refreshed snapshot."""
    import hashlib as _h
    import secrets as _s
    import authn
    if not one('SELECT id FROM artifacts WHERE id=?', (id,)):
        return jsonify({'error': 'Not found'}), 404
    f = one('SELECT html FROM artifact_files WHERE artifact_id=? '
            'ORDER BY version DESC LIMIT 1', (id,))
    if not f:
        return jsonify({'error': 'Render the artifact before creating a public link'}), 409
    b = request.get_json() or {}
    token = _s.token_urlsafe(24)
    hours = b.get('expires_in_hours')
    pw = b.get('password')
    lid = execute('INSERT INTO share_links (artifact_id, token_hash, password_hash, '
                  "expires_at, snapshot_html, snapshot_at, created_by) "
                  "VALUES (?,?,?, CASE WHEN ? IS NULL THEN NULL ELSE datetime('now', ?) END, "
                  "?, datetime('now'), ?)",
                  (id, _h.sha256(token.encode()).hexdigest(),
                   authn.hash_password(pw) if pw else None,
                   hours, f'+{hours} hours' if hours else None,
                   f['html'], getattr(g, 'user_email', None)))
    log_action('share_link.created', 'artifact', id,
               {'share_link_id': lid, 'expires_in_hours': hours,
                'password_protected': bool(pw)})
    return jsonify({'id': lid, 'token': token, 'url': f'/api/public/{token}',
                    'expires_in_hours': hours}), 201


@app.post('/api/artifacts/<int:id>/share_links/revoke')
@require_role('admin', 'analyst')
def revoke_share_links(id):
    """R30S3E4-US1: the share modal's red "Revoke link" — expiry-based
    revocation (no schema change), audited; the public route 410s instantly."""
    if not one('SELECT id FROM artifacts WHERE id=?', (id,)):
        return jsonify({'error': 'Not found'}), 404
    n = one("SELECT COUNT(*) AS c FROM share_links WHERE artifact_id=? "
            "AND (expires_at IS NULL OR expires_at > datetime('now'))", (id,))['c']
    execute("UPDATE share_links SET expires_at = datetime('now', '-1 hour') "
            "WHERE artifact_id=? AND (expires_at IS NULL OR expires_at > datetime('now'))",
            (id,))
    log_action('share_link.revoked', 'artifact', id, {'revoked': n})
    return jsonify({'revoked': n})


@app.get('/embed/<token>')
def embed_render(token):
    """R19S1E1: embed render route with server-side allowed_origins
    enforcement (arch §11.2/§14.1) — origin checked on every request."""
    import embed_tokens as et
    payload = et.verify(token, et.workspace_secret('default'))
    if not payload:
        return jsonify({'error': 'Invalid or expired embed token'}), 401
    origin = request.headers.get('Origin') or request.headers.get('Referer', '')
    allowed = payload.get('allowed_origins') or []
    if '*' not in allowed and not any(origin.startswith(a) for a in allowed):
        log_action('embed.origin_denied', 'artifact', payload.get('artifact_id'),
                   {'origin': origin})
        return jsonify({'error': f'Origin {origin!r} is not permitted for this embed'}), 403
    f = one('SELECT html FROM artifact_files WHERE artifact_id=? ORDER BY version DESC LIMIT 1',
            (payload.get('artifact_id'),))
    if not f:
        return jsonify({'error': 'Artifact has no rendered file'}), 404
    return Response(f['html'], mimetype='text/html')


@app.get('/api/public/<token>/meta')
def public_artifact_meta(token):
    """R19S1E1: branded-viewer metadata (title, freshness, expiry)."""
    import hashlib as _h
    row = one('SELECT * FROM share_links WHERE token_hash=?',
              (_h.sha256(token.encode()).hexdigest(),))
    if not row:
        return jsonify({'error': 'Unknown link'}), 404
    art = one('SELECT * FROM artifacts WHERE id=?', (row['artifact_id'],))
    fresh = one("SELECT CAST((julianday('now') - julianday(created_at)) * 24 AS INTEGER) AS h "
                'FROM artifact_files WHERE artifact_id=? ORDER BY version DESC LIMIT 1',
                (row['artifact_id'],))
    return jsonify({'title': art['title'] if art else 'Artifact',
                    'freshness_hours': fresh['h'] if fresh else None,
                    'expires_at': row['expires_at'],
                    'password_protected': bool(row['password_hash']),
                    'view_count': row['view_count']})


@app.get('/api/public/<token>')
@limiter.limit('60/minute')
def public_artifact(token):
    import hashlib as _h
    import authn
    row = one('SELECT * FROM share_links WHERE token_hash=?',
              (_h.sha256(token.encode()).hexdigest(),))
    if not row:
        return jsonify({'error': 'Unknown link'}), 404
    if row['expires_at'] and row['expires_at'] <= one("SELECT datetime('now') AS n")['n']:
        return jsonify({'error': 'This link has expired'}), 410
    if row['password_hash']:
        pw = request.args.get('password') or request.headers.get('X-Link-Password') or ''
        if not authn.verify_password(pw, row['password_hash']):
            return jsonify({'error': 'Password required'}), 401
    db = get_db()
    # refresh the snapshot at most once per hour from the latest render
    stale = one("SELECT 1 AS s FROM share_links WHERE id=? "
                "AND snapshot_at <= datetime('now', '-1 hour')", (row['id'],))
    if stale:
        latest = one('SELECT html FROM artifact_files WHERE artifact_id=? '
                     'ORDER BY version DESC LIMIT 1', (row['artifact_id'],))
        if latest:
            db.execute("UPDATE share_links SET snapshot_html=?, snapshot_at=datetime('now') "
                       'WHERE id=?', (latest['html'], row['id']))
            row = one('SELECT * FROM share_links WHERE id=?', (row['id'],))
    db.execute('UPDATE share_links SET view_count=view_count+1 WHERE id=?', (row['id'],))
    db.commit()
    return Response(row['snapshot_html'], mimetype='text/html')


@app.get('/api/artifacts/<int:id>/shares')
def list_shares(id):
    return jsonify(many('SELECT * FROM artifact_shares WHERE artifact_id=? ORDER BY shared_at DESC', (id,)))

SHARE_ROLES = ('Viewer', 'Editor', 'Owner')
_EMAIL_RE = __import__('re').compile(r'^[\w.+-]+@[\w-]+\.[\w.-]+$')


@app.post('/api/artifacts/<int:id>/shares')
@require_role('admin', 'analyst')
def add_share(id):
    b = request.get_json() or {}
    email = (b.get('email') or '').strip()
    role = b.get('role', 'Viewer')
    if not email or not _EMAIL_RE.match(email):
        return jsonify({'error': 'A valid email address is required'}), 400
    if role not in SHARE_ROLES:
        return jsonify({'error': f'role must be one of {SHARE_ROLES}'}), 400
    art = one('SELECT * FROM artifacts WHERE id=?', (id,))
    if not art:
        return jsonify({'error': 'Not found'}), 404
    if one('SELECT id FROM artifact_shares WHERE artifact_id=? AND email=?', (id, email)):
        return jsonify({'error': f'{email} already has access to this artifact'}), 409

    lid = execute('INSERT INTO artifact_shares (artifact_id,email,role) VALUES (?,?,?)',
                  (id, email, role))
    send_email(to=email, subject=f"{art['owner']} shared an artifact with you",
               html=(f"<h2>{art['title']}</h2>"
                     f"<p>You've been given <strong>{role}</strong> access in AnalytIQ.</p>"
                     f"<p>Log in to view the live artifact.</p>"))
    log_action('share.added', 'artifact', id,
               {'share_id': lid, 'email': email, 'role': role, 'notified': True})
    record_activity(id, 'shared', email)
    return jsonify(one('SELECT * FROM artifact_shares WHERE id=?', (lid,))), 201

@app.delete('/api/artifacts/<int:art_id>/shares/<int:share_id>')
@require_role('admin', 'analyst')
def remove_share(art_id, share_id):
    if not one('SELECT id FROM artifact_shares WHERE id=? AND artifact_id=?',
               (share_id, art_id)):
        return jsonify({'error': 'Share not found'}), 404
    execute('DELETE FROM artifact_shares WHERE id=? AND artifact_id=?', (share_id, art_id))
    log_action('share.removed', 'artifact', art_id, {'share_id': share_id})
    return '', 204


@app.post('/api/artifacts/<int:id>/refresh')
@limiter.limit('10/minute')
@require_role('admin', 'analyst')
def refresh_artifact(id):
    import hashlib
    import artifact_gen as ag
    art = one('SELECT * FROM artifacts WHERE id=?', (id,))
    if not art:
        return jsonify({'error': 'Not found'}), 404
    if not art['pipeline_run_id']:
        return jsonify({'error': 'Artifact has no pipeline run to refresh'}), 409

    old_run = one('SELECT session_id FROM pipeline_runs WHERE id=?', (art['pipeline_run_id'],))
    session_id = old_run['session_id'] if old_run else None

    # synchronous re-score: new pipeline run + fresh chart data
    new_run_id = execute("INSERT INTO pipeline_runs (session_id,status,current_step,log_entries) "
                         "VALUES (?,'running',0,'[]')", (session_id,))
    db = get_db()
    db.executemany(
        'INSERT INTO chart_data (pipeline_run_id,day_index,date,actual,predicted,ci_low,ci_high,is_forecast) '
        'VALUES (:pipeline_run_id,:day_index,:date,:actual,:predicted,:ci_low,:ci_high,:is_forecast)',
        generate_chart_data(new_run_id))
    db.execute("UPDATE pipeline_runs SET status='done',current_step=4,mape=8.9,features_count=34,"
               "rows_count=12847,completed_at=datetime('now') WHERE id=?", (new_run_id,))
    db.execute('UPDATE artifacts SET pipeline_run_id=? WHERE id=?', (new_run_id, id))
    db.commit()

    # regenerate the self-contained file
    art = one('SELECT * FROM artifacts WHERE id=?', (id,))
    rows = many('SELECT * FROM chart_data WHERE pipeline_run_id=? ORDER BY day_index', (new_run_id,))
    html = ag.generate_artifact_html(art, rows, compute_kpis(rows))
    validation = ag.validate_artifact(html)
    prev = one('SELECT MAX(version) AS v FROM artifact_files WHERE artifact_id=?', (id,))
    fversion = (prev['v'] or 0) + 1
    _store_artifact_file(id, fversion, html, validation)

    # R6S2E2: evaluate metric alert subscriptions against fresh KPIs
    kpis_now = compute_kpis(rows)
    kpi_map = {'mape': kpis_now['mape'], 'avg_actual': kpis_now['avgActual'],
               'forecast_avg': kpis_now['forecast14Avg']}
    for sub in many('SELECT * FROM metric_subscriptions WHERE artifact_id=?', (id,)):
        value = kpi_map.get(sub['metric'])
        if value is None:
            continue
        fired = value > sub['threshold'] if sub['direction'] == 'above' else value < sub['threshold']
        if fired:
            execute('INSERT INTO alerts (type, connection_id, subject, detail_json) '
                    "VALUES ('metric', NULL, ?, ?)",
                    (f"Metric alert: {sub['metric']} {sub['direction']} {sub['threshold']}",
                     json.dumps({'artifact_id': id, 'metric': sub['metric'],
                                 'value': value, 'threshold': sub['threshold']})))
            send_email(to=sub['subscriber'] or 'viewer@analytiq.dev',
                       subject=f"Alert: {sub['metric']} is {value} "
                               f"({sub['direction']} {sub['threshold']})",
                       html=(f'<p>Your subscribed metric fired.</p>'
                             f'<p><a href="/artifacts/{id}">Open the artifact</a></p>'))

    # R5S2E2: drift check — rolling 30-day MAPE vs training baseline
    drift = None
    sess = one('SELECT session_id FROM pipeline_runs WHERE id=?', (new_run_id,))
    if sess and sess['session_id'] is not None:
        card = one("SELECT * FROM model_cards WHERE session_id=? AND status='promoted' "
                   'ORDER BY id DESC LIMIT 1', (sess['session_id'],)) or \
               one('SELECT * FROM model_cards WHERE session_id=? ORDER BY id DESC LIMIT 1',
                   (sess['session_id'],))
        if card:
            baseline = (json.loads(card['metrics_json'] or '{}') or {}).get('val_mape')
            recent = [r_ for r_ in rows if r_['actual'] is not None][-30:]
            if baseline and recent:
                rolling = round(sum(abs(r_['actual'] - r_['predicted']) / r_['actual']
                                    for r_ in recent) / len(recent) * 100, 2)
                drifting = rolling > baseline * 1.5
                drift = {'rolling_mape': rolling, 'baseline_mape': baseline,
                         'drifting': drifting, 'model_card_id': card['id']}
                if drifting:
                    execute('INSERT INTO alerts (type, connection_id, subject, detail_json) '
                            "VALUES ('model_drift', NULL, ?, ?)",
                            (f'Model drift on artifact {id}',
                             json.dumps({**drift, 'artifact_id': id,
                                         'session_id': sess['session_id']})))
                    send_email(to='workspace-admins@analytiq.dev',
                               subject='Model drift detected',
                               html=f'<p>Rolling MAPE {rolling}% vs baseline {baseline}%.</p>')

    log_action('artifact.refreshed', 'artifact', id,
               {'pipeline_run_id': new_run_id, 'file_version': fversion,
                'drift': drift})
    return jsonify({'artifact_id': id, 'pipeline_run_id': new_run_id,
                    'file_version': fversion, 'kpis': compute_kpis(rows),
                    'drift': drift})


@app.get('/api/artifacts/<int:id>/drift')
def artifact_drift(id):
    if not one('SELECT id FROM artifacts WHERE id=?', (id,)):
        return jsonify({'error': 'Not found'}), 404
    rows = many("SELECT * FROM alerts WHERE type='model_drift' ORDER BY id DESC LIMIT 50")
    history = []
    for r_ in rows:
        d = json.loads(r_['detail_json'] or '{}')
        if d.get('artifact_id') == id:
            history.append({**r_, 'detail': d})
    for h in history:
        h.pop('detail_json', None)
    return jsonify({'artifact_id': id,
                    'status': 'drifting' if history else 'healthy',
                    'history': history})


@app.post('/api/models/<int:session_id>/retrain')
@require_role('admin', 'analyst')
def one_click_retrain(session_id):
    """R5S2E2: reuse the spec with an advanced rolling window; archive champion."""
    from datetime import date as _date, timedelta as _td
    if not one('SELECT id FROM sessions WHERE id=?', (session_id,)):
        return jsonify({'error': 'Session not found'}), 404
    spec = _spec_for_session(session_id)
    if not spec:
        return jsonify({'error': 'No confirmed spec to retrain from'}), 409
    horizon = spec.get('prediction_horizon') or 14
    dr = spec.get('date_range') or {}
    try:
        new_range = {
            'start': (_date.fromisoformat(dr['start']) + _td(days=horizon)).isoformat(),
            'end': (_date.fromisoformat(dr['end']) + _td(days=horizon)).isoformat(),
        }
    except Exception:
        return jsonify({'error': 'Spec has no parseable date_range'}), 422
    new_spec = {**spec, 'date_range': new_range}

    import hashlib as _h
    prev = one('SELECT MAX(spec_version) AS v FROM session_specs WHERE session_id=?',
               (session_id,))
    execute('INSERT INTO session_specs (session_id, spec_version, payload_hash, spec_json) '
            'VALUES (?,?,?,?)',
            (session_id, (prev['v'] or 0) + 1,
             _h.sha256(json.dumps(new_spec, sort_keys=True).encode()).hexdigest(),
             json.dumps(new_spec)))

    db = get_db()
    db.execute("UPDATE model_registry SET status='archived' "
               "WHERE session_id=? AND status='active'", (session_id,))
    db.commit()

    gold = one("SELECT * FROM gold_tables WHERE session_id=? AND status='written' "
               'ORDER BY version DESC LIMIT 1', (session_id,))
    if not gold:
        return jsonify({'error': 'No gold table to retrain on'}), 409
    job_id = execute('INSERT INTO training_jobs (session_id, gold_table_id, status) '
                     "VALUES (?,?,'queued')", (session_id, gold['id']))
    _execute_training_job(job_id, session_id, gold, horizon)
    log_action('model.retrained', 'training_job', job_id,
               {'session_id': session_id, 'new_date_range': new_range})
    return jsonify({'jobId': job_id, 'date_range': new_range}), 201


# ─────────────────────────────────────────────────────────
# Routes — Artifact Schedules
# ─────────────────────────────────────────────────────────
@app.get('/api/artifacts/<int:id>/schedule')
def get_schedule(id):
    row = one('SELECT * FROM artifact_schedules WHERE artifact_id=?', (id,))
    return jsonify(row or {})


@app.put('/api/artifacts/<int:id>/schedule')
def upsert_schedule(id):
    b = request.get_json() or {}
    cron_expr = b.get('cron_expr')
    if not cron_expr:
        return jsonify({'error': 'cron_expr required'}), 400
    tz_name = b.get('timezone', 'UTC')
    try:
        from zoneinfo import ZoneInfo
        tz = ZoneInfo(tz_name)
    except Exception:
        return jsonify({'error': f'Unknown timezone {tz_name!r}'}), 400
    try:
        # R7S2E2: compute the next run in the workspace timezone, store as UTC
        from datetime import timezone as _tz
        base = datetime.now(tz)
        next_local = croniter(cron_expr, base).get_next(datetime)
        next_run = next_local.astimezone(_tz.utc).strftime('%Y-%m-%d %H:%M:%S')
    except (ValueError, KeyError):
        return jsonify({'error': 'Invalid cron expression'}), 400

    org_id = request.headers.get('X-Org-Id')
    existing = one('SELECT id FROM artifact_schedules WHERE artifact_id=?', (id,))
    db = get_db()
    if existing:
        db.execute(
            'UPDATE artifact_schedules SET cron_expr=?, next_run_at=?, timezone=?, '
            'enabled=1 WHERE artifact_id=?',
            (cron_expr, next_run, tz_name, id),
        )
    else:
        db.execute(
            'INSERT INTO artifact_schedules (artifact_id,org_id,cron_expr,next_run_at,timezone) '
            'VALUES (?,?,?,?,?)',
            (id, org_id, cron_expr, next_run, tz_name),
        )
    db.commit()
    log_action('schedule.upserted', 'artifact', id, {'cron_expr': cron_expr})
    return jsonify(one('SELECT * FROM artifact_schedules WHERE artifact_id=?', (id,)))


@app.delete('/api/artifacts/<int:id>/schedule')
def disable_schedule(id):
    db = get_db()
    db.execute('UPDATE artifact_schedules SET enabled=0 WHERE artifact_id=?', (id,))
    db.commit()
    log_action('schedule.disabled', 'artifact', id)
    return '', 204


# ─────────────────────────────────────────────────────────
# Routes — Audit Logs
@app.get('/api/audit-logs')
def list_audit_logs():
    org_id = getattr(g, 'org_id', None)
    resource_type = request.args.get('resource_type')
    resource_id = request.args.get('resource_id')
    limit = min(int(request.args.get('limit', 50)), 500)

    clauses = ['org_id IS ?']
    params = [org_id]

    if resource_type:
        clauses.append('resource_type = ?')
        params.append(resource_type)
    if resource_id:
        clauses.append('resource_id = ?')
        params.append(resource_id)
    action = request.args.get('action')
    if action:
        clauses.append('action = ?')
        params.append(action)

    params.append(limit)
    sql = f"SELECT * FROM audit_logs WHERE {' AND '.join(clauses)} ORDER BY created_at DESC LIMIT ?"
    return jsonify(many(sql, tuple(params)))


# ─────────────────────────────────────────────────────────
# Routes — Billing
# ─────────────────────────────────────────────────────────
@app.post('/api/billing/checkout')
def billing_checkout():
    b = request.get_json() or {}
    org_id = b.get('org_id') or request.headers.get('X-Org-Id')
    if not org_id:
        return jsonify({'error': 'org_id required'}), 400
    if not STRIPE_SECRET_KEY:
        return jsonify({'error': 'Billing not configured'}), 503

    sub = one("SELECT stripe_customer_id FROM subscriptions WHERE org_id=? AND stripe_customer_id IS NOT NULL LIMIT 1", (org_id,))
    if sub:
        customer_id = sub['stripe_customer_id']
    else:
        customer = stripe.Customer.create(metadata={'org_id': org_id})
        customer_id = customer.id

    session = stripe.checkout.Session.create(
        customer=customer_id,
        line_items=[{'price': STRIPE_PRO_PRICE_ID, 'quantity': 1}],
        mode='subscription',
        success_url=os.environ.get('CLIENT_URL', 'http://localhost:5173') + '?billing=success',
        cancel_url=os.environ.get('CLIENT_URL', 'http://localhost:5173') + '?billing=cancel',
        metadata={'org_id': org_id},
    )
    return jsonify({'url': session.url})


@app.post('/api/billing/portal')
def billing_portal():
    b = request.get_json() or {}
    org_id = b.get('org_id') or request.headers.get('X-Org-Id')
    if not org_id:
        return jsonify({'error': 'org_id required'}), 400
    if not STRIPE_SECRET_KEY:
        return jsonify({'error': 'Billing not configured'}), 503

    sub = one(
        "SELECT stripe_customer_id FROM subscriptions "
        "WHERE org_id=? AND stripe_customer_id IS NOT NULL LIMIT 1",
        (org_id,),
    )
    if not sub:
        return jsonify({'error': 'No billing account found'}), 404

    session = stripe.billing_portal.Session.create(
        customer=sub['stripe_customer_id'],
        return_url=os.environ.get('CLIENT_URL', 'http://localhost:5173'),
    )
    return jsonify({'url': session.url})


@app.post('/api/billing/webhook')
@limiter.exempt
def billing_webhook():
    payload = request.data
    sig = request.headers.get('Stripe-Signature')
    try:
        event = stripe.Webhook.construct_event(payload, sig, STRIPE_WEBHOOK_SECRET)
    except Exception:
        return '', 400

    if event['type'] == 'checkout.session.completed':
        session_obj = event['data']['object']
        org_id = session_obj.get('metadata', {}).get('org_id')
        customer_id = session_obj.get('customer')
        subscription_id = session_obj.get('subscription')

        sub_data = stripe.Subscription.retrieve(subscription_id)
        period_end = datetime.utcfromtimestamp(sub_data['current_period_end'])

        existing = one("SELECT id FROM subscriptions WHERE org_id=?", (org_id,))
        db = get_db()
        if existing:
            db.execute(
                "UPDATE subscriptions SET stripe_customer_id=?, stripe_subscription_id=?, "
                "plan='pro', status='active', current_period_end=? WHERE org_id=?",
                (customer_id, subscription_id, period_end, org_id),
            )
        else:
            db.execute(
                "INSERT INTO subscriptions (org_id, stripe_customer_id, stripe_subscription_id, plan, status, current_period_end) "
                "VALUES (?, ?, ?, 'pro', 'active', ?)",
                (org_id, customer_id, subscription_id, period_end),
            )
        db.commit()

    elif event['type'] == 'customer.subscription.deleted':
        subscription_obj = event['data']['object']
        subscription_id = subscription_obj['id']
        db = get_db()
        db.execute(
            "UPDATE subscriptions SET plan='free', status='canceled' WHERE stripe_subscription_id=?",
            (subscription_id,),
        )
        db.commit()

    return '', 200


# Scheduled refresh
def _run_scheduled_refresh(conn, sched):
    session_id = None
    if sched['pipeline_run_id']:
        row = conn.execute('SELECT session_id FROM pipeline_runs WHERE id=?', (sched['pipeline_run_id'],)).fetchone()
        if row:
            session_id = row['session_id']

    ins = conn.execute(
        "INSERT INTO pipeline_runs (session_id,status,current_step,log_entries) VALUES (?,'running',0,'[]')",
        (session_id,),
    )
    new_run_id = ins.lastrowid
    conn.commit()

    conn.executemany(
        'INSERT INTO chart_data (pipeline_run_id,day_index,date,actual,predicted,ci_low,ci_high,is_forecast) '
        'VALUES (:pipeline_run_id,:day_index,:date,:actual,:predicted,:ci_low,:ci_high,:is_forecast)',
        generate_chart_data(new_run_id),
    )
    conn.execute(
        "UPDATE pipeline_runs SET status='done',current_step=4,mape=8.9,features_count=34,"
        "rows_count=12847,completed_at=datetime('now') WHERE id=?",
        (new_run_id,),
    )
    conn.execute('UPDATE artifacts SET pipeline_run_id=? WHERE id=?', (new_run_id, sched['artifact_id']))

    cron = croniter(sched['cron_expr'], datetime.utcnow())
    next_run = cron.get_next(datetime)
    conn.execute(
        "UPDATE artifact_schedules SET next_run_at=?, last_run_at=datetime('now') WHERE id=?",
        (next_run, sched['id']),
    )
    conn.commit()
    log_action_bg(conn, 'schedule.executed', 'artifact', sched['artifact_id'],
                  {'pipeline_run_id': new_run_id})

    if sched.get('artifact_owner'):
        send_email(
            to=sched['artifact_owner'],
            subject=f'Forecast refreshed: {sched["artifact_title"]}',
            html=(
                '<h2>Scheduled refresh complete</h2>'
                '<p>Your artifact <strong>' + sched["artifact_title"] + '</strong> has been refreshed.</p>'
                '<p>Log in to AnalytIQ to view the updated forecast.</p>'
            ),
        )


def scheduler_loop():
    while True:
        time.sleep(60)
        conn = thread_db()
        try:
            due = conn.execute(
                "SELECT s.*, a.title AS artifact_title, a.owner AS artifact_owner, a.pipeline_run_id "
                "FROM artifact_schedules s "
                "JOIN artifacts a ON a.id = s.artifact_id "
                "WHERE s.enabled = 1 AND s.next_run_at <= datetime('now')"
            ).fetchall()
            for sched in due:
                try:
                    _run_scheduled_refresh(conn, dict(sched))
                except Exception:
                    pass
            _run_due_api_polls(conn)
        except Exception:
            pass
        finally:
            put_db(conn)
            try:
                conn.close()
            except Exception:
                pass


# ─────────────────────────────────────────────────────────
# Routes — Unified Artifact Store (R8S1E1 / §17.3.2)
# ─────────────────────────────────────────────────────────
@app.get('/api/uas/artifacts')
def uas_list():
    import uas
    run_id = request.args.get('run_id', type=int)
    rows = uas.list_artifacts(get_db(), artifact_type=request.args.get('type'),
                              run_id=run_id,
                              workspace_id=request.args.get('workspace'))
    return jsonify({'artifacts': [{k: v for k, v in r.items() if k != 'payload_json'}
                                  for r in rows]})


@app.get('/api/uas/artifacts/<uid>')
def uas_get(uid):
    import uas
    row = uas.get_by_uid(get_db(), uid)
    if not row:
        return jsonify({'error': 'UAS artifact not found'}), 404
    payload = json.loads(row.pop('payload_json'))
    ups = [uas.get_by_uid(get_db(), u) for u in row['upstream_artifact_ids']]
    row['upstream'] = [{'artifact_uid': u['artifact_uid'], 'artifact_type': u['artifact_type'],
                        'version': u['version']} for u in ups if u]
    return jsonify({**row, 'payload': payload})


@app.get('/api/uas/artifacts/<uid>/versions')
def uas_versions(uid):
    import uas
    rows = uas.versions_of(get_db(), uid)
    if rows is None:
        return jsonify({'error': 'UAS artifact not found'}), 404
    return jsonify({'versions': [{k: v for k, v in r.items() if k != 'payload_json'}
                                 for r in rows]})


@app.get('/api/diff')
def artifact_diff():
    """R11S2E4: structural diff between two versions of a versioned object."""
    import diff_engine as de
    import uas
    kind, a, b = request.args.get('kind'), request.args.get('a'), request.args.get('b')
    if kind not in ('semantic_schema', 'dashboard_plan', 'governance_manifest', 'model_card'):
        return jsonify({'error': 'kind must be one of semantic_schema, dashboard_plan, '
                                 'governance_manifest, model_card'}), 400
    if not a or not b:
        return jsonify({'error': 'a and b version references required'}), 400

    def load(ref):
        if kind == 'semantic_schema':
            row = one("SELECT schema_json AS j FROM semantic_schemas WHERE workspace_id='default' "
                      'AND version=?', (ref,))
            return json.loads(row['j']) if row else None
        if kind == 'dashboard_plan':
            node = uas.get_by_uid(get_db(), ref)
            return json.loads(node['payload_json']) if node and                 node['artifact_type'] == 'dashboard_plan' else None
        if kind == 'governance_manifest':
            row = one('SELECT manifest_json AS j FROM governance_manifests WHERE id=?', (ref,))
            return json.loads(row['j']) if row else None
        row = one('SELECT * FROM model_cards WHERE id=?', (ref,))
        if not row:
            return None
        return {'algorithm': row['algorithm'], 'status': row['status'],
                'gold_table_name': row['gold_table_name'],
                'metrics': json.loads(row['metrics_json'] or '{}'),
                'gates': json.loads(row['gates_json'] or '{}'),
                'hyperparams': json.loads(row['hyperparams_json'] or '{}')}

    da, db_ = load(a), load(b)
    if da is None or db_ is None:
        return jsonify({'error': f'{kind} version not found'}), 404
    out = {'kind': kind, 'a': a, 'b': b, 'structural': de.structural_diff(da, db_)}
    if kind == 'semantic_schema':
        out['summary'] = de.semantic_summary(da, db_)
    return jsonify(out)


@app.get('/api/pipeline/<int:run_id>/replay')
def pipeline_replay(run_id):
    """R11S2E3: step-by-step replay of a run — each DAG node with its stored
    UAS payload and gate results, plus retained repair attempts. Read-only:
    served straight from the store, no re-execution."""
    import dag
    import uas
    g_ = dag.graph(get_db(), run_id)
    if g_ is None:
        return jsonify({'error': 'No DAG recorded for this run'}), 404
    run = one('SELECT * FROM pipeline_runs WHERE id=?', (run_id,))
    sess = one('SELECT * FROM sessions WHERE id=?', (run['session_id'],)) if run else None
    ns = _uas_ns(sess) if sess else 'default'
    sid = sess['id'] if sess else None

    payload_keys = {'session_plan': f'{ns}:dashboard_plan:s{sid}',
                    'viz_specs': f'{ns}:vega_lite_specs:s{sid}',
                    'gold_build': f'{ns}:gold_predictions_ref:r{run_id}'}
    gates_by_node = {}
    for e in g_['edges']:
        gates_by_node.setdefault(e['to_key'], []).append(
            {'gate': e['gate_name'], 'status': e['gate_status'], 'detail': e['gate_detail']})

    steps = []
    for n in g_['nodes']:
        payload = None
        lk = payload_keys.get(n['node_key'])
        if lk:
            node_art = uas.latest_by_logical(get_db(), lk)
            if node_art:
                payload = json.loads(node_art['payload_json'])
        elif n.get('uas_artifact_id'):
            node_art = uas.get_by_uid(get_db(), n['uas_artifact_id'])
            if node_art:
                payload = json.loads(node_art['payload_json'])
        steps.append({'node_key': n['node_key'], 'status': n['status'],
                      'cached': bool(n['cached']), 'prior_run_id': n['prior_run_id'],
                      'content_hash': n['content_hash'],
                      'started_at': n['started_at'], 'completed_at': n['completed_at'],
                      'gates': gates_by_node.get(n['node_key'], []),
                      'uas_payload': payload})
    attempts = [{'cycle': r['cycle'], 'scope': r['scope'],
                 'resolved': bool(r['resolved']),
                 **json.loads(r['detail_json'])}
                for r in many('SELECT * FROM repair_attempts WHERE run_id=? ORDER BY id',
                              (run_id,))]
    return jsonify({'run_id': run_id, 'steps': steps, 'repair_attempts': attempts})


PLANS = {
    'starter':    {'tokens': 100_000,   'public_links': False, 'predictive_models': False,
                   'sso': False, 'rls': False},
    'team':       {'tokens': 1_000_000, 'public_links': True, 'predictive_models': True,
                   'sso': False, 'rls': False},
    'business':   {'tokens': 5_000_000, 'public_links': True, 'predictive_models': True,
                   'sso': True, 'rls': True},
    'enterprise': {'tokens': 50_000_000, 'public_links': True, 'predictive_models': True,
                   'sso': True, 'rls': True},
}


def _current_plan():
    row = one("SELECT value FROM platform_settings WHERE key='plan'")
    return (row['value'] if row and row['value'] in PLANS else 'team')


def _entitled(feature):
    return PLANS[_current_plan()].get(feature, False)


MARK_RULES = {  # §8.3 chart-type → mark mapping, ranked per data shape
    'timeseries_ci': ['line', 'area', 'bar'],
    'forecast': ['area', 'line'],
    'dimension_breakdown': ['bar', 'horizontal_bar', 'heatmap'],
    'feature_importance': ['bar', 'horizontal_bar'],
}


@app.post('/api/workspace/observability_report')
@require_role('admin', 'analyst')
def observability_report():
    """Evo #21: the platform monitors itself as a native artifact."""
    import artifact_gen as ag
    import cache_hier
    import orchestrator
    cache = cache_hier.stats(get_db())
    disp = orchestrator.aggregate(get_db())
    repairs = one('SELECT COUNT(*) AS c FROM repair_attempts')['c']
    events = one('SELECT COUNT(*) AS c FROM platform_events')['c']
    rows = [{'day_index': i, 'date': k, 'actual': v['hits'] + v['misses'],
             'predicted': v['hits'], 'ci_low': 0, 'ci_high': v['hits'] + v['misses'],
             'is_forecast': 0} for i, (k, v) in enumerate(cache.items())]
    aid = execute('INSERT INTO artifacts (title, type, owner, dq_status, is_sandbox) '
                  "VALUES ('Platform Observability', 'Report', ?, 'pass', 0)", (_me(),))
    art_row = one('SELECT * FROM artifacts WHERE id=?', (aid,))
    html = ag.generate_artifact_html(art_row, rows or [
        {'day_index': 0, 'date': 'n/a', 'actual': 0, 'predicted': 0,
         'ci_low': 0, 'ci_high': 0, 'is_forecast': 0}], compute_kpis(rows or [
        {'day_index': 0, 'date': 'n/a', 'actual': 0, 'predicted': 0,
         'ci_low': 0, 'ci_high': 0, 'is_forecast': 0}]))
    section = ('<section id="observability"><h2>Platform telemetry</h2>'
               f'<p>Dispatch cost ${disp["est_cost_total"]} over {disp["count"]} dispatches; '
               f'cache layers: ' + ', '.join(f'{k} {round(v["hit_rate"] * 100)}%'
                                             for k, v in cache.items()) +
               f'; repair attempts {repairs}; platform events {events}.</p></section>')
    html = html.replace('</body>', section + '</body>', 1)
    _obs_att = []
    html, _c, validation = ag.validate_and_repair(html, attempt_log=_obs_att)
    _persist_repair_attempts(get_db(), 'artifact_render', aid, None, _obs_att,
                             validation['status'] == 'PASS')
    _store_artifact_file(aid, 1, html, validation)
    log_action('workspace.observability_report', 'artifact', aid, {})
    return jsonify(one('SELECT * FROM artifacts WHERE id=?', (aid,))), 201


@app.get('/api/metrics/<slug>/benchmarks')
def metric_benchmarks(slug):
    """Evo #13: historical from gold; peer/budget only when registered."""
    hist = one("SELECT AVG(actual) AS mean, MIN(actual) AS lo, MAX(actual) AS hi "
               'FROM gold_predictions WHERE actual IS NOT NULL')
    refs = {r['kind']: {'value': r['value']} for r in
            many('SELECT * FROM benchmark_refs WHERE metric=? ORDER BY id DESC', (slug,))}
    return jsonify({'metric': slug,
                    'historical': {'mean': round(hist['mean'], 2) if hist['mean'] else None,
                                   'min': hist['lo'], 'max': hist['hi']},
                    'peer': refs.get('peer'), 'budget': refs.get('budget'),
                    'seasonal': refs.get('seasonal')})


@app.post('/api/metrics/<slug>/benchmarks')
@require_role('admin', 'analyst')
def register_benchmark(slug):
    b = request.get_json() or {}
    if b.get('kind') not in ('peer', 'budget', 'seasonal'):
        return jsonify({'error': "kind must be peer|budget|seasonal"}), 400
    execute('INSERT INTO benchmark_refs (metric, kind, value) VALUES (?,?,?)',
            (slug, b['kind'], float(b.get('value') or 0)))
    log_action('benchmark.registered', 'metric', slug, {'kind': b['kind']})
    return jsonify({'metric': slug, 'kind': b['kind'], 'value': b.get('value')}), 201


@app.get('/api/artifacts/<int:id>/viz_alternates')
def viz_alternates(id):
    """Evo #31: ranked alternates from the mark-mapping rules + data shape."""
    if not one('SELECT id FROM artifacts WHERE id=?', (id,)):
        return jsonify({'error': 'Artifact not found'}), 404
    section = request.args.get('section', 'timeseries_ci')
    marks = MARK_RULES.get(section, ['line', 'bar'])
    return jsonify({'section': section,
                    'alternates': [{'rank': i + 1, 'mark': m,
                                    'rationale': f'{m} ranks #{i + 1} for {section} by the '
                                                 f'chart-type mapping (§8.3)'}
                                   for i, m in enumerate(marks)]})


@app.post('/api/artifacts/<int:id>/viz_swap')
def viz_swap(id):
    """Evo #31: one-click swap routes through the semantic-edit path."""
    b = request.get_json() or {}
    with app.test_request_context(json={'chart_type': b.get('mark')}):
        pass
    # reuse the section-edit machinery directly
    request_json = {'chart_type': b.get('mark')}
    art = one('SELECT * FROM artifacts WHERE id=?', (id,))
    if not art:
        return jsonify({'error': 'Artifact not found'}), 404
    import requests as _unused  # noqa
    return edit_artifact_section.__wrapped__(id, b.get('section')) if False else         _viz_swap_apply(id, b.get('section'), b.get('mark'))


def _viz_swap_apply(id, sid, mark):
    from flask import jsonify as _j
    art = one('SELECT * FROM artifacts WHERE id=?', (id,))
    layout = json.loads(art.get('layout_json') or '{"sections": []}')
    section = next((s for s in layout['sections'] if s['id'] == sid), None)
    if not section:
        return _j({'error': f'Unknown section {sid!r}'}), 404
    section['mark'] = mark
    execute('UPDATE artifacts SET layout_json=? WHERE id=?', (json.dumps(layout), id))
    import artifact_gen as ag
    rows = many('SELECT * FROM chart_data WHERE pipeline_run_id=? ORDER BY day_index',
                (art['pipeline_run_id'],))
    if rows:
        html = ag.generate_artifact_html(art, rows, compute_kpis(rows))
        _sw = []
        html, _c, validation = ag.validate_and_repair(html, attempt_log=_sw)
        prev = one('SELECT MAX(version) AS v FROM artifact_files WHERE artifact_id=?', (id,))
        _store_artifact_file(id, (prev['v'] or 0) + 1, html, validation)
    log_action('artifact.viz_swapped', 'artifact', id, {'section': sid, 'mark': mark})
    return _j({'artifact_id': id, 'section': sid, 'mark': mark, 'edit_class': 'semantic'})


@app.post('/api/plugins/validators')
@require_role('admin', 'analyst')
def register_plugin_validator():
    """Evo #14: plugin validators join the gate set — sandboxed to a safe
    declarative check (row-count floor on a whitelisted table)."""
    import re as _re
    b = request.get_json() or {}
    name, table = b.get('name', ''), b.get('table', '')
    if not _re.fullmatch(r'[a-z_][a-z0-9_]*', name or '') or        not _re.fullmatch(r'gold_[a-z_]+', table or ''):
        return jsonify({'error': 'name must be snake_case; table must be a gold_* table'}), 400
    execute('INSERT OR REPLACE INTO plugin_validators (name, table_name, min_rows) '
            'VALUES (?,?,?)', (name, table, int(b.get('min_rows') or 1)))
    log_action('plugin.validator_registered', 'plugin', name, {'table': table})
    return jsonify({'name': name, 'table': table}), 201


@app.post('/api/integrations/outbound')
def outbound_action():
    """Evo #35: workflow triggers — console/outbox fallback, audited."""
    b = request.get_json() or {}
    if b.get('kind') not in ('slack', 'teams', 'jira', 'email'):
        return jsonify({'error': 'kind must be slack|teams|jira|email'}), 400
    oid = execute('INSERT INTO outbound_actions (kind, target, message) VALUES (?,?,?)',
                  (b['kind'], b.get('target', ''), b.get('message', '')))
    if b['kind'] == 'email':
        send_email(to=b.get('target', 'admin@acme.com'), subject='AnalytIQ alert',
                   html=f"<p>{b.get('message', '')}</p>")
    log_action('outbound.dispatched', 'outbound_action', oid, {'kind': b['kind']},
               )
    return jsonify(one('SELECT * FROM outbound_actions WHERE id=?', (oid,))), 201


@app.post('/api/templates/package')
@require_role('admin', 'analyst')
def package_template():
    """Evo #33: package a session's plan with parameterized semantic refs."""
    b = request.get_json() or {}
    sess = one('SELECT * FROM sessions WHERE id=?', (b.get('session_id'),))
    if not sess:
        return jsonify({'error': 'Session not found'}), 404
    plan = {'target_metric': '$METRIC', 'grain': sess['grain'],
            'prediction_horizon': sess['horizon'],
            'output_type': 'forecast_dashboard',
            'panels': ['kpi_row', 'timeseries_ci', 'forecast']}
    tid = execute('INSERT INTO dashboard_templates (name, plan_template_json, created_by) '
                  'VALUES (?,?,?)', (b.get('name', f"{sess['metric']} template"),
                                     json.dumps(plan), _me()))
    log_action('template.packaged', 'template', tid, {'session_id': sess['id']})
    return jsonify({'id': tid, 'plan_template': plan}), 201


@app.post('/api/templates/<int:tid>/apply')
def apply_template(tid):
    """Evo #33: applying re-resolves refs and re-runs plan validation."""
    import planner
    tpl = one('SELECT * FROM dashboard_templates WHERE id=?', (tid,))
    if not tpl:
        return jsonify({'error': 'Template not found'}), 404
    b = request.get_json() or {}
    metric = b.get('metric')
    if not metric:
        return jsonify({'error': 'metric required to re-resolve the template'}), 400
    plan = json.loads(tpl['plan_template_json'])
    spec = {'intent': 'predictive', 'intent_confidence': 0.9,
            'analytic_goal': f"Template: {tpl['name']} → {metric}",
            'target_metric': metric, 'feature_candidates': [],
            'date_range': {'start': '2023-01-01', 'end': '2023-12-31'},
            'grain': plan.get('grain', 'Location · Day'),
            'output_type': plan.get('output_type', 'forecast_dashboard'),
            'prediction_horizon': plan.get('prediction_horizon', 14),
            'explores_used': [], 'semantic_layer_version': '1.0.0',
            'governance_manifest_version': '1.0.0'}
    errs = planner.validate_session_spec(spec)
    if errs:
        return jsonify({'error': 'Template failed validation for this context',
                        'validation': {'errors': errs}}), 422
    sid = execute('INSERT INTO sessions (metric, grain, horizon, status) '
                  "VALUES (?, ?, ?, 'pending')",
                  (metric, spec['grain'], spec['prediction_horizon']))
    prev = one('SELECT MAX(spec_version) AS v FROM session_specs WHERE session_id=?', (sid,))
    import hashlib as _hl
    execute('INSERT INTO session_specs (session_id, spec_version, payload_hash, spec_json) '
            'VALUES (?,?,?,?)',
            (sid, (prev['v'] or 0) + 1,
             _hl.sha256(json.dumps(spec, sort_keys=True).encode()).hexdigest(),
             json.dumps(spec)))
    log_action('template.applied', 'template', tid, {'session_id': sid, 'metric': metric})
    return jsonify({'session_id': sid, 'spec': spec, 'validation': {'errors': []}}), 201


@app.get('/api/billing/usage')
def billing_usage():
    """R20S1E1: token metering rollup with soft thresholds + hard cap."""
    plan = _current_plan()
    included = PLANS[plan]['tokens']
    used = one('SELECT COALESCE(SUM(tokens),0) AS t FROM task_dispatches')['t']
    pct = used / included * 100 if included else 0
    status = 'ok'
    for thr in (50, 75, 90):
        if pct >= thr:
            status = f'soft_{thr}'
    if pct >= 90:
        status = 'capped' if pct >= 100 else 'soft_90'
    by_cap = many('SELECT task_kind AS capability, SUM(tokens) AS tokens, COUNT(*) AS calls '
                  'FROM task_dispatches GROUP BY task_kind ORDER BY tokens DESC')
    by_user = many("SELECT COALESCE(workspace_id,'default') AS consumer, SUM(tokens) AS tokens "
                   'FROM task_dispatches GROUP BY workspace_id')
    return jsonify({'plan': plan,
                    'cycle': {'tokens_used': used, 'included': included,
                              'pct': round(pct, 2)},
                    'by_capability': by_cap, 'by_consumer': by_user,
                    'thresholds': {'soft': [50, 75, 90], 'hard_cap_pct': 90,
                                   'status': status},
                    'overage': {'rate_usd_per_100k': 8}})


@app.put('/api/billing/plan')
@require_role('admin')
def put_billing_plan():
    b = request.get_json() or {}
    plan = b.get('plan')
    if plan not in PLANS:
        return jsonify({'error': f'plan must be one of {sorted(PLANS)}'}), 400
    execute("INSERT INTO platform_settings (key, value) VALUES ('plan', ?) "
            "ON CONFLICT(key) DO UPDATE SET value=excluded.value", (plan,))
    log_action('billing.plan_changed', 'platform', None, {'plan': plan})
    return jsonify({'plan': plan, 'entitlements': PLANS[plan]})


@app.post('/api/admin/rls')
@require_role('admin')
def create_rls_policy():
    """R20S1E1: RLS policies — safe expression subset, enforced on gold reads."""
    import re as _re
    b = request.get_json() or {}
    expr = (b.get('expression') or '').strip()
    if not _re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*\s*(=|<|>|<=|>=|!=)\s*('[^']*'|\d+(\.\d+)?)", expr):
        return jsonify({'error': "expression must be a safe comparison (col op literal)"}), 400
    pid = execute('INSERT INTO rls_policies (table_name, expression, status) VALUES (?,?,?)',
                  (b.get('table_name'), expr, b.get('status', 'on')))
    log_action('rls.policy_created', 'rls_policy', pid,
               {'table': b.get('table_name'), 'expression': expr}, )
    return jsonify(one('SELECT * FROM rls_policies WHERE id=?', (pid,))), 201


@app.post('/api/admin/rls/simulate')
def simulate_rls():
    b = request.get_json() or {}
    table = b.get('table_name')
    pol = one("SELECT * FROM rls_policies WHERE table_name=? AND status='on' "
              'ORDER BY id DESC LIMIT 1', (table,))
    if not pol:
        return jsonify({'policy': None, 'visible_rows': None,
                        'note': 'No active policy for this table'})
    n = one(f'SELECT COUNT(*) AS c FROM "{table}" WHERE {pol["expression"]}')['c']
    return jsonify({'policy': dict(pol), 'visible_rows': n,
                    'note': f'A viewer under this policy sees {n} row(s).'})


@app.get('/api/audit-logs/export')
def export_audit():
    fmt = request.args.get('format', 'json')
    if fmt not in ('csv', 'json'):
        return jsonify({'error': "format must be 'csv' or 'json'"}), 400
    rows = many('SELECT id, action, COALESCE(severity, "info") AS severity, resource_type, '
                'resource_id, user_email, created_at FROM audit_logs ORDER BY id DESC LIMIT 1000')
    if fmt == 'json':
        return jsonify({'events': rows})
    import io, csv as _csv
    buf = io.StringIO()
    w = _csv.DictWriter(buf, fieldnames=['id', 'action', 'severity', 'resource_type',
                                         'resource_id', 'user_email', 'created_at'])
    w.writeheader()
    for r in rows:
        w.writerow(r)
    return Response(buf.getvalue(), mimetype='text/csv',
                    headers={'Content-Disposition': 'attachment; filename=audit.csv'})


def _me():
    return getattr(g, 'user_email', None) or 'admin@acme.com'


# ── R22S1E1: workspace home aggregate ──────────────────────────────────────
@app.get('/api/home/summary')
def home_summary():
    """R22S1E1-US1: one call feeds the 8 home widgets (App Home Frame 01).
    Composed strictly from existing tables; usage block is admin-only."""
    import datetime as _dt
    db = get_db()

    email = _me()
    first = (getattr(g, 'auth_user', None) or {}).get('name') or email.split('@')[0]
    first = first.replace('.', ' ').split()[0].capitalize()
    now = _dt.datetime.now()
    date_line = now.strftime('%a · %b %d, %Y · %H:%M').replace(' 0', ' ')

    def _age(iso):
        if not iso:
            return ''
        try:
            t = _dt.datetime.fromisoformat(str(iso).replace(' ', 'T'))
        except ValueError:
            return ''
        mins = max(0, int((_dt.datetime.utcnow() - t).total_seconds() // 60))
        if mins < 60:
            return f'{mins}m ago'
        if mins < 60 * 24:
            return f'{mins // 60}h ago'
        return f'{mins // (60 * 24)}d ago'

    recents = []
    for a in many('SELECT id, title, dq_status, confidence, created_at FROM artifacts '
                  'WHERE COALESCE(is_sandbox,0)=0 ORDER BY id DESC LIMIT 3'):
        warn = (a['dq_status'] or '').lower() not in ('', 'pass', 'ok', 'healthy') \
               or (a['confidence'] is not None and a['confidence'] < 0.5)
        recents.append({'id': a['id'], 'title': a['title'],
                        'health': 'warnings' if warn else 'healthy',
                        'age': _age(a['created_at'])})

    tbl = one('SELECT COUNT(*) AS n, COALESCE(AVG(health_score),100) AS avg, '
              "SUM(CASE WHEN health_score >= 80 THEN 1 ELSE 0 END) AS ok "
              'FROM cataloged_tables')
    alert_counts = {r['type']: r['n'] for r in
                    many('SELECT type, COUNT(*) AS n FROM alerts GROUP BY type')}
    drift_n = sum(n for t, n in alert_counts.items() if 'drift' in t)
    fresh_n = sum(n for t, n in alert_counts.items() if 'fresh' in t)
    pii_n = sum(n for t, n in alert_counts.items() if 'pii' in t)
    health = {
        'score': round(tbl['avg'] or 100),
        'rows': [
            {'label': 'Sources healthy', 'value': f"{tbl['ok'] or 0}/{tbl['n'] or 0}"},
            {'label': 'Freshness SLAs', 'value': 'met' if not fresh_n else f'{fresh_n} breached'},
            {'label': 'Schema drift', 'value': f'{drift_n} table' + ('' if drift_n == 1 else 's')},
            {'label': 'PII flags', 'value': f'{pii_n} open'},
        ],
    }

    runs = []
    for r in many('SELECT p.id, p.current_step, p.started_at, s.metric FROM pipeline_runs p '
                  'LEFT JOIN sessions s ON s.id = p.session_id '
                  'WHERE p.completed_at IS NULL ORDER BY p.id DESC LIMIT 3'):
        runs.append({'id': r['id'], 'title': r['metric'] or f"Run {r['id']}",
                     'stage': r['current_step'] or 0, 'total': 9,
                     'started_at': r['started_at']})

    sev = {'revenue': 'HIGH', 'anomaly': 'MED', 'drift': 'MED'}
    alerts = [{'severity': sev.get(a['type'], 'MED') if 'fresh' not in a['type'] else 'LOW',
               'message': a['subject'] or a['type'], 'age': _age(a['created_at'])}
              for a in many('SELECT * FROM alerts ORDER BY id DESC LIMIT 3')]

    pend = many("SELECT id, type, definition FROM semantic_definitions "
                "WHERE status='pending' ORDER BY id DESC LIMIT 3")
    chip = {'metric': 'DEF', 'dimension': 'DEF', 'pii': 'PII', 'drift': 'DRIFT'}
    review = {
        'count': one("SELECT COUNT(*) AS n FROM semantic_definitions WHERE status='pending'")['n'],
        'items': [{'id': p['id'],
                   'label': (p['definition'] or p['type'] or '')[:60],
                   'chip': chip.get((p['type'] or '').lower(), 'DEF')} for p in pend],
    }

    suggested = [{'id': o['id'], 'prompt': o['question'] or o['headline']}
                 for o in many("SELECT id, headline, question FROM opportunities "
                               "WHERE status='pending' ORDER BY id DESC LIMIT 3")]

    viewed = [{'id': r['id'], 'title': r['title'], 'age': _age(r['created_at'])}
              for r in many('SELECT id, title, created_at FROM artifacts '
                            'WHERE COALESCE(is_sandbox,0)=0 ORDER BY id DESC LIMIT 3')]

    usage = None
    if current_role() in ('admin', 'owner'):
        used = one('SELECT COALESCE(SUM(tokens),0) AS t FROM task_dispatches')['t']
        included = PLANS[_current_plan()]['tokens']
        # R31S2E2: 7-day series + w/w delta for the frame's usage mini chart
        daily = [one("SELECT COALESCE(SUM(tokens),0) AS t FROM task_dispatches "
                     "WHERE date(created_at) = date('now', ?)", (f'-{d} days',))['t']
                 for d in range(6, -1, -1)]
        this_wk = one("SELECT COALESCE(SUM(tokens),0) AS t FROM task_dispatches "
                      "WHERE created_at >= datetime('now', '-7 days')")['t']
        last_wk = one("SELECT COALESCE(SUM(tokens),0) AS t FROM task_dispatches "
                      "WHERE created_at >= datetime('now', '-14 days') "
                      "AND created_at < datetime('now', '-7 days')")['t']
        wow = round((this_wk - last_wk) / last_wk * 100) if last_wk else None
        usage = {'tokens_used': used,
                 'pct': round(used / included * 100, 2) if included else 0,
                 'daily': daily, 'wow_delta': wow}

    return jsonify({'greeting': f'Good morning, {first}', 'date_line': date_line,
                    'recents': recents, 'health': health, 'runs': runs,
                    'alerts': alerts, 'review': review, 'suggested': suggested,
                    'recently_viewed': viewed, 'usage': usage})


@app.get('/api/activity')
def activity_feed():
    """R31S2E1-US1: typed activity projection over audit_logs (kind buckets,
    cursor pagination) for /app/activity."""
    import activity as act
    kind = request.args.get('kind') or None
    cursor = request.args.get('cursor') or None
    try:
        limit = max(1, min(100, int(request.args.get('limit', 20))))
    except ValueError:
        return jsonify({'error': 'limit must be an integer'}), 400
    return jsonify(act.project(get_db(), kind=kind, cursor=cursor, limit=limit))


@app.get('/api/notifications')
def list_notifications():
    """R18S1E1: in-app inbox with unread count."""
    me = _me()
    rows = many('SELECT * FROM notifications WHERE user_id=? ORDER BY id DESC LIMIT 100', (me,))
    unread = one('SELECT COUNT(*) AS c FROM notifications WHERE user_id=? AND read=0', (me,))['c']
    return jsonify({'notifications': rows, 'unread': unread})


@app.post('/api/notifications/<int:nid>/read')
def read_notification(nid):
    execute('UPDATE notifications SET read=1 WHERE id=? AND user_id=?', (nid, _me()))
    return jsonify({'ok': True})


@app.post('/api/notifications/read_all')
def read_all_notifications():
    execute('UPDATE notifications SET read=1 WHERE user_id=?', (_me(),))
    return jsonify({'ok': True})


@app.get('/api/workspace/activity')
def workspace_activity():
    """R18S1E1: unified workspace activity feed joined from audit events."""
    KIND_MAP = {'pipeline.': 'build', 'artifact.': 'build', 'governance': 'governance',
                'semantic.': 'governance', 'uas.': 'build', 'share': 'sharing',
                'embed': 'sharing', 'auth.': 'team', 'memory.': 'team',
                'opportunity.': 'insight', 'insights.': 'insight'}
    rows = many('SELECT action, user_email, resource_type, resource_id, created_at '
                'FROM audit_logs ORDER BY id DESC LIMIT ?',
                (min(200, request.args.get('limit', 50, type=int)),))
    events = []
    for r in rows:
        kind = next((v for k, v in KIND_MAP.items() if (r['action'] or '').startswith(k)), 'other')
        events.append({'kind': kind, 'summary': r['action'],
                       'actor': r['user_email'] or 'system',
                       'resource': f"{r['resource_type']}:{r['resource_id']}",
                       'created_at': r['created_at']})
    want = request.args.get('kind')
    if want:
        events = [e for e in events if e['kind'] == want]
    return jsonify({'events': events})


@app.post('/api/artifacts/<int:id>/comments')
def post_comment(id):
    """R18S1E2: threaded, section-anchored comments with @mention fan-out."""
    import notifications as nt
    if not one('SELECT id FROM artifacts WHERE id=?', (id,)):
        return jsonify({'error': 'Artifact not found'}), 404
    b = request.get_json() or {}
    if not (b.get('body') or '').strip():
        return jsonify({'error': 'body required'}), 400
    cid = execute('INSERT INTO comments (artifact_id, parent_id, section_id, author, body) '
                  'VALUES (?,?,?,?,?)',
                  (id, b.get('parent_id'), b.get('section_id'), _me(), b['body'].strip()))
    nt.fan_out_mentions(get_db(), b['body'], f'/app/artifacts/{id}', _me())
    log_action('comment.created', 'artifact', id, {'comment_id': cid,
                                                   'section': b.get('section_id')})
    return jsonify(one('SELECT * FROM comments WHERE id=?', (cid,))), 201


@app.get('/api/artifacts/<int:id>/comments')
def list_comments(id):
    rows = many('SELECT * FROM comments WHERE artifact_id=? ORDER BY id', (id,))
    roots = [dict(r, replies=[]) for r in rows if not r['parent_id']]
    by_id = {r['id']: r for r in roots}
    for r in rows:
        if r['parent_id'] and r['parent_id'] in by_id:
            by_id[r['parent_id']]['replies'].append(r)
    return jsonify({'comments': roots})


@app.post('/api/comments/<int:cid>/resolve')
def resolve_comment(cid):
    if not one('SELECT id FROM comments WHERE id=?', (cid,)):
        return jsonify({'error': 'Not found'}), 404
    execute('UPDATE comments SET resolved=1 WHERE id=?', (cid,))
    return jsonify({'id': cid, 'resolved': True})


@app.get('/api/comments/inbox')
def comments_inbox():
    tab = request.args.get('tab', 'open')
    me = _me()
    if tab == 'resolved':
        rows = many('SELECT * FROM comments WHERE resolved=1 ORDER BY id DESC LIMIT 100')
    elif tab == 'mentioned':
        rows = many("SELECT * FROM comments WHERE body LIKE ? ORDER BY id DESC LIMIT 100",
                    (f'%@{me}%',))
    else:
        rows = many('SELECT * FROM comments WHERE resolved=0 ORDER BY id DESC LIMIT 100')
    return jsonify({'comments': rows})


SEAT_TOTAL = 25


@app.post('/api/team/invites')
@require_role('admin', 'analyst')
def create_invites():
    """R18S1E2: invitation lifecycle with seat accounting."""
    import secrets as _secrets
    b = request.get_json() or {}
    emails = [e.strip() for e in (b.get('emails') or []) if e.strip()]
    if not emails:
        return jsonify({'error': 'emails required'}), 400
    used = one("SELECT COUNT(*) AS c FROM users")['c'] +         one("SELECT COUNT(*) AS c FROM team_invites WHERE status='pending'")['c']
    if used + len(emails) > SEAT_TOTAL:
        return jsonify({'error': f'Seat limit reached ({SEAT_TOTAL})'}), 409
    out = []
    for email in emails:
        token = _secrets.token_urlsafe(24)
        execute('INSERT INTO team_invites (email, role, token, invited_by) VALUES (?,?,?,?)',
                (email, b.get('role', 'analyst'), token, _me()))
        send_email(to=email, subject="You're invited to AnalytIQ",
                   html=f'<p>{_me()} invited you. Accept: /api/team/invites/{token}/accept</p>')
        out.append({'email': email, 'token': token})
    log_action('team.invited', 'team', None, {'count': len(out)})
    return jsonify({'invites': out}), 201


@app.post('/api/team/invites/<token>/accept')
def accept_invite(token):
    import authn
    inv = one('SELECT * FROM team_invites WHERE token=?', (token,))
    if not inv:
        return jsonify({'error': 'Invalid invite'}), 404
    if inv['status'] != 'pending':
        return jsonify({'error': 'Invite already used'}), 410
    b = request.get_json() or {}
    pw = b.get('password') or ''
    if len(pw) < 8:
        return jsonify({'error': 'password of 8+ chars required'}), 400
    if not one('SELECT id FROM users WHERE email=?', (inv['email'],)):
        execute('INSERT INTO users (email, password_hash, role) VALUES (?,?,?)',
                (inv['email'], authn.hash_password(pw), inv['role']))
    execute("UPDATE team_invites SET status='accepted', accepted_at=datetime('now') "
            'WHERE id=?', (inv['id'],))
    log_action('team.joined', 'team', inv['id'], {'email': inv['email']})
    return jsonify({'email': inv['email'], 'role': inv['role']})


@app.get('/api/team/roster')
def team_roster():
    members = [{'email': u['email'], 'role': u['role'], 'status': 'active',
                'joined_at': u['created_at']} for u in
               many('SELECT email, role, created_at FROM users ORDER BY id')]
    members += [{'email': i['email'], 'role': i['role'], 'status': 'invited',
                 'joined_at': i['created_at']} for i in
                many("SELECT * FROM team_invites WHERE status='pending'")]
    used = len(members)
    return jsonify({'members': members,
                    'seats': {'used': used, 'total': SEAT_TOTAL}})


@app.get('/api/pipeline/<int:run_id>/contracts')
def pipeline_contracts(run_id):
    """R17S1E1: the run's per-component query + data contracts."""
    if not one('SELECT id FROM pipeline_runs WHERE id=?', (run_id,)):
        return jsonify({'error': 'Run not found'}), 404
    qcs = many('SELECT * FROM query_contracts WHERE run_id=? ORDER BY id', (run_id,))
    dcs = many('SELECT * FROM component_data_contracts WHERE run_id=? ORDER BY id', (run_id,))
    for r in dcs:
        r['actual_columns'] = json.loads(r.pop('actual_columns_json'))
        r['numeric_ranges'] = json.loads(r.pop('numeric_ranges_json'))
    for r in qcs:
        r['expected_columns'] = json.loads(r.pop('expected_columns_json'))
    return jsonify({'query_contracts': qcs, 'data_contracts': dcs})


@app.get('/api/gold/catalog')
def gold_catalog():
    """R17S1E1: workspace-wide gold catalog."""
    out = []
    for table in ('gold_predictions', 'gold_forecast'):
        for r in many(f'SELECT pipeline_run_id AS run_id, session_id, COUNT(*) AS row_count '
                      f'FROM {table} GROUP BY pipeline_run_id ORDER BY pipeline_run_id DESC LIMIT 100'):
            run = one('SELECT status, mape FROM pipeline_runs WHERE id=?', (r['run_id'],))
            sess = one('SELECT metric FROM sessions WHERE id=?', (r['session_id'],))
            out.append({'table': table, 'run_id': r['run_id'], 'session_id': r['session_id'],
                        'metric': sess['metric'] if sess else None,
                        'row_count': r['row_count'],
                        'gate_status': 'PASS' if run and run['status'] == 'done' else 'WARN',
                        'mape': run['mape'] if run else None})
    return jsonify({'tables': out})


@app.get('/api/pipeline/<int:run_id>/dag')
def pipeline_dag(run_id):
    """R8S2E3: execution graph — nodes with cached flags, edges with gate results."""
    import dag
    g = dag.graph(get_db(), run_id)
    if g is None:
        return jsonify({'error': 'No DAG recorded for this run'}), 404
    return jsonify(g)


@app.post('/api/artifacts/<int:id>/monitor')
def monitor_artifact(id):
    """R12S2E4: model monitoring resolved from an artifact's lineage."""
    import model_monitor as mm
    art = one('SELECT * FROM artifacts WHERE id=?', (id,))
    if not art or not art.get('pipeline_run_id'):
        return jsonify({'error': 'Artifact not found'}), 404
    run = one('SELECT * FROM pipeline_runs WHERE id=?', (art['pipeline_run_id'],))
    out = mm.check(get_db(), run['session_id']) if run else None
    if out is None:
        return jsonify({'error': 'No session lineage'}), 404
    log_action('model.monitored', 'artifact', id, {'triggers': out['triggers']})
    return jsonify(out)


@app.post('/api/models/<int:session_id>/monitor')
def monitor_model(session_id):
    """R12S2E4: continuous model monitoring — importance + input drift."""
    import model_monitor as mm
    out = mm.check(get_db(), session_id)
    if out is None:
        return jsonify({'error': 'Session not found'}), 404
    log_action('model.monitored', 'session', session_id,
               {'triggers': out['triggers']})
    return jsonify(out)


@app.post('/api/platform/self_improve')
def run_self_improve():
    """R12S2E3: run the self-improvement miner now (also a job kind)."""
    import self_improve
    n = self_improve.mine(get_db())
    log_action('platform.self_improved', 'platform', None, {'signals': n})
    return jsonify({'new_signals': n})


@app.get('/api/platform/signals')
def list_platform_signals():
    q, args = 'SELECT * FROM platform_signals WHERE 1=1', []
    for field, col in (('consumer', 'consumer'), ('kind', 'signal_kind')):
        if request.args.get(field):
            q += f' AND {col}=?'; args.append(request.args[field])
    q += ' ORDER BY id DESC LIMIT 100'
    return jsonify({'signals': many(q, tuple(args))})


@app.get('/api/platform/feedback')
def platform_feedback():
    """R12S1E2: acceptance rates per recommendation type."""
    import feedback_loop as fb
    return jsonify({'types': fb.acceptance_rates(get_db())})


@app.post('/api/feedback')
def post_feedback():
    """R12S1E2: generic feedback recording (benchmarks, suggestions, …)."""
    import feedback_loop as fb
    b = request.get_json() or {}
    try:
        fb.record(get_db(), b.get('rec_type', ''), b.get('rec_id'),
                  b.get('decision', ''), category=b.get('category'),
                  user=getattr(g, 'user_email', None) or 'default',
                  signal=b.get('signal'))
    except ValueError as e:
        return jsonify({'error': str(e)}), 400
    return jsonify({'recorded': True}), 201


@app.get('/api/artifacts/<int:id>/opportunities')
def list_opportunities(id):
    if not one('SELECT id FROM artifacts WHERE id=?', (id,)):
        return jsonify({'error': 'Artifact not found'}), 404
    rows = many('SELECT * FROM opportunities WHERE artifact_id=? ORDER BY id DESC', (id,))
    for r in rows:
        r['detail'] = json.loads(r.pop('detail_json') or '{}')
    return jsonify({'opportunities': rows})


@app.post('/api/artifacts/<int:id>/opportunities/evaluate')
def evaluate_opportunities(id):
    """R12S1E1: re-run the opportunity evaluation for an artifact."""
    import opportunity
    ins, err = _compute_artifact_insights(id)
    n = opportunity.evaluate(get_db(), id, ins or [])
    log_action('opportunity.evaluated', 'artifact', id, {'new': n})
    return jsonify({'new_opportunities': n})


def _decide_opportunity(oid, decision):
    row = one('SELECT * FROM opportunities WHERE id=?', (oid,))
    if not row:
        return jsonify({'error': 'Opportunity not found'}), 404
    if row['status'] != 'open':
        return jsonify({'error': f"Opportunity already {row['status']}"}), 409
    execute("UPDATE opportunities SET status=?, decided_at=datetime('now') WHERE id=?",
            (decision, oid))
    # R12S1E2: every decision feeds the recommendation feedback loop
    import feedback_loop as fb
    _detail = json.loads(row['detail_json'] or '{}')
    fb.record(get_db(), 'opportunity', oid,
              'accept' if decision == 'accepted' else 'dismiss',
              category=row['kind'],
              user=getattr(g, 'user_email', None) or 'default',
              signal=_detail.get('edge_weight'))
    out = {'id': oid, 'status': decision}
    detail = json.loads(row['detail_json'] or '{}')
    if decision == 'accepted' and row['kind'] in ('forecast_gap', 'causal_candidate'):
        # route to Stage 7 by creating a pre-seeded session — the user still
        # confirms the run; nothing generates without confirmation (§17.4.1)
        metric = (detail.get('metric') or detail.get('related_metric') or
                  'Net Revenue').replace('_', ' ').title()
        sid = execute('INSERT INTO sessions (metric, grain, horizon, status) '
                      "VALUES (?, 'Location · Day', 14, 'pending')", (metric,))
        out['session_id'] = sid
    log_action(f'opportunity.{decision}', 'opportunity', oid,
               {'kind': row['kind'], 'artifact_id': row['artifact_id']})
    return jsonify(out)


@app.post('/api/opportunities/<int:oid>/accept')
def accept_opportunity(oid):
    return _decide_opportunity(oid, 'accepted')


@app.post('/api/opportunities/<int:oid>/dismiss')
def dismiss_opportunity(oid):
    return _decide_opportunity(oid, 'dismissed')


@app.get('/api/artifacts/<int:id>/narrative')
def artifact_narrative(id):
    """R19S1E1 / Evo #25: grounded, audience-tailored narrative."""
    import narrative
    out = narrative.generate(get_db(), id, request.args.get('audience', 'executive'))
    if out is None:
        return jsonify({'error': 'Artifact not found'}), 404
    return jsonify(out)


@app.post('/api/artifacts/<int:id>/duplicate')
def duplicate_artifact(id):
    """R19S1E1: duplicate keeps lineage; new artifact row + copied layout."""
    art = one('SELECT * FROM artifacts WHERE id=?', (id,))
    if not art:
        return jsonify({'error': 'Artifact not found'}), 404
    nid = execute('INSERT INTO artifacts (title, type, mape, owner, dq_status, '
                  'pipeline_run_id, is_sandbox, layout_json, confidence, confidence_json) '
                  'VALUES (?,?,?,?,?,?,?,?,?,?)',
                  (f"{art['title']} (copy)", art['type'], art['mape'], _me(),
                   art['dq_status'], art['pipeline_run_id'], art['is_sandbox'],
                   art.get('layout_json'), art.get('confidence'), art.get('confidence_json')))
    f = one('SELECT * FROM artifact_files WHERE artifact_id=? ORDER BY version DESC LIMIT 1', (id,))
    if f:
        _store_artifact_file(nid, 1, f['html'], json.loads(f['validator_json'] or '{}'))
    log_action('artifact.duplicated', 'artifact', nid, {'source': id})
    return jsonify(one('SELECT * FROM artifacts WHERE id=?', (nid,))), 201


@app.get('/api/artifacts/<int:id>/export/pdf')
def export_artifact_pdf(id):
    """R19S1E1: demo-grade one-page PDF summary (pure-python writer)."""
    import narrative
    out = narrative.generate(get_db(), id, 'executive')
    if out is None:
        return jsonify({'error': 'Artifact not found'}), 404
    art = one('SELECT * FROM artifacts WHERE id=?', (id,))
    lines = [f"AnalytIQ — {art['title']}", '',
             out['narrative'], '',
             f"History rows: {out['facts']['history_days']} · Forecast days: "
             f"{out['facts']['forecast_days']} · MAPE: {out['facts']['mape']}%"]
    def esc(s):
        return s.replace('\\', r'\\').replace('(', r'\(').replace(')', r'\)')
    content = 'BT /F1 11 Tf 40 760 Td 14 TL\n' + '\n'.join(
        f'({esc(l[:110])}) Tj T*' for l in lines) + '\nET'
    objs = ['<< /Type /Catalog /Pages 2 0 R >>',
            '<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
            '<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] '
            '/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>',
            f'<< /Length {len(content)} >>\nstream\n{content}\nendstream',
            '<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>']
    pdf, offsets = '%PDF-1.4\n', []
    for i, o in enumerate(objs, 1):
        offsets.append(len(pdf.encode('latin1', 'replace')))
        pdf += f'{i} 0 obj\n{o}\nendobj\n'
    xref_at = len(pdf.encode('latin1', 'replace'))
    pdf += f'xref\n0 {len(objs) + 1}\n0000000000 65535 f \n'
    pdf += ''.join(f'{off:010d} 00000 n \n' for off in offsets)
    pdf += f'trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF'
    return Response(pdf.encode('latin1', 'replace'), mimetype='application/pdf',
                    headers={'Content-Disposition': f'attachment; filename=artifact_{id}.pdf'})


@app.get('/api/artifacts/<int:id>/export/png')
def export_artifact_png(id):
    """R19S1E1: PNG chart snapshot rendered with Pillow from real rows."""
    from PIL import Image, ImageDraw
    art = one('SELECT * FROM artifacts WHERE id=?', (id,))
    if not art or not art['pipeline_run_id']:
        return jsonify({'error': 'Artifact not found'}), 404
    rows = many('SELECT * FROM chart_data WHERE pipeline_run_id=? ORDER BY day_index',
                (art['pipeline_run_id'],))
    W, H = 900, 360
    img = Image.new('RGB', (W, H), '#ffffff')
    d = ImageDraw.Draw(img)
    d.text((20, 12), f"AnalytIQ · {art['title']}", fill='#0f172a')
    vals = [r['actual'] if r['actual'] is not None else r['predicted'] for r in rows]
    if vals:
        lo, hi = min(vals), max(vals)
        pts = [(30 + i * (W - 60) / max(1, len(vals) - 1),
                H - 40 - (v - lo) / (hi - lo or 1) * (H - 90)) for i, v in enumerate(vals)]
        d.line(pts, fill='#2563eb', width=2)
    import io
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    return Response(buf.getvalue(), mimetype='image/png',
                    headers={'Content-Disposition': f'attachment; filename=artifact_{id}.png'})


@app.get('/api/artifacts/<int:id>/health')
def artifact_health(id):
    """R11S2E5: per-dashboard health breakdown."""
    import dashboard_health as dh
    out = dh.score(get_db(), id)
    if out is None:
        return jsonify({'error': 'Artifact not found'}), 404
    return jsonify(out)


@app.get('/api/workspace/dashboard_health')
def workspace_dashboard_health():
    """R11S2E5: admin rollup across all production dashboards."""
    import dashboard_health as dh
    arts = many('SELECT id, title FROM artifacts WHERE is_sandbox=0 ORDER BY id DESC LIMIT 50')
    scored = []
    for a in arts:
        s = dh.score(get_db(), a['id'])
        if s:
            scored.append({'artifact_id': a['id'], 'title': a['title'],
                           'score': s['score'], 'components': s['components']})
    avg = round(sum(s['score'] for s in scored) / len(scored)) if scored else None
    return jsonify({'average': avg, 'artifacts': scored})


LAYOUT_FIELDS = {'title', 'position'}
SEMANTIC_FIELDS = {'chart_type', 'top_n'}


@app.patch('/api/artifacts/<int:id>/sections/<sid>')
def edit_artifact_section(id, sid):
    """R16S2E4 / Evolution #32: conversational-canvas edits. Layout-only
    changes apply deterministically; semantic changes re-render through the
    validated assembly path. Every edit is versioned, never destructive."""
    art = one('SELECT * FROM artifacts WHERE id=?', (id,))
    if not art:
        return jsonify({'error': 'Artifact not found'}), 404
    layout = json.loads(art.get('layout_json') or '{"sections": []}')
    section = next((s for s in layout['sections'] if s['id'] == sid), None)
    if not section:
        return jsonify({'error': f'Unknown section {sid!r}'}), 404
    b = request.get_json() or {}
    fields = set(b.keys())
    if not fields or not fields <= (LAYOUT_FIELDS | SEMANTIC_FIELDS):
        return jsonify({'error': f'Editable fields: {sorted(LAYOUT_FIELDS | SEMANTIC_FIELDS)}'}), 400

    edit_class = 'semantic' if fields & SEMANTIC_FIELDS else 'layout'
    if 'title' in b:
        section['title'] = str(b['title'])[:120]
    if 'position' in b:
        new_pos = int(b['position'])
        for s in layout['sections']:
            if s['position'] == new_pos:
                s['position'] = section['position']
        section['position'] = new_pos
    if 'chart_type' in b:
        section['mark'] = str(b['chart_type'])
    if 'top_n' in b:
        section['top_n'] = int(b['top_n']) if b['top_n'] is not None else None
    layout['sections'].sort(key=lambda s: s['position'])
    execute('UPDATE artifacts SET layout_json=? WHERE id=?', (json.dumps(layout), id))

    # versioned in the store — history stays replayable/diffable
    import uas
    ns = 'sandbox:default' if art.get('is_sandbox') else 'default'
    uas.register(get_db(), 'artifact_layout', layout,
                 logical_key=f'{ns}:artifact_layout:a{id}', workspace_id=ns,
                 agent='canvas_editor', run_id=art.get('pipeline_run_id'))

    rendered_version = None
    if edit_class == 'semantic' and art.get('pipeline_run_id'):
        # re-route through the validated assembly path (§17.6.3)
        import artifact_gen as ag
        rows = many('SELECT * FROM chart_data WHERE pipeline_run_id=? ORDER BY day_index',
                    (art['pipeline_run_id'],))
        if rows:
            html = ag.generate_artifact_html(art, rows, compute_kpis(rows))
            _att = []
            html, _c, validation = ag.validate_and_repair(html, attempt_log=_att)
            _persist_repair_attempts(get_db(), 'artifact_render', id,
                                     art.get('pipeline_run_id'), _att,
                                     validation['status'] == 'PASS')
            prev = one('SELECT MAX(version) AS v FROM artifact_files WHERE artifact_id=?', (id,))
            rendered_version = (prev['v'] or 0) + 1
            _store_artifact_file(id, rendered_version, html, validation)
    log_action('artifact.edited', 'artifact', id,
               {'section': sid, 'fields': sorted(fields), 'class': edit_class})
    return jsonify({'artifact_id': id, 'section': sid, 'edit_class': edit_class,
                    'layout': layout, 'rendered_version': rendered_version})


@app.get('/api/artifacts/<int:id>/explain')
def explain_artifact(id):
    """R11S1E1: the explain affordance — pure composition over data the
    system already produced (§17.5.1). Lineage from the UAS/DAG, exact SQL
    from the data modeler, semantic definitions, viz field bindings; model
    card + gates + top features when a prediction backs the artifact."""
    import uas
    art = one('SELECT * FROM artifacts WHERE id=?', (id,))
    if not art:
        return jsonify({'error': 'Artifact not found'}), 404
    run = one('SELECT * FROM pipeline_runs WHERE id=?', (art['pipeline_run_id'],))         if art.get('pipeline_run_id') else None
    sess = one('SELECT * FROM sessions WHERE id=?', (run['session_id'],)) if run else None

    # lineage — source tables (manifest when connected), gold tables, UAS chain
    source_tables = []
    if sess and sess.get('connection_id'):
        mrow = _manifest_row(sess['connection_id'])
        if mrow:
            source_tables = [t['name'] for t in
                             json.loads(mrow['manifest_json']).get('tables', [])]
    gold_tables = []
    if run:
        for t in ('gold_predictions', 'gold_forecast'):
            if one(f'SELECT 1 AS x FROM {t} WHERE pipeline_run_id=? LIMIT 1', (run['id'],)):
                gold_tables.append(t)
    ns = _uas_ns(sess) if sess else 'default'
    terminal = uas.latest_by_logical(get_db(), f'{ns}:artifact_html_ref:a{id}')
    chain = uas.provenance_chain(get_db(), terminal['artifact_uid']) if terminal else []

    # exact generated SQL from the data-modeler output
    gt = one('SELECT * FROM gold_tables WHERE session_id=? ORDER BY id DESC LIMIT 1',
             (sess['id'],)) if sess else None
    sql_section = {'ddl': gt['ddl'] if gt else None,
                   'insert_sql': gt['insert_sql'] if gt else None,
                   'gold_read_example': f"GET /api/gold/default/gold_predictions"
                                        f"?filter_col=pipeline_run_id&filter_val={run['id']}"
                                        if run else None}

    # semantic definitions for the artifact's metric
    semantic = []
    if sess:
        from modeler import _slug as _mslug
        mslug = _mslug(sess['metric'])
        for d in many("SELECT name, type, definition, confidence, status FROM semantic_definitions"):
            if _mslug(d['name']) == mslug:
                semantic.append(d)

    # field bindings from the registered viz specs (incl. consulted format)
    spec_node = uas.latest_by_logical(get_db(), f"{ns}:vega_lite_specs:s{sess['id']}") if sess else None
    panels, metric_format = [], None
    if spec_node:
        payload = json.loads(spec_node['payload_json'])
        panels = payload.get('specs', [])
        metric_format = payload.get('metric_format')
    component = request.args.get('component')
    if component:
        panels = [p for p in panels if p.get('panel') == component]

    # model section for predictive artifacts
    card = one('SELECT * FROM model_cards WHERE session_id=? ORDER BY id DESC LIMIT 1',
               (sess['id'],)) if sess else None
    model = None
    if card:
        model = {'model_card_id': card['id'], 'algorithm': card['algorithm'],
                 'status': card['status'],
                 'metrics': json.loads(card['metrics_json'] or '{}'),
                 'gates': json.loads(card['gates_json'] or '{}'),
                 'top_features': json.loads(card['metrics_json'] or '{}').get('top_features')
                                 or json.loads(card['lineage_json'] or '{}').get('top_features')}

    log_action('artifact.explained', 'artifact', id, {'component': component})
    return jsonify({
        'artifact_id': id,
        'lineage': {'run_id': run['id'] if run else None,
                    'source_tables': source_tables,
                    'gold_tables': gold_tables,
                    'provenance_chain': [{'artifact_type': c['artifact_type'],
                                          'version': c['version'],
                                          'content_hash': c['content_hash']} for c in chain]},
        'sql': sql_section,
        'semantic': semantic,
        'field_bindings': {'panels': panels, 'metric_format': metric_format},
        'model': model,
        'confidence': json.loads(art['confidence_json']) if art.get('confidence_json') else None,
        'intent_confidence': None if not sess else _latest_spec_confidence(sess['id']),
    })


def _latest_spec_confidence(session_id):
    row = one('SELECT spec_json FROM session_specs WHERE session_id=? '
              'ORDER BY spec_version DESC LIMIT 1', (session_id,))
    return json.loads(row['spec_json']).get('intent_confidence') if row else None


@app.post('/api/artifacts/<int:id>/promote')
@require_role('admin', 'analyst')
def promote_sandbox_artifact(id):
    """R9S2E6: promote a sandbox artifact to production. Re-runs the full
    deterministic gate set against its run — sandbox mode is never a
    governance bypass (§17.2.8)."""
    import dag
    art = one('SELECT * FROM artifacts WHERE id=?', (id,))
    if not art:
        return jsonify({'error': 'Artifact not found'}), 404
    if not art.get('is_sandbox'):
        return jsonify({'error': 'Artifact is already in production'}), 409
    run = one('SELECT * FROM pipeline_runs WHERE id=?', (art['pipeline_run_id'],))
    sess = one('SELECT * FROM sessions WHERE id=?', (run['session_id'],)) if run else None
    if not sess:
        return jsonify({'error': 'No session lineage for this artifact'}), 409

    ctx = {'session': dict(sess)}
    results = []
    for (f, t) in dag.EDGES:
        for gate_name, fn in dag.DAG_EDGE_GATES.get((f, t), []):
            status, detail = fn(get_db(), art['pipeline_run_id'], ctx)
            results.append({'edge': f'{f}->{t}', 'gate': gate_name,
                            'status': status, 'detail': detail})
    if any(g['status'] == 'BLOCK' for g in results):
        log_action('sandbox.promotion_blocked', 'artifact', id,
                   {'gates': [g['gate'] for g in results if g['status'] == 'BLOCK']})
        return jsonify({'error': 'Promotion blocked by validation gates',
                        'gates': results}), 409

    execute('UPDATE artifacts SET is_sandbox=0 WHERE id=?', (id,))
    file_row = one('SELECT * FROM artifact_files WHERE artifact_id=? '
                   'ORDER BY version DESC LIMIT 1', (id,))
    file_info = {'sha256': file_row.get('sha256'), 'size_bytes': file_row.get('size_bytes'),
                 'validation_status': 'pass', 'version': file_row.get('version')} if file_row else None
    import uas
    gov, sem = _uas_context_versions(get_db(), sess.get('connection_id'))
    uas.register(get_db(), 'artifact_html_ref',
                 {'artifact_id': id, 'title': art['title'], 'promoted_from': 'sandbox',
                  **({'sha256': file_info['sha256']} if file_info else {})},
                 logical_key=f'default:artifact_html_ref:a{id}',
                 gov_version=gov, sem_version=sem, agent='artifact_assembler',
                 workspace_id='default', run_id=art['pipeline_run_id'])
    import search as search_mod
    search_mod.index_artifact(get_db(), one('SELECT * FROM artifacts WHERE id=?', (id,)))
    log_action('sandbox.promoted', 'artifact', id, {'gates_passed': len(results)})
    return jsonify({'promoted': True, 'gates': results})


@app.get('/api/artifacts/<int:id>/provenance')
def artifact_provenance(id):
    import uas
    if not one('SELECT id FROM artifacts WHERE id=?', (id,)):
        return jsonify({'error': 'Artifact not found'}), 404
    art_row = one('SELECT is_sandbox FROM artifacts WHERE id=?', (id,))
    prefix = 'sandbox:default' if (art_row or {}).get('is_sandbox') else 'default'
    terminal = uas.latest_by_logical(get_db(), f'{prefix}:artifact_html_ref:a{id}')
    if not terminal:
        return jsonify({'chain': []})
    chain = uas.provenance_chain(get_db(), terminal['artifact_uid'])
    import dag
    art = one('SELECT pipeline_run_id FROM artifacts WHERE id=?', (id,))
    g = dag.graph(get_db(), art['pipeline_run_id']) if art and art.get('pipeline_run_id') else None
    return jsonify({'chain': [{'artifact_uid': r['artifact_uid'],
                               'artifact_type': r['artifact_type'],
                               'version': r['version'],
                               'content_hash': r['content_hash'],
                               'created_by_agent': r['created_by_agent'],
                               'created_at': r['created_at']} for r in chain],
                    'dag': g})


# ─────────────────────────────────────────────────────────
# Client bundle serving (zero-key one-process boot)
# Serves client/dist when built. /api/* never falls through here.
# ─────────────────────────────────────────────────────────
_CLIENT_DIST = Path(__file__).resolve().parent.parent / 'client' / 'dist'


@app.get('/', defaults={'spa_path': ''})
@app.get('/<path:spa_path>')
def serve_client(spa_path):
    from flask import abort, send_from_directory
    if spa_path.startswith('api/') or spa_path == 'api' or spa_path.startswith('embed/'):
        abort(404)
    if not _CLIENT_DIST.exists():
        abort(404)
    target = _CLIENT_DIST / spa_path
    if spa_path and target.is_file():
        return send_from_directory(_CLIENT_DIST, spa_path)
    return send_from_directory(_CLIENT_DIST, 'index.html')


# ─────────────────────────────────────────────────────────
# Entry point (restored — was lost in a prior working-tree edit)
# ─────────────────────────────────────────────────────────
def start_background_services():
    """Job worker + refresh scheduler. Called by the live server only —
    never at import, so tests keep full control."""
    import jobs
    jobs.ensure_worker(_new_conn)
    threading.Thread(target=scheduler_loop, daemon=True).start()


if __name__ == '__main__':
    init_db()
    start_background_services()
    print(f'\n[OK]  AnalytIQ API  ->  http://localhost:{PORT}\n')
    app.run(host='0.0.0.0', port=PORT, threaded=True, debug=False)
